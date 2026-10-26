"""النموذج الموحَّد model/wac.py يعيد أسعار أغسطس وسبتمبر الحالية نفسها، ويسمّى أى ملف شهر ناقص.

expected_prices.json مستخرج من المخرجات المعتمدة (بعد إعادة احتسابها بـ LibreOffice):
  تسعير جرد اخر شهر ٨ سنة ٢٠٢٦.xlsx، حساب متوسط التكلفة سبتمبر ٢٠٢٦.xlsx، تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx
"""
import json
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import wac as W  # noqa: E402

EXPECTED = json.loads((HERE / "expected_prices.json").read_text(encoding="utf-8"))
needs_pdftotext = pytest.mark.skipif(shutil.which("pdftotext") is None, reason="pdftotext (poppler) غير مثبت")


def same(a, b):
    return a == b or (isinstance(a, float) and isinstance(b, float) and abs(a - b) < 1e-9)


def check(got: dict, exp: dict):
    assert set(got) == set(exp)
    bad = [(k, got[k], v) for k, v in exp.items() if not same(got[k], v)]
    assert not bad, bad[:10]


@needs_pdftotext
@pytest.mark.parametrize("month", ["2026-08", "2026-09"])
def test_reproduces_current_prices(month):
    R, E = W.run(month), EXPECTED[month]
    check({x["name"]: x["wac"] for x in R["items"].values()}, E["wac"])
    check({f["name"]: f["price"] for f in R["fg"]}, E["fg"])
    if month == "2026-08":
        check(R["kg"], E["kg"])
    else:
        check(R["kg_open"], {k: v[0] for k, v in E["kg"].items()})
        check(R["kg"], {k: v[1] for k, v in E["kg"].items()})
        check({x["name"]: x["file_price"] for x in R["items"].values()}, E["file_price"])


@needs_pdftotext
def test_month_totals_match_the_reports():
    a, s = W.run("2026-08"), W.run("2026-09")
    assert (a["open_value"], a["printed"], a["raw_value"], a["fg_value"], a["total_value"]) == \
        (2016665.64, 2727152.34, 2699872.00, 435605.75, 3135477.74)
    assert (s["printed"], s["raw_value"], s["fg_value"], s["total_value"]) == \
        (2217803.14, 2202826.80, 319590.86, 2522417.67)


@needs_pdftotext
def test_priced_count_equals_the_issued_september_file(tmp_path):
    out = W.write_priced_count(W.run("2026-09"), tmp_path / "priced.xlsx")
    a = openpyxl.load_workbook(out).worksheets[0]
    for col, rows in EXPECTED["2026-09"]["priced_count"].items():
        bad = [(r, a[f"{col}{r}"].value, v) for r, v in rows.items() if a[f"{col}{r}"].value != v]
        assert not bad, bad[:5]
    issued = openpyxl.load_workbook(W.ROOT / "تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx").worksheets[0]
    assert [c.coordinate for row in issued.iter_rows() for c in row if a[c.coordinate].value != c.value] == []


@needs_pdftotext
@pytest.mark.parametrize("month", ["2026-08", "2026-09"])
def test_committed_outputs_are_reproducible(month, tmp_path):
    W.write_outputs(W.run(month), tmp_path)
    for f in tmp_path.iterdir():
        assert f.read_bytes() == (W.month_dir(month) / f.name).read_bytes(), f.name


def test_missing_files_are_named(tmp_path, monkeypatch):
    monkeypatch.setattr(W, "MONTHLY", tmp_path)
    monkeypatch.setattr(W, "REFERENCE", tmp_path / "reference")
    assert [p for p, _ in W.check_month("2026-10")] == [
        "monthly/reference/recipes.xlsx", "monthly/reference/june_prices.xlsx",
        "monthly/2026-10/month.json", "monthly/2026-10/purchases.pdf", "monthly/2026-10/stock_count.xlsx"]
    shutil.copytree(HERE.parents[1] / "monthly" / "reference", tmp_path / "reference")
    (tmp_path / "2026-10").mkdir()
    (tmp_path / "2026-10" / "month.json").write_text('{"month": "2026-10"}', encoding="utf-8")
    (tmp_path / "2026-10" / "purchases.pdf").write_bytes(b"")
    assert W.check_month("2026-10") == [("monthly/2026-10/stock_count.xlsx", W.REQUIRED["stock_count.xlsx"])]
    with pytest.raises(W.MissingFiles, match="stock_count.xlsx"):
        W.require("2026-10")
    assert W.main(["--check", "2026-10"]) == 1


@needs_pdftotext
def test_unmapped_purchase_code_stops_the_month(tmp_path, monkeypatch):
    """كود شراء جديد بلا ربط يوقف الشهر ويسمّى الكود — لا يُحمَّل على صنف تخمينًا."""
    src = HERE.parents[1] / "monthly"
    for d in ("reference", "2026-07", "2026-08", "2026-09"):
        shutil.copytree(src / d, tmp_path / d)
    cfg = json.loads((tmp_path / "2026-09" / "month.json").read_text(encoding="utf-8"))
    del cfg["new_codes"]["219"]
    (tmp_path / "2026-09" / "month.json").write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(W, "MONTHLY", tmp_path)
    monkeypatch.setattr(W, "REFERENCE", tmp_path / "reference")
    W.run.cache_clear()
    W.recipes.cache_clear()
    try:
        with pytest.raises(ValueError, match="كود 219"):
            W.run("2026-09")
    finally:
        W.run.cache_clear()
        W.recipes.cache_clear()
