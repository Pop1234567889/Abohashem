"""الهانك فى نموذج التكاليف (تأكيد الإدارة 26/10/2026).

كل كيلو سجق عليه 1/15 هانك، وأوراق الوصفات فى «سعر انتاج تام-1.xlsx» تضيف تكلفته إلى سعر
العجينة. MIX فى النموذج سعر عجينة فقط للهنا واسبشيال، فيُضاف الهانك لعبوات السجق وحدها
(لا للعجينة التى تشترك فى الوصفة)، ويبقى أبو هاشم كما هو لأن سعره يشمله.
"""
import contextlib
import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
with contextlib.redirect_stdout(io.StringIO()):
    import costing_pricing_model as M  # noqa: E402

STANDARD = HERE.parents[1] / "سعر انتاج تام-1.xlsx"


@pytest.fixture(scope="module")
def rows():
    return {r["code"]: r for r in M.build()[0]}


def _casing_cell(ws, label, header=None):
    """The number right of `label` (under `header` when given) in a recipe sheet."""
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    col = None
    if header:
        col = next(c.column for c in cells if c.value.strip() == header)
    lab = next(c for c in cells if c.value.strip() == label and (col is None or c.column == col))
    return next(float(ws.cell(lab.row, c).value) for c in range(lab.column + 1, lab.column + 4)
                if isinstance(ws.cell(lab.row, c).value, (int, float)))


def test_casing_per_kilo_is_the_standard_workbook():
    if not STANDARD.exists():
        pytest.skip("سعر انتاج تام-1.xlsx غير موجود")
    wb = openpyxl.load_workbook(STANDARD, data_only=True)
    assert _casing_cell(wb["الهنا"], "هانك") == pytest.approx(M.CASING_PER_KG["سجق الهنا"])
    assert _casing_cell(wb["اسبشيال 2026"], "هانك") == pytest.approx(M.CASING_PER_KG["سجق اسبشيال"])
    # أبو هاشم: the sheet adds 20/kg too, and MIX already carries it (91.08 vs the sheet's 90.28)
    assert _casing_cell(wb["الهاشم"], "هنكات للكيلو", "تكلفه سجق ابوهاشم") == 20
    assert "سجق ابوهاشم" not in M.CASING_PER_KG


def test_casing_goes_on_sausage_packs_only(rows):
    expected = {257: 100.0, 108: 6.0, 99: 14.0, 405: 5.0, 406: 250 / 15 * 0.7}
    for code, r in rows.items():
        assert r["std_casing_u"] == pytest.approx(expected.get(code, 0.0)), (code, r["name"])
    assert rows[292]["cat"] == "عجينة" and rows[292]["std_casing_u"] == 0      # paste shares the recipe
    for code in (257, 108, 99, 405, 406):
        r = rows[code]
        assert r["std_mat_u"] == pytest.approx(r["weight"] * M.MIX[r["recipe"]] + r["std_casing_u"])


def test_statement_still_ties_to_the_cent(rows):
    assert sum(r["act_mat_total"] for r in rows.values()) == pytest.approx(M.INCOME["cogs"], abs=0.01)
    assert sum(r["gross_profit"] for r in rows.values()) == pytest.approx(1_701_838.92, abs=0.01)
    assert sum(r["op_profit"] for r in rows.values()) == pytest.approx(418_841.92, abs=0.01)


def test_committed_model_rows_are_current(rows):
    committed = json.loads((HERE.parent / "model_rows.json").read_text(encoding="utf-8"))
    for code, r in rows.items():
        c = committed[str(code)]
        for k in ("std_casing_u", "std_mat_u", "act_mat_total", "gross_profit", "op_profit"):
            assert c[k] == pytest.approx(round(r[k], 4), abs=1e-4), (code, k)
