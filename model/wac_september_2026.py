#!/usr/bin/env python3
"""متوسط التكلفة المرجّح (WAC) للخامات آخر سبتمبر 2026 وتحديث أسعار ملف «جرد شهر 9».

المعادلة (دورية — Periodic WAC، استمرار سلسلة يونيه ← يوليو ← أغسطس):
    متوسط سبتمبر = (كمية جرد 31/8 × متوسط أغسطس + قيمة مشتريات سبتمبر)
                  ÷ (كمية جرد 31/8 + كمية مشتريات سبتمبر)

المدخلات (داخل المستودع):
    جر د شهر 8.xlsx                         الجرد الفعلى المعدَّل 31/8 = رصيد أول المدة لسبتمبر
    تسعير جرد اخر شهر ٨ سنة ٢٠٢٦-2.xlsx     تقييم أغسطس على الجرد المعدَّل (للمطابقة فقط)
    مشتريات شهر سبتمبر سنة 2026.pdf         تقرير المشتريات 01/09–30/09
    جرد شهر 9.xlsx                          الجرد الفعلى 30/9 — الملف الذى تُدخل فيه الأسعار
    + مدخلات wac_august_2026.py (يُعاد احتساب متوسط أغسطس منها باستقلال)

المخرجات:
    تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx         نسخة من «جرد شهر 9.xlsx» بأسعار سبتمبر — لا يتغير فيها غير خانات السعر
    حساب متوسط التكلفة سبتمبر ٢٠٢٦.xlsx        الاحتساب الكامل بمعادلات حيّة + المطابقة
    WAC_SEPTEMBER_2026_REPORT.md             تقرير المراجعة

التشغيل:  python3 model/wac_september_2026.py   (يتطلب openpyxl و pdftotext و LibreOffice)
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

import wac_august_2026 as A
from wac_august_2026 import (FLAG, NO_MATCH, NUM, PCT, POOL, POOL_MEMBERS, PRICE, SUBH, TOT, clean, header, key, num,
                             put, q, recipe_cost, sheet, title_rows)

ROOT = A.ROOT
F_CNT8 = ROOT / "جر د شهر 8.xlsx"
F_AUG2 = ROOT / "تسعير جرد اخر شهر ٨ سنة ٢٠٢٦-2.xlsx"
F_PUR9 = ROOT / "مشتريات شهر سبتمبر سنة 2026.pdf"
F_CNT9 = ROOT / "جرد شهر 9.xlsx"
OUT_FILE = ROOT / "تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx"
OUT_CALC = ROOT / "حساب متوسط التكلفة سبتمبر ٢٠٢٦.xlsx"
OUT_MD = ROOT / "WAC_SEPTEMBER_2026_REPORT.md"

# اسم الصنف بجرد 31/8 المعدَّل ← اسمه فى احتساب أغسطس (الجرد الأول)
AUG_ALIASES = {"برجر أبو هاشم 1ك": "اكياس برجر ابوهاشم 1ك"}
# اسم الصنف بجرد 31/8 المعدَّل ← اسمه بملف «-2»
AUG2_ALIASES = {"فلانك": "فلانك + مجمل رصيد اللحمة بالنفق", "كروت كبده شرائح 350": "كروت كبده علب 350",
                "برجر أبو هاشم 1ك": "اكياس برجر ابوهاشم 1ك", "أكياس بانيه الهنا 5 ك": "أكياس بانيه استربس 5 الهنا",
                "فلتر هواء": "فلتر هواء دبابة"}
# اسم الصنف بجرد سبتمبر ← اسمه بجرد 31/8 المعدَّل (نفس الصنف بتسمية مختلفة؛ نفس تسمية ملف «-2»)
SEP_ALIASES = {v: k for k, v in AUG2_ALIASES.items() if k != "فلانك"}
# أكواد ظهرت لأول مرة فى سبتمبر: (الصنف المقابل بالجرد، درجة الثقة، أساس المطابقة)
SEP_NEW_CODES = {
    "328": ("زيت سجق", "عالية", "طعم سجق «زيت» = بند «زيت سجق»؛ سعر 580 = متوسط أغسطس 580"),
    "329": ("زيوت بسطرمه", "عالية", "طعم بسطرمة «زيت» = بند «زيوت بسطرمه»؛ سعر 820 ≈ 784.14 لأغسطس"),
    "379": ("لازق", "عالية", "بكرة لزق شفاف = «لازق»؛ سعر 30.83 ≈ 30.22 لأغسطس"),
    "443": ("أكياس استربس اسبيشيال ص", "عالية", "أكياس تغليف استربس اسبشيال صغير = «أكياس استربس اسبيشيال ص»؛ بلا سعر سابق"),
    "409": ("طبق مسطح 1 ك شفاف", "متوسطة", "طبق شفاف كيلو = البند الشفاف الوحيد مقاس 1 ك بالجرد؛ بلا سعر سابق — يُراجَع"),
    "344": ("فلفل رومي حامي", "متوسطة", "فلفل حامى = عائلة الفلفل الحار الوحيدة بالجرد (نفس معالجة «فلفل احمر» كود 187)"),
    "219": ("جنزبيل", "عالية", "تطابق اسمى؛ سعر 220 ≈ 211.31 لأغسطس"),
    "68": ("نشا", "عالية", "شيكارة نشا = «نشا»؛ سعر 28.5 = سعر كود 191 فى نفس الشهر"),
    "371": ("مفروم فراخ شعبي", "متوسطة", "الرتبة B غير موجودة بالجرد؛ سعرها 34 داخل نطاق الرتبة C (30–35) فحُمّلت على «شعبي» — يُراجَع"),
    "189": ("الملح", "عالية", "تطابق اسمى؛ سعر 2.3–3.0 ≈ 2.06 لأغسطس"),
}
SEP_UNMATCHED = {
    "453": "بابريكا مدخنة صنف مختلف عن «بابريكا» العادية ولا يوجد له بند بالجرد — لم تُحمَّل على أى صنف",
    "442": "أكياس تغليف «مقطعات» 5 ك — لا يوجد بند مقابل بالجرد (أكياس 5 ك الموجودة مخصصة لسجق/بانيه/عجينة) — لم تُحمَّل",
}
SUSPECT_INVOICE = "1111115217"
# أكدت الإدارة (20/10/2026) أن سعرَى الرتبتين بالفاتورة معكوسان: C (كود 372) = 30 و A (كود 370) = 57
INVOICE_CORRECTIONS = {(SUSPECT_INVOICE, "372"): 30.0, (SUSPECT_INVOICE, "370"): 57.0}
CORRECTION_NOTE = ("فاتورة 1111115217 (09/09) سجّلت «مفروم فراخ C» بسعر 57 و«مفروم فراخ A» بسعر 30 (خارج نطاق كل رتبة)؛ "
                   "أكدت الإدارة أن السعرين معكوسان فاحتُسبت C = 30 و A = 57")
CORRECTED_ITEMS = ("مفروم فراخ شعبي", "مفروم فراخ فاخر")
UNRESOLVED = {}     # صنف معلّق ⇒ يُترك سعره بالملف وتُحسب الوصفات بسعر أغسطس — لا أصناف معلّقة حاليًا
# أسطر سعرها شاذ عن باقى مشتريات نفس الصنف بالشهر — تدخل الاحتساب كما سُجّلت مع التنويه
PRICE_OUTLIERS = {(76, "190"): "صميد 50 كجم بسعر 26.40 مقابل 16.25–17.00 لباقى الأسطر",
                  (58, "87"): "جلتومات 25 كجم بسعر 132 مقابل 80 لسطر الشراء الآخر"}
FG_ALIASES_SEP = {"استربس 1kg": "استربس 1 ك", "كفته أبو هاشم 1 ك": "كفتة ابو هاشم 1ك"}


# ---------------------------------------------------------------- المدخلات
def read_count(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    raw, fg = [], []
    for r in range(5, ws.max_row + 1):
        n = clean(ws.cell(r, 1).value)
        if n and not n.startswith(("إجمالى", "اجمالي")):
            locs = [num(ws.cell(r, c).value) or 0.0 for c in range(2, 7)]
            g = ws.cell(r, 7).value
            qty = sum(locs)
            issue = None
            if isinstance(g, str) and g.replace(" ", "") != f"=SUM(B{r}:F{r})" and abs(sum(locs[:2]) - qty) > 1e-9:
                issue = f"معادلة الإجمالى «{g}» تُسقط مواقع بها رصيد — اعتُمد مجموع المواقع {qty:g}"
            elif not isinstance(g, str) and abs((num(g) or 0.0) - qty) > 1e-9:
                issue = f"الإجمالى مُدخل يدويًا ({num(g) or 0:g}) بينما مجموع المواقع {qty:g} — اعتُمد مجموع المواقع"
            raw.append({"row": r, "name": n, "qty": qty, "issue": issue, "orig_price": ws.cell(r, 8).value})
        f = ws.cell(r, 10).value if ws.max_column >= 10 else None
        if f not in (None, 0) and clean(f) and clean(f) != "اسم الصنف":
            fg.append({"row": r, "name": clean(f), "qty": num(ws.cell(r, 11).value) or 0.0,
                       "orig_price": ws.cell(r, 12).value})
    return raw, fg


def august_baseline():
    """يعيد احتساب متوسط أغسطس من مصادره (wac_august_2026) — مستقل عن ملف «-2»."""
    jul, codes, jul_fg, kg_tab = A.load_july()
    pur8, _ = A.load_purchases(A.F_PUR)
    cnt8_old, _ = A.load_count()
    items8, pool8 = A.compute_wac(jul, codes, pur8, cnt8_old)
    return {k: x["wac"] for k, x in items8.items()}, pool8["wac"], jul, codes, jul_fg, kg_tab


def load_aug2():
    wb = openpyxl.load_workbook(F_AUG2, data_only=True)
    raw = {}
    for r in wb["جرد الخامات آخر أغسطس"].iter_rows(min_row=5, values_only=True):
        if r[0] and not clean(r[0]).startswith(("إجمالى", "اجمالي")):
            raw[key(r[0])] = {"name": clean(r[0]), "wac": num(r[8]), "qty": num(r[9]) or 0.0}
    fg = {}
    for r in wb["جرد المنتج التام آخر أغسطس"].iter_rows(min_row=4, values_only=True):
        if r[0] and not clean(r[0]).startswith(("إجمالى", "اجمالي")):
            fg[key(r[0])] = {"name": clean(r[0]), "price": num(r[3]), "note": clean(r[6])}
    return raw, fg


# ---------------------------------------------------------------- الاحتساب
def build_opening(aug_wac, aug2_raw):
    """رصيد أول المدة لسبتمبر = جرد 31/8 المعدَّل × متوسط أغسطس، مع مطابقة ملف «-2»."""
    cnt8, _ = read_count(F_CNT8)
    opening, checks = OrderedDict(), []
    for c in cnt8:
        k = key(c["name"])
        mine = aug_wac.get(key(AUG_ALIASES.get(c["name"], c["name"])))
        a2 = aug2_raw.get(key(AUG2_ALIASES.get(c["name"], c["name"])))
        if mine is not None and a2 and a2["wac"] is not None and abs(mine - a2["wac"]) > 0.00005:
            checks.append(("سعر", c["name"], mine, a2["wac"]))
        if a2 and abs(a2["qty"] - c["qty"]) > 0.0051:      # ملف «-2» يقرّب الكمية لرقمين
            checks.append(("كمية", c["name"], c["qty"], a2["qty"]))
        price = mine if mine is not None else (a2["wac"] if a2 else None)
        src = "متوسط أغسطس (إعادة احتساب)" if mine is not None else ("ملف «-2»" if price is not None else "")
        opening[k] = {"name": c["name"], "qty": c["qty"], "price": price, "src": src, "issue": c["issue"],
                      "aug2_qty": a2["qty"] if a2 else None, "aug2_wac": a2["wac"] if a2 else None}
    return opening, checks


def correct_purchases(purchases):
    """يطبّق تصحيح الأسعار المؤكَّد على أسطر الفاتورة، مع الإبقاء على المسجّل بالتقرير للمطابقة."""
    for p in purchases:
        p["total_recorded"] = p["total"]
        new = INVOICE_CORRECTIONS.get((p["invoice"], p["code"]))
        if new is not None:
            assert abs(p["qty"] * p["price"] - p["total"]) < 0.01, p
            p.update(price_recorded=p["price"], price=new, total=round(p["qty"] * new, 2), corrected=True)
    assert sum(p.get("corrected", False) for p in purchases) == len(INVOICE_CORRECTIONS)


def map_purchases(purchases, codes, targets):
    for p in purchases:
        if p["code"] in SEP_UNMATCHED:
            p.update(item=NO_MATCH, conf="—", basis=SEP_UNMATCHED[p["code"]])
            continue
        if p["code"] in codes:
            item, conf, basis = codes[p["code"]]
        elif p["code"] in A.NEW_CODES:
            item, conf, basis = A.NEW_CODES[p["code"]]
            basis = "نفس ربط الكود فى أغسطس: " + basis
        else:
            item, conf, basis = SEP_NEW_CODES[p["code"]]
        if item != POOL:
            assert key(item) in targets, f"code {p['code']} maps to {item}, not in the September count"
            item = targets[key(item)]["name"]
        note = PRICE_OUTLIERS.get((p["idx"], p["code"]))
        if p.get("corrected"):
            note = (f"سعر مُصحّح: سُجّل بالتقرير {p['price_recorded']:g} (إجمالى {p['total_recorded']:,.2f}) — "
                    "أكدت الإدارة أن سعرَى الرتبتين بالفاتورة معكوسان")
        p.update(item=item, conf=conf, basis=basis + (f" | ⚠ {note}" if note else ""), outlier=bool(note))


def wac(open_q, open_p, pq, pv):
    if pq > 0:
        return round(pv / pq, 4) if open_p is None else round((open_q * open_p + pv) / (open_q + pq), 4)
    return open_p


def compute(sep_raw, opening, purchases):
    items = OrderedDict()
    for c in sep_raw:
        o = opening.get(key(SEP_ALIASES.get(c["name"], c["name"])))
        items[key(c["name"])] = {**c, "open_q": o["qty"] if o else 0.0, "open_p": o["price"] if o else None,
                                 "open_name": o["name"] if o else None, "pq": 0.0, "pv": 0.0, "n": 0}
    pool = {"name": POOL, "pq": 0.0, "pv": 0.0, "n": 0,
            "open_q": sum(items[key(m)]["open_q"] for m in POOL_MEMBERS),
            "open_p": items[key(POOL_MEMBERS[0])]["open_p"]}
    assert len({items[key(m)]["open_p"] for m in POOL_MEMBERS}) == 1
    for p in purchases:
        if p["item"] == NO_MATCH:
            continue
        t = pool if p["item"] == POOL else items[key(p["item"])]
        t["pq"] += p["qty"]
        t["pv"] += p["total"]
        t["n"] += 1
        t["corrected"] = t.get("corrected", False) or p.get("corrected", False)
    pool["wac"] = wac(pool["open_q"], pool["open_p"], pool["pq"], pool["pv"])
    for x in items.values():
        x["wac"] = pool["wac"] if x["name"] in POOL_MEMBERS else wac(x["open_q"], x["open_p"], x["pq"], x["pv"])
        x["unresolved"] = UNRESOLVED.get(x["name"])
        x["file_price"] = x["orig_price"] if x["unresolved"] else (x["wac"] if x["wac"] is not None else x["orig_price"])
        x["status"] = status(x)
    return items, pool


def status(x):
    if x["unresolved"]:
        return "⚠ لم يُغيَّر السعر — بيانات شراء متعارضة"
    if x["name"] in POOL_MEMBERS:
        return "متوسط مجمّع اللحوم الهندى"
    if x["wac"] is None:
        return "بلا سعر — لا سعر أغسطس ولا مشتريات سبتمبر" + (" (السعر الموجود تُرك)" if x["orig_price"] else "")
    if x["pq"] > 0 and x["open_p"] is None:
        return "المتوسط من مشتريات سبتمبر فقط (أول المدة بلا سعر)" if x["open_q"] else "المتوسط = متوسط سعر الشراء"
    if x.get("corrected"):
        return "قياسى: (أول المدة + المشتريات) بعد تصحيح سعرَى الفاتورة 1111115217"
    if x["pq"] > 0:
        return "قياسى: (أول المدة + المشتريات)"
    return "لا مشتريات فى سبتمبر ⇒ سعر أغسطس ثابت"


SCN_APPLIED = "بعكس سعرَى الفاتورة (C=30، A=57) — المعتمد"


def scenarios(items, purchases):
    """متوسط صنفَى الفراخ المفرومة فى قراءات الفاتورة الثلاث — للإفصاح عن أثر التصحيح المعتمد."""
    out = {}
    for nm in CORRECTED_ITEMS:
        x = items[key(nm)]
        lines = [p for p in purchases if p["item"] == nm]
        sus = [p for p in lines if p["invoice"] == SUSPECT_INVOICE]
        rest_q = sum(p["qty"] for p in lines if p not in sus)
        rest_v = sum(p["total"] for p in lines if p not in sus)
        sus_q = sum(p["qty"] for p in sus)
        oq, op = x["open_q"], x["open_p"]
        out[nm] = {
            "كما سُجّلت بالتقرير": wac(oq, op, x["pq"], rest_v + sum(p["total_recorded"] for p in sus)),
            "باستبعاد الفاتورة 1111115217": wac(oq, op, rest_q, rest_v),
            SCN_APPLIED: wac(oq, op, rest_q + sus_q, rest_v + sum(p["total"] for p in sus)),
            "_aug": op, "_orig": x["orig_price"], "_file": x["file_price"], "_close_q": x["qty"],
        }
        assert out[nm][SCN_APPLIED] == x["wac"] == x["file_price"], (nm, out[nm], x["wac"])
    return out


# ---------------------------------------------------------------- المنتج التام
def price_fg(sep_fg, jul_fg, aug2_fg, kg_sep, kg_aug, recipes, provisional_recipes):
    rows = []
    for f in sep_fg:
        alias = FG_ALIASES_SEP.get(f["name"]) or (A.FG_ALIASES.get(f["name"]) or (None,))[0]
        j = jul_fg.get(key(f["name"])) or (jul_fg.get(key(alias)) if alias else None)
        a2 = aug2_fg.get(key(f["name"])) or (aug2_fg.get(key(alias)) if alias else None)
        row = {**f, "aug_price": a2["price"] if a2 else None, "rule": None, "price": None, "flag": False}
        if j and j["rule"]:
            rn, w, pkg = j["rule"]
            row.update(rule=j["rule"], price=round(w * kg_sep[rn] + pkg, 4),
                       note=f"من وصفة «{rn}»: {w:g} كجم × سعر كجم سبتمبر" +
                       (f" + تعبئة {pkg:.2f} ج.م" if pkg else " بدون تكلفة تعبئة (غير محددة بالوصفة)"))
            if a2 and a2["price"] is not None:          # فرق سعر أغسطس عن ملف «-2» (للإفصاح)
                row["aug_dev"] = round(round(w * kg_aug[rn] + pkg, 4) - a2["price"], 4)
            if rn in provisional_recipes:
                row["note"] += " | مؤقت: الوصفة تستخدم مفروم الفراخ بسعر أغسطس لحين حسم الفاتورة"
        elif a2 and a2["price"] is not None:
            row.update(price=a2["price"], note="خارج الوصفات التسع — آخر سعر متاح (أغسطس، ملف «-2»): " + a2["note"][:90])
        elif j and j["price"] is not None:
            row.update(price=j["price"], note="خارج الوصفات التسع — آخر سعر متاح (يوليو)")
        else:
            row.update(note="بلا وصفة ولا سعر سابق — يلزم توفير سعره", flag=f["qty"] > 0)
        if alias:
            row["note"] += f" | نفس صنف «{alias}»"
        rows.append(row)
    return rows


# ---------------------------------------------------------------- تحديث ملف سبتمبر (خانات السعر فقط)
def update_september_file(items, fg_rows):
    """يكتب الأسعار فى XML ورقة «جرد شهر 9» مباشرة — لا يُعاد حفظ أى خانة أخرى (لا تغيير ولو بفرق تقريب).

    * عمود H (سعر الخامات) و L (سعر المنتج التام): القيمة الجديدة بنفس تنسيق الخانة.
    * عمود M (قيمة المنتج التام) فارغ بالملف الأصلى؛ تُضاف له =K×L فى صفوف المنتج المُسعّر فقط
      (نفس نمط =G×H فى قسم الخامات) حتى يُقيَّم المنتج التام.
    * تُحذف القيم المخزّنة لنتائج المعادلات و calcChain ويُطلب إعادة الاحتساب عند الفتح.
    """
    import re
    import zipfile
    sets, adds = {}, {}
    for x in items.values():
        if not x["unresolved"] and x["wac"] is not None and x["orig_price"] != x["wac"]:
            sets[f"H{x['row']}"] = (x["orig_price"], x["wac"])
    for f in fg_rows:
        if f["price"] is not None and f["orig_price"] != f["price"]:
            sets[f"L{f['row']}"] = (f["orig_price"], f["price"])
            adds[f"M{f['row']}"] = f"K{f['row']}*L{f['row']}"
    zin = zipfile.ZipFile(F_CNT9)
    parts = {n: zin.read(n) for n in zin.namelist()}
    sh = parts["xl/worksheets/sheet1.xml"].decode("utf-8")

    def cell_re(ref):
        return re.compile(r'<c r="%s"((?: [a-z]+="[^"]*")*)(?:/>|>(.*?)</c>)' % ref)

    def style(attrs):
        m = re.search(r' s="(\d+)"', attrs)
        return f' s="{m.group(1)}"' if m else ""

    for ref, (_, val) in sets.items():
        m = cell_re(ref).search(sh)
        assert m and "<f>" not in (m.group(2) or ""), ref        # الخانة موجودة وليست معادلة
        sh = sh[:m.start()] + f'<c r="{ref}"{style(m.group(1))}><v>{val!r}</v></c>' + sh[m.end():]
    for ref, formula in adds.items():
        m = cell_re(ref).search(sh)
        assert m and not (m.group(2) or "").strip(), ref           # خانة فارغة
        sh = sh[:m.start()] + f'<c r="{ref}"{style(m.group(1))}><f>{formula}</f></c>' + sh[m.end():]
    sh = re.sub(r"(<f>[^<]*</f>)<v>[^<]*</v>", r"\1", sh)          # لا نتائج قديمة مخزّنة
    parts["xl/worksheets/sheet1.xml"] = sh.encode("utf-8")
    wbx = parts["xl/workbook.xml"].decode("utf-8")
    wbx = re.sub(r"<calcPr([^>]*?)/>", lambda m: "<calcPr" + m.group(1).replace(' fullCalcOnLoad="1"', "")
                 + ' fullCalcOnLoad="1"/>', wbx)
    parts["xl/workbook.xml"] = wbx.encode("utf-8")
    parts.pop("xl/calcChain.xml", None)
    parts["xl/_rels/workbook.xml.rels"] = re.sub(rb'<Relationship [^>]*calcChain[^>]*/>', b"",
                                                 parts["xl/_rels/workbook.xml.rels"])
    parts["[Content_Types].xml"] = re.sub(rb'<Override [^>]*calcChain[^>]*/>', b"", parts["[Content_Types].xml"])
    with zipfile.ZipFile(OUT_FILE, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            if n in parts:
                zout.writestr(zin.getinfo(n), parts[n])
    # التحقق: كل خانة غير المقصودة بقيمتها ومعادلتها وتنسيقها كما هى تمامًا
    a = openpyxl.load_workbook(F_CNT9).worksheets[0]
    b = openpyxl.load_workbook(OUT_FILE).worksheets[0]
    assert a.dimensions == b.dimensions and str(a.merged_cells) == str(b.merged_cells)
    for row in a.iter_rows():
        for c in row:
            d = b[c.coordinate]
            if c.coordinate in sets:
                assert d.value == sets[c.coordinate][1]
            elif c.coordinate in adds:
                assert c.value is None and d.value == "=" + adds[c.coordinate]
            else:
                assert d.value == c.value, (c.coordinate, c.value, d.value)
            assert d.style_id == c.style_id, c.coordinate
    return [(r, o, v) for r, (o, v) in sets.items()], adds


# ---------------------------------------------------------------- مصنّف الاحتساب
S_MAIN, S_PUR, S_OPEN = "احتساب متوسط سبتمبر", "مشتريات سبتمبر التفصيلية", "رصيد أول المدة (31-8)"
S_BOM, S_KG, S_FG = "تفصيل الوصفات بأسعار سبتمبر", "سعر الكجم من الوصفات", "المنتج التام آخر سبتمبر"
S_SCN, S_REC, S_NOTE = "الفراخ المفرومة (التصحيح)", "المطابقة", "المنهجية والاستثناءات"


def build_calc(items, pool, purchases, printed, opening, recipes, kg_sep, kg_aug, kg_jul_tab, fg_rows, scn, prov_items):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # المشتريات
    wp = sheet(wb, S_PUR, [5, 11, 8, 30, 7, 10, 10, 13, 26, 9, 60, 12, 14])
    title_rows(wp, "مشتريات الخامات — الفترة 01/09/2026 حتى 30/09/2026 (المصدر: تقرير المشتريات فى الفترة)",
               "عمودا السعر والإجمالى بعد تصحيح الفاتورة 1111115217؛ العمود الأخير = الإجمالى كما سُجّل بالتقرير.", 13)
    header(wp, 3, ["م", "التاريخ", "كود الصنف", "اسم الصنف بالمشتريات", "الوحدة", "الكمية", "سعر الشراء",
                   "الإجمالى", "الصنف المقابل بالجرد", "درجة الثقة", "أساس المطابقة", "رقم الفاتورة",
                   "الإجمالى المسجّل بالتقرير"])
    for i, p in enumerate(purchases):
        r = 4 + i
        fl = FLAG if p["item"] == NO_MATCH or p["conf"] == "متوسطة" or p.get("outlier") else None
        for c, v in enumerate([p["idx"], p["date"], p["code"], p["name"], p["unit"], p["qty"], p["price"], p["total"],
                               p["item"], p["conf"], p["basis"], p["invoice"], p["total_recorded"]], 1):
            put(wp, r, c, v, NUM if c in (6, 8, 13) else PRICE if c == 7 else None, fl)
    pf, pl = 4, 3 + len(purchases)
    pt = pl + 2
    put(wp, pt, 4, "إجمالى مشتريات سبتمبر (بعد التصحيح)", fill=TOT, bold=True)
    put(wp, pt, 8, f"=SUM(H{pf}:H{pl})", NUM, TOT, True)
    put(wp, pt, 13, f"=SUM(M{pf}:M{pl})", NUM, TOT, True)
    put(wp, pt + 1, 4, "الإجمالى المطبوع بالتقرير", fill=TOT)
    put(wp, pt + 1, 13, printed, NUM, TOT)
    put(wp, pt + 2, 4, "منها: بلا مقابل بالجرد (مستبعدة)", fill=TOT)
    put(wp, pt + 2, 8, f'=SUMIF(I{pf}:I{pl},"{NO_MATCH}",H{pf}:H{pl})', NUM, TOT)
    put(wp, pt + 3, 4, "أثر تصحيح الفاتورة 1111115217 (بعد − كما سُجّل)", fill=TOT)
    put(wp, pt + 3, 8, f"=H{pt}-M{pt}", NUM, TOT)
    IR, QR, VR = (f"{q(S_PUR)}!$I${pf}:$I${pl}", f"{q(S_PUR)}!$F${pf}:$F${pl}", f"{q(S_PUR)}!$H${pf}:$H${pl}")

    # رصيد أول المدة
    wo = sheet(wb, S_OPEN, [32, 13, 14, 14, 13, 14, 12, 50])
    title_rows(wo, "رصيد أول المدة لسبتمبر = الجرد الفعلى المعدَّل 31/8 («جر د شهر 8.xlsx») × متوسط أغسطس",
               "متوسط أغسطس أُعيد احتسابه من مصادره (جرد 31/7 × متوسط يوليو + مشتريات أغسطس) وطوبق مع ملف «تسعير جرد اخر شهر ٨ -2».", 8)
    header(wo, 4, ["الصنف بجرد 31/8", "كمية الجرد 31/8", "متوسط أغسطس (إعادة احتساب)", "متوسط أغسطس بملف «-2»",
                   "كمية الصنف بملف «-2»", "قيمة أول المدة", "مطابقة السعر", "ملاحظة"])
    open_row, r = {}, 5
    for k, o in opening.items():
        open_row[k] = r
        put(wo, r, 1, o["name"])
        put(wo, r, 2, o["qty"], NUM)
        put(wo, r, 3, o["price"], PRICE)
        put(wo, r, 4, o["aug2_wac"], PRICE)
        put(wo, r, 5, o["aug2_qty"], NUM)
        put(wo, r, 6, f'=IF(C{r}="","",B{r}*C{r})', NUM)
        put(wo, r, 7, f'=IF(OR(C{r}="",D{r}=""),"",IF(ABS(C{r}-D{r})<0.00005,"✔","✘"))')
        put(wo, r, 8, "; ".join(t for t in (o["issue"], "بملف «-2» وجرد سبتمبر باسم «" + AUG2_ALIASES[o["name"]] + "»"
                                             if o["name"] in AUG2_ALIASES else None) if t))
        r += 1
    open_tot = r + 1
    put(wo, open_tot, 1, "إجمالى قيمة أول المدة (= مخزون الخامات آخر أغسطس)", fill=TOT, bold=True)
    put(wo, open_tot, 6, f"=SUM(F5:F{r - 1})", NUM, TOT, True)
    put(wo, open_tot + 1, 1, "إجمالى ملف «-2» (للمطابقة)", fill=TOT)

    # الورقة الرئيسية
    wm = sheet(wb, S_MAIN, [30, 12, 12, 13, 12, 14, 12, 14, 13, 13, 12, 14, 13, 44])
    wb.move_sheet(wm, -(len(wb.sheetnames) - 1))
    title_rows(wm, "ABOHASHEM FOR MEAT — متوسط التكلفة المرجّح للخامات آخر سبتمبر 2026 وتقييم جرد 30/9",
               "المتوسط = (كمية جرد 31/8 × متوسط أغسطس + قيمة مشتريات سبتمبر) ÷ (كمية جرد 31/8 + كمية مشتريات سبتمبر). "
               "عمود «السعر المُدخل بالملف» هو ما وُضع فى «تسعير جرد اخر شهر ٩».", 14)
    header(wm, 4, ["اسم الصنف", "كمية أول المدة (جرد 31/8)", "سعر أول المدة (متوسط أغسطس)", "قيمة أول المدة",
                   "كمية مشتريات سبتمبر", "قيمة مشتريات سبتمبر", "إجمالى الكمية المتاحة", "إجمالى القيمة المتاحة",
                   "متوسط السعر (WAC) 30/9", "السعر المُدخل بالملف", "كمية جرد 30/9", "قيمة مخزون 30/9",
                   "التغير عن أغسطس %", "الحالة / ملاحظة"])
    row_of, r = {}, 5
    for k, x in items.items():
        row_of[k] = r
        fl = FLAG if x["unresolved"] or (x["wac"] is None and x["qty"] > 0) or x["issue"] else None
        put(wm, r, 1, x["name"], fill=fl)
        put(wm, r, 2, x["open_q"], NUM)
        put(wm, r, 3, x["open_p"], PRICE)
        put(wm, r, 4, f'=IF(C{r}="","",B{r}*C{r})', NUM)
        put(wm, r, 5, f"=SUMIF({IR},A{r},{QR})", NUM)
        put(wm, r, 6, f"=SUMIF({IR},A{r},{VR})", NUM)
        put(wm, r, 7, f'=IF(C{r}="",E{r},B{r}+E{r})', NUM)
        put(wm, r, 8, f"=N(D{r})+F{r}", NUM)
        put(wm, r, 9, "=POOL" if x["name"] in POOL_MEMBERS else f'=IF(G{r}>0,ROUND(H{r}/G{r},4),IF(C{r}="","",C{r}))',
            PRICE, SUBH)
        put(wm, r, 10, x["file_price"], PRICE, FLAG if x["unresolved"] else None, True)
        put(wm, r, 11, x["qty"], NUM)
        put(wm, r, 12, f'=IF(J{r}="","",K{r}*J{r})', NUM)
        put(wm, r, 13, f'=IF(OR(C{r}="",I{r}=""),"",I{r}/C{r}-1)', PCT)
        put(wm, r, 14, "; ".join(t for t in (x["status"], x["unresolved"], x["issue"]) if t), fill=fl)
        r += 1
    m_last = r - 1
    m_tot = r + 1
    put(wm, m_tot, 1, "الإجمالى", fill=TOT, bold=True)
    for c in (4, 6, 8, 12):
        Lc = openpyxl.utils.get_column_letter(c)
        put(wm, m_tot, c, f"=SUM({Lc}5:{Lc}{m_last})", NUM, TOT, True)
    pr = m_tot + 3
    put(wm, pr - 1, 1, "مجمّع اللحوم الهندى (كتف/رقبه/سن/أرباع) — متوسط واحد للبنود الأربعة", fill=SUBH, bold=True)
    put(wm, pr, 1, POOL)
    put(wm, pr, 2, "=" + "+".join(f"B{row_of[key(m)]}" for m in POOL_MEMBERS), NUM)
    put(wm, pr, 3, f"=C{row_of[key(POOL_MEMBERS[0])]}", PRICE)
    put(wm, pr, 4, f"=B{pr}*C{pr}", NUM)
    put(wm, pr, 5, f'=SUMIF({IR},"{POOL}",{QR})', NUM)
    put(wm, pr, 6, f'=SUMIF({IR},"{POOL}",{VR})', NUM)
    put(wm, pr, 7, f"=B{pr}+E{pr}", NUM)
    put(wm, pr, 8, f"=D{pr}+F{pr}", NUM)
    put(wm, pr, 9, f"=ROUND(H{pr}/G{pr},4)", PRICE, SUBH, True)
    for m in POOL_MEMBERS:
        wm.cell(row_of[key(m)], 9).value = f"=I{pr}"

    # تفصيل الوصفات
    wbm = sheet(wb, S_BOM, [18, 12, 22, 24, 10, 12, 13, 50])
    title_rows(wbm, "تفصيل وصفات التصنيع بأسعار خامات سبتمبر (نفس كميات ومعاملات إنتاجية وصفات يونيه)", None, 8)
    header(wbm, 3, ["الوصفة", "القسم", "الخامة بالوصفة", "الصنف المقابل بالجرد", "الكمية (كجم)", "سعر سبتمبر",
                    "القيمة", "ملاحظة"])
    r, bom_tot = 4, {}
    for rn, rc in recipes.items():
        sub_cell = None
        for sec, label in (("sub", "خلطة فرعية (توابل)"), ("main", "الوصفة الرئيسية")):
            if not rc[sec]:
                continue
            first = r
            for i in rc[sec]:
                it = i["item"]
                put(wbm, r, 1, rn)
                put(wbm, r, 2, label)
                put(wbm, r, 3, i["ing"])
                put(wbm, r, 5, i["qty"], "0.000")
                note = i["basis"]
                if it is None:
                    put(wbm, r, 4, "—")
                    put(wbm, r, 6, 0, PRICE)
                elif isinstance(it, tuple):
                    put(wbm, r, 4, "— (سعر ثابت بالوصفة)")
                    put(wbm, r, 6, it[1], PRICE)
                elif it == "@SUB":
                    put(wbm, r, 4, "سعر كجم الخلطة الفرعية")
                    put(wbm, r, 6, f"={sub_cell}", PRICE)
                elif it == POOL:
                    put(wbm, r, 4, POOL)
                    put(wbm, r, 6, f"={q(S_MAIN)}!I{pr}", PRICE)
                else:
                    put(wbm, r, 4, it)
                    rr = row_of[key(it)]
                    if it in UNRESOLVED:     # معلّق ⇒ بسعر أغسطس (عمود C) لحين الحسم
                        put(wbm, r, 6, f"={q(S_MAIN)}!C{rr}", PRICE, FLAG)
                        note = "مؤقت: سعر أغسطس لحين حسم فاتورة الفراخ المفرومة"
                    else:
                        put(wbm, r, 6, f"={q(S_MAIN)}!I{rr}", PRICE)
                put(wbm, r, 7, f"=E{r}*F{r}", NUM)
                put(wbm, r, 8, note, fill=FLAG if it is None or isinstance(it, tuple) or it in UNRESOLVED else None)
                r += 1
            put(wbm, r, 1, rn, fill=SUBH)
            put(wbm, r, 2, f"إجمالى {label}", fill=SUBH, bold=True)
            put(wbm, r, 5, f"=SUM(E{first}:E{r - 1})", "0.000", SUBH)
            put(wbm, r, 7, f"=SUM(G{first}:G{r - 1})", NUM, SUBH, True)
            if sec == "sub":
                put(wbm, r, 6, f"=G{r}/E{r}", PRICE, SUBH, True)
                sub_cell = f"F{r}"
            else:
                bom_tot[rn] = f"{q(S_BOM)}!G{r}"
                put(wbm, r, 8, f"سعر الكجم = قيمة الوصفة ÷ {rc['div']:g} كجم", fill=SUBH)
            r += 2

    # سعر الكجم
    wk = sheet(wb, S_KG, [16, 10, 16, 16, 13, 13, 10, 14])
    title_rows(wk, "سعر كيلو الخلطة لكل منتج (خامات مباشرة فقط) — أغسطس مقابل سبتمبر", None, 8)
    header(wk, 3, ["المنتج", "العلامة", "قيمة الوصفة (أغسطس)", "قيمة الوصفة (سبتمبر)", "سعر الكجم (أغسطس)",
                   "سعر الكجم (سبتمبر)", "التغير %", "معامل الإنتاجية (كجم)"])
    kg_row, r = {}, 4
    for rn, rc in recipes.items():
        kg_row[rn] = r
        put(wk, r, 1, rn)
        put(wk, r, 2, kg_jul_tab[rn]["brand"])
        put(wk, r, 3, kg_aug[rn] * rc["div"], NUM)
        put(wk, r, 4, "=" + bom_tot[rn], NUM)
        put(wk, r, 5, kg_aug[rn], PRICE)
        put(wk, r, 6, f"=D{r}/H{r}", PRICE, SUBH, True)
        put(wk, r, 7, f"=F{r}/E{r}-1", PCT)
        put(wk, r, 8, rc["div"], "0.0000")
        r += 1

    # المنتج التام
    wf = sheet(wb, S_FG, [30, 12, 13, 14, 10, 14, 70])
    title_rows(wf, "تقييم مخزون المنتج التام آخر سبتمبر 2026 (الوصفات بأسعار خامات سبتمبر + تكلفة التعبئة)", None, 7)
    header(wf, 3, ["اسم الصنف", "كمية الجرد 30/9", "سعر أغسطس", "سعر سبتمبر", "التغير %", "قيمة المخزون 30/9", "ملاحظة"])
    r = 4
    for f in fg_rows:
        f["row_calc"] = r
        put(wf, r, 1, f["name"], fill=FLAG if f["flag"] else None)
        put(wf, r, 2, f["qty"], NUM)
        put(wf, r, 3, f["aug_price"], PRICE)
        if f["rule"]:
            rn, w, pkg = f["rule"]
            put(wf, r, 4, f"=ROUND({w:g}*{q(S_KG)}!F{kg_row[rn]}+{pkg:g},4)", PRICE)
        else:
            put(wf, r, 4, f["price"], PRICE)
        put(wf, r, 5, f'=IF(OR(C{r}="",D{r}=""),"",D{r}/C{r}-1)', PCT)
        put(wf, r, 6, f'=IF(D{r}="","",B{r}*D{r})', NUM)
        put(wf, r, 7, f["note"], fill=FLAG if f["flag"] else None)
        r += 1
    fg_tot = r + 1
    put(wf, fg_tot, 1, "إجمالى قيمة مخزون المنتج التام آخر سبتمبر (الأصناف المُسعّرة)", fill=TOT, bold=True)
    put(wf, fg_tot, 6, f"=SUM(F4:F{r - 1})", NUM, TOT, True)
    put(wf, fg_tot + 2, 1, "إجمالى قيمة المخزون الكلى آخر سبتمبر (خامات + منتج تام)", fill=TOT, bold=True)
    put(wf, fg_tot + 2, 6, f"=F{fg_tot}+{q(S_MAIN)}!L{m_tot}", NUM, TOT, True)

    # السيناريوهات
    wsn = sheet(wb, S_SCN, [52, 16, 16])
    title_rows(wsn, "صنفا الفراخ المفرومة — تصحيح الفاتورة 1111115217 (معتمد بتأكيد الإدارة)", None, 3)
    header(wsn, 3, ["القراءة", "مفروم فراخ شعبي", "مفروم فراخ فاخر"])
    labels = [k for k in scn["مفروم فراخ شعبي"] if not k.startswith("_")]
    r = 4
    for lab in labels + ["_aug", "_orig", "_file", "_close_q"]:
        put(wsn, r, 1, {"_aug": "متوسط أغسطس (للمقارنة)", "_orig": "السعر الذى كان بملف جرد سبتمبر",
                        "_file": "السعر المُدخل بملف سبتمبر الآن", "_close_q": "كمية جرد 30/9"}
            .get(lab, "متوسط سبتمبر — " + lab), fill=SUBH if lab in (SCN_APPLIED, "_file") else None,
            bold=lab in (SCN_APPLIED, "_file"))
        for c, nm in ((2, "مفروم فراخ شعبي"), (3, "مفروم فراخ فاخر")):
            put(wsn, r, c, scn[nm][lab], NUM if lab == "_close_q" else PRICE,
                SUBH if lab in (SCN_APPLIED, "_file") else None)
        r += 1
    put(wsn, r + 1, 1, CORRECTION_NOTE)
    wsn.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    wsn.row_dimensions[r + 1].height = 48

    # المطابقة
    wc = sheet(wb, S_REC, [66, 18, 18, 14, 12])
    title_rows(wc, "المطابقة النهائية — كل فحص يجب أن يكون الفرق صفرًا", None, 5)
    header(wc, 3, ["الفحص", "القيمة المحسوبة", "القيمة المرجعية", "الفرق", "النتيجة"])
    checks = [
        ("إجمالى أسطر مشتريات سبتمبر كما سُجّلت = الإجمالى المطبوع بالتقرير",
         f"={q(S_PUR)}!M{pt}", f"={q(S_PUR)}!M{pt + 1}", 0.005),
        ("أثر تصحيح الفاتورة 1111115217 = Σ الكمية × (السعر المصحّح − المسجّل)", f"={q(S_PUR)}!H{pt + 3}",
         sum(p["qty"] * (p["price"] - p["price_recorded"]) for p in purchases if p.get("corrected")), 0.005),
        ("المشتريات المحمّلة على الأصناف + المجمّع + المستبعدة = إجمالى المشتريات",
         f"={q(S_MAIN)}!F{m_tot}+{q(S_MAIN)}!F{pr}+{q(S_PUR)}!H{pt + 2}", f"={q(S_PUR)}!H{pt}", 0.005),
        ("قيمة أول المدة بورقة الاحتساب = قيمة أصناف جرد 31/8 المنقولة لسبتمبر",
         f"={q(S_MAIN)}!D{m_tot}",
         "=" + "+".join(f"N({q(S_OPEN)}!F{open_row[key(SEP_ALIASES.get(x['name'], x['name']))]})"
                        for x in items.values() if key(SEP_ALIASES.get(x["name"], x["name"])) in open_row), 0.005),
        ("قيمة مخزون الخامات 30/9 = Σ (كمية جرد 30/9 × السعر المُدخل بالملف)",
         f"={q(S_MAIN)}!L{m_tot}", f"=SUMPRODUCT({q(S_MAIN)}!K5:K{m_last},{q(S_MAIN)}!J5:J{m_last})", 0.01),
    ]
    for i, (lab, a, b, tol) in enumerate(checks):
        rr = 4 + i
        put(wc, rr, 1, lab)
        put(wc, rr, 2, a, NUM)
        put(wc, rr, 3, b, NUM)
        put(wc, rr, 4, f"=B{rr}-C{rr}", NUM)
        put(wc, rr, 5, f'=IF(ABS(D{rr})<={tol},"✔ مطابق","✘ فرق")')

    wb.save(OUT_CALC)
    return {"m_tot": m_tot, "fg_tot": fg_tot, "row_of": row_of, "kg_row": kg_row, "open_tot": open_tot}


def add_notes_and_checks(rules, flags, extra_checks, idx, aug2_total):
    wb = openpyxl.load_workbook(OUT_CALC)
    wo = wb[S_OPEN]
    put(wo, idx["open_tot"] + 1, 6, aug2_total, NUM, TOT)
    wc = wb[S_REC]
    for lab, a, b, ok in extra_checks:
        rr = wc.max_row + 1
        put(wc, rr, 1, lab)
        put(wc, rr, 2, a, NUM)
        put(wc, rr, 3, b, NUM)
        put(wc, rr, 4, (a - b) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else None, NUM)
        put(wc, rr, 5, "✔ مطابق" if ok else "✘ فرق")
    ws = sheet(wb, S_NOTE, [24, 55, 75])
    title_rows(ws, "المنهجية، الافتراضات، والاستثناءات — سبتمبر 2026", None, 3)
    r = 3
    for a, b in rules:
        put(ws, r, 1, a, bold=True)
        put(ws, r, 2, b)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 45
        r += 1
    r += 1
    header(ws, r, ["نوع الملاحظة", "الصنف / البند", "التفصيل"])
    ws.freeze_panes = None
    for t in flags:
        r += 1
        for c, v in enumerate(t, 1):
            put(ws, r, c, v, fill=FLAG)
    wb.save(OUT_CALC)


def recalc(path):
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([shutil.which("soffice"), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, str(path)],
                       capture_output=True, check=True, timeout=180)
        return openpyxl.load_workbook(Path(d) / path.name, data_only=True)


# ---------------------------------------------------------------- التشغيل
def main():
    aug_wac, aug_pool_wac, jul, codes, jul_fg, kg_jul_tab = august_baseline()
    aug2_raw, aug2_fg = load_aug2()
    aug2_total = openpyxl.load_workbook(F_AUG2, data_only=True)["جرد الخامات آخر أغسطس"]
    aug2_total = next(num(r[10]) for r in aug2_total.iter_rows(values_only=True)
                      if r[0] and clean(r[0]).startswith("إجمالى قيمة مخزون الخامات"))
    opening, open_checks = build_opening(aug_wac, aug2_raw)
    price_mismatch = [c for c in open_checks if c[0] == "سعر"]
    assert not price_mismatch, price_mismatch                  # متوسط أغسطس = ملف «-2» لكل صنف مشترك

    sep_raw, sep_fg = read_count(F_CNT9)
    purchases, printed = A.load_purchases(F_PUR9)
    correct_purchases(purchases)
    assert abs(sum(p["total_recorded"] for p in purchases) - printed) < 0.005
    targets = {key(x["name"]): x for x in sep_raw}
    map_purchases(purchases, codes, targets)
    items, pool = compute(sep_raw, opening, purchases)
    scn = scenarios(items, purchases)

    # استقلال: إعادة احتساب المتوسط سطرًا سطرًا من أسطر الشراء
    agg = defaultdict(lambda: [0.0, 0.0])
    for p in purchases:
        if p["item"] != NO_MATCH:
            agg[key(p["item"])][0] += p["qty"]
            agg[key(p["item"])][1] += p["total"]
    pq, pv = agg[key(POOL)]
    oq = sum(opening[key(m)]["qty"] for m in POOL_MEMBERS)
    op = opening[key(POOL_MEMBERS[0])]["price"]
    pool_ind = round((oq * op + pv) / (oq + pq), 4)
    for k, x in items.items():
        o = opening.get(key(SEP_ALIASES.get(x["name"], x["name"])))
        q_, v_ = agg.get(k, (0.0, 0.0))
        if x["name"] in POOL_MEMBERS:
            exp = pool_ind
        elif q_ and o and o["price"] is not None:
            exp = round((o["qty"] * o["price"] + v_) / (o["qty"] + q_), 4)
        elif q_:
            exp = round(v_ / q_, 4)
        else:
            exp = o["price"] if o else None
        assert exp == x["wac"], (x["name"], exp, x["wac"])

    # الوصفات والمنتج التام
    recipes = A.load_recipes(A.load_june_prices())
    sep_price = {k: (x["open_p"] if x["unresolved"] else x["wac"]) for k, x in items.items()}
    aug_price = {k: x["open_p"] for k, x in items.items()}

    def kg(prices):
        return {rn: recipe_cost(rc, lambda it: prices.get(key(POOL_MEMBERS[0] if it == POOL else it)))[1]
                for rn, rc in recipes.items()}
    kg_sep, kg_aug = kg(sep_price), kg(aug_price)
    prov_recipes = {rn for rn, rc in recipes.items() for i in rc["sub"] + rc["main"] if i["item"] in UNRESOLVED}
    fg_rows = price_fg(sep_fg, jul_fg, aug2_fg, kg_sep, kg_aug, recipes, prov_recipes)

    idx = build_calc(items, pool, purchases, printed, opening, recipes, kg_sep, kg_aug, kg_jul_tab, fg_rows, scn,
                     prov_recipes)
    changed, added = update_september_file(items, fg_rows)

    # مطابقة ملف سبتمبر بعد إعادة الاحتساب
    wf = recalc(OUT_FILE).worksheets[0]
    raw_val = fg_val = 0.0
    n_ok = n_all = 0
    for x in items.values():
        h = wf.cell(x["row"], 8).value
        n_all += 1
        n_ok += (h == x["file_price"]) or (h in (None, "") and x["file_price"] is None)
        iv = wf.cell(x["row"], 9).value
        exp = x["qty"] * (x["file_price"] or 0)
        assert abs((iv or 0) - exp) < 0.005, (x["name"], iv, exp)
        raw_val += iv or 0
    for f in fg_rows:
        lv = wf.cell(f["row"], 12).value
        assert (lv == f["price"]) or (lv in (None, "") and f["price"] is None), (f["name"], lv, f["price"])
        mv = wf.cell(f["row"], 13).value
        fg_val += mv if isinstance(mv, (int, float)) else 0
    assert n_ok == n_all, (n_ok, n_all)

    open_val = sum(o["qty"] * o["price"] for o in opening.values() if o["price"] is not None)
    qty_effect = sum((o["qty"] - o["aug2_qty"]) * o["price"] for o in opening.values()
                     if o["price"] is not None and o["aug2_qty"] is not None)
    carried = sum(o["qty"] * o["price"] for k, o in opening.items() if o["price"] is not None and
                  k in {key(SEP_ALIASES.get(x["name"], x["name"])) for x in items.values()})
    extra = [("أسعار ملف «تسعير جرد اخر شهر ٩» = الاحتساب المستقل (عدد الأصناف)", n_ok, n_all, n_ok == n_all),
             ("قيمة مخزون الخامات بملف سبتمبر بعد إعادة الاحتساب (Σ عمود القيمة)", round(raw_val, 2),
              round(sum(x["qty"] * (x["file_price"] or 0) for x in items.values()), 2), True),
             ("قيمة جرد 31/8 المعدَّل × متوسط أغسطس = إجمالى ملف «-2» + أثر فروق الكميات بينهما "
              f"({qty_effect:+,.2f}، أغلبه «زيت سيارات» 2 بالجرد مقابل 16 بملف «-2»)",
              round(open_val, 2), round(aug2_total + qty_effect, 2), abs(open_val - aug2_total - qty_effect) < 1.0),
             ("أصناف جرد 31/8 ذات قيمة غير موجودة بجرد سبتمبر (قيمة)", round(open_val - carried, 2), 0.0,
              abs(open_val - carried) < 0.01)]
    rules, flags = notes(items, purchases, fg_rows, open_checks, opening, scn)
    add_notes_and_checks(rules, flags, extra, idx, aug2_total)

    wc = recalc(OUT_CALC)
    errs = [(ws.title, c.coordinate) for ws in wc for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
    assert not errs, errs[:5]
    wm = wc[S_MAIN]
    for k, x in items.items():
        v = wm.cell(idx["row_of"][k], 9).value
        assert (round(v, 4) if isinstance(v, (int, float)) else None) == x["wac"], (x["name"], v, x["wac"])
        j = wm.cell(idx["row_of"][k], 10).value
        assert (j if isinstance(j, (int, float)) else None) == x["file_price"], (x["name"], j, x["file_price"])
    wfg = wc[S_FG]
    for f in fg_rows:
        v = wfg.cell(f["row_calc"], 4).value
        assert (round(v, 4) if isinstance(v, (int, float)) else None) == f["price"], (f["name"], v, f["price"])
    rec = [(wc[S_REC].cell(r, 1).value, wc[S_REC].cell(r, 5).value) for r in range(4, wc[S_REC].max_row + 1)]
    bad = [r for r in rec if r[1] != "✔ مطابق"]
    raw_calc = wm.cell(idx["m_tot"], 12).value
    fg_calc = wc[S_FG].cell(idx["fg_tot"], 6).value
    assert abs(raw_calc - raw_val) < 0.01 and abs(fg_calc - fg_val) < 0.01, (raw_calc, raw_val, fg_calc, fg_val)

    write_report(items, pool, purchases, printed, opening, open_val, aug2_total, fg_rows, kg_sep, kg_aug, scn,
                 raw_val, fg_val, changed, added, rec, flags)
    print(f"September purchases: {len(purchases)} lines, {printed:,.2f}; opening value {open_val:,.2f} "
          f"(file -2: {aug2_total:,.2f})")
    print(f"raw 30/9 {raw_val:,.2f} | finished goods {fg_val:,.2f} | prices written {len(changed)} | "
          f"verified {n_ok}/{n_all} | reconciliation not matching: {bad}")


def notes(items, purchases, fg_rows, open_checks, opening, scn):
    mapped = sum(p["total"] for p in purchases if p["item"] != NO_MATCH)
    fix = sum(p["total"] - p["total_recorded"] for p in purchases)
    rules = [
        ("نطاق العمل", "متوسط التكلفة المرجّح (WAC) للخامات فقط — لم تُستخدم FIFO / LIFO / التكلفة المعيارية."),
        ("المعادلة", "متوسط سبتمبر = (كمية جرد 31/8 × متوسط أغسطس + قيمة مشتريات سبتمبر) ÷ (كمية جرد 31/8 + كمية "
                     "مشتريات سبتمبر). يُحتسب بالدقة الكاملة ويُقرَّب السعر المُدخل لأربعة أرقام عشرية."),
        ("رصيد أول المدة", "الجرد الفعلى المعدَّل «جر د شهر 8.xlsx» (مجموع المواقع الخمسة) × متوسط أغسطس المُعاد احتسابه "
                           "من مصادره؛ طوبق السعر مع ملف «تسعير جرد اخر شهر ٨ -2» لكل صنف مشترك (لا فرق)."),
        ("المشتريات", f"تقرير المشتريات 01/09–30/09/2026: {len(purchases)} سطرًا، الإجمالى مطابق للمطبوع. قيمة الشراء = "
                      f"«الإجمالى» بالفاتورة (7 أسطر بها فرق تقريب بين الكمية × السعر والإجمالى)، بعد تصحيح سطرَى "
                      f"الفاتورة 1111115217 (أثر {fix:+,.2f} ج.م). المحمَّل على الأصناف {mapped:,.2f} ج.م."),
        ("ربط الأكواد", "الكود المستخدم فى يوليو/أغسطس يُربط بنفس الصنف؛ الأكواد الجديدة رُبطت بالاسم ونطاق السعر بدرجة ثقة "
                        "مُسجّلة. مجمّع اللحوم الهندى متوسط واحد للبنود الأربعة."),
        ("تصحيح فاتورة", f"{CORRECTION_NOTE}. صنفا الفراخ المفرومة (شعبي/فاخر) سُعّرا بالمتوسط بعد التصحيح، "
                         f"والوصفات تستخدم هذا السعر؛ القراءات البديلة للإفصاح بورقة «{S_SCN}»."),
        ("أصناف بلا مشتريات", "يبقى سعر أغسطس كما هو. رصيد بلا سعر يُستبعد من البسط والمقام."),
        ("ملف سبتمبر", "«تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx» نسخة من «جرد شهر 9.xlsx» تغيّرت فيها خانات السعر فقط "
                       "(عمود H للخامات و L للمنتج التام)، وأُضيفت =K×L بعمود M الفارغ لصفوف المنتج التام المُسعّرة؛ "
                       "تحقق آلى من عدم تغيّر أى خانة أخرى. الأصناف بلا سعر تُركت كما هى."),
        ("المنتج التام", "الوزن × سعر كجم الوصفة بأسعار سبتمبر + تكلفة التعبئة؛ الأصناف خارج الوصفات بآخر سعر متاح."),
    ]
    flags, seen = [], set()
    for p in purchases:
        if p.get("corrected"):
            flags.append(("تصحيح سعر فاتورة", f"سطر {p['idx']} كود {p['code']} «{p['name']}» ← {p['item']}",
                          f"المسجّل {p['price_recorded']:g} × {p['qty']:,.0f} = {p['total_recorded']:,.2f}؛ المعتمد "
                          f"{p['price']:g} × {p['qty']:,.0f} = {p['total']:,.2f} (تأكيد الإدارة: السعران معكوسان)"))
    for nm, s in scn.items():
        flags.append(("أثر التصحيح", nm, f"السعر بالملف {s['_orig']} ← {s['_file']:.4f}؛ القراءات: " + "، ".join(
            f"{k}: {v:.4f}" for k, v in s.items() if not k.startswith("_"))))
    for p in purchases:
        if p["code"] in seen and not p.get("outlier"):
            continue
        seen.add(p["code"])
        if p["item"] == NO_MATCH:
            flags.append(("شراء بلا مقابل", f"كود {p['code']} «{p['name']}» — {p['total']:,.2f} ج.م", p["basis"]))
        elif p.get("outlier") and p["invoice"] != SUSPECT_INVOICE:
            flags.append(("سعر شاذ (دخل كما سُجّل)", f"سطر {p['idx']} كود {p['code']} «{p['name']}»", p["basis"]))
        elif p["conf"] == "متوسطة":
            flags.append(("ربط بثقة متوسطة", f"كود {p['code']} «{p['name']}» ← {p['item']}", p["basis"]))
    for c in open_checks:
        o = opening[key(c[1])]
        eff = (c[2] - c[3]) * (o["price"] or 0)
        if abs(eff) < 0.01:                       # فرق كمية لصنف بلا سعر — لا أثر على القيمة
            continue
        flags.append(("فرق كمية مع ملف «-2»", c[1], f"الجرد المعدَّل «جر د شهر 8» = {c[2]:g} مقابل {c[3]:g} بملف «-2» — "
                      f"اعتُمد الجرد الفعلى؛ الأثر على قيمة أول المدة {eff:+,.2f} ج.م"))
    for x in items.values():
        if x["wac"] is None and x["qty"] > 0:
            flags.append(("صنف بلا سعر", f"{x['name']} — كمية جرد 30/9 = {x['qty']:,.2f}", x["status"]))
        if x["issue"]:
            flags.append(("ملاحظة على ملف الجرد", x["name"], x["issue"]))
    for k, o in opening.items():
        if o["issue"]:
            flags.append(("ملاحظة على جرد 31/8", o["name"], o["issue"]))
    for f in fg_rows:
        if f["flag"]:
            flags.append(("منتج تام بلا سعر", f"{f['name']} — كمية {f['qty']:,.2f}", f["note"]))
    return rules, flags


def write_report(items, pool, purchases, printed, opening, open_val, aug2_total, fg_rows, kg_sep, kg_aug, scn,
                 raw_val, fg_val, changed, added, rec, flags):
    L = []
    w = L.append
    mapped = sum(p["total"] for p in purchases if p["item"] != NO_MATCH)
    w("# ABOHASHEM FOR MEAT — متوسط التكلفة المرجّح للخامات آخر سبتمبر 2026 وتقييم جرد 30/9\n")
    w("**الملفات:** `تسعير جرد اخر شهر ٩ سنة ٢٠٢٦.xlsx` (ملف «جرد شهر 9» بعد إدخال الأسعار فقط) — "
      "`حساب متوسط التكلفة سبتمبر ٢٠٢٦.xlsx` (الاحتساب بمعادلات حيّة + المطابقة) — `model/wac_september_2026.py`.\n")
    w("## ١. المعادلة\n")
    w("```\nمتوسط سبتمبر = (كمية جرد 31/8 × متوسط أغسطس + قيمة مشتريات سبتمبر)\n"
      "              ÷ (كمية جرد 31/8 + كمية مشتريات سبتمبر)\n```\n")
    w("## ٢. الإجماليات\n")
    w("| البند | القيمة (ج.م) |\n|---|---:|")
    w(f"| رصيد أول المدة (جرد 31/8 المعدَّل × متوسط أغسطس) | {open_val:,.2f} |")
    w(f"| إجمالى ملف «تسعير جرد اخر شهر ٨ -2» (الفرق = فروق كميات بينه وبين الجرد المعدَّل، أغلبها «زيت سيارات» 2 مقابل 16) "
      f"| {aug2_total:,.2f} |")
    fix = sum(p["total"] - p["total_recorded"] for p in purchases)
    w(f"| مشتريات سبتمبر (إجمالى التقرير، {len(purchases)} سطرًا) | {printed:,.2f} |")
    w(f"| أثر تصحيح الفاتورة 1111115217 | {fix:,.2f} |")
    w(f"| مشتريات سبتمبر بعد التصحيح | {printed + fix:,.2f} |")
    w(f"| منها محمّلة على أصناف الجرد | {mapped:,.2f} |")
    w(f"| منها بلا مقابل بالجرد (مستبعدة) | {printed + fix - mapped:,.2f} |")
    w(f"| **مخزون الخامات 30/9 (بالأسعار المُدخلة بالملف)** | **{raw_val:,.2f}** |")
    w(f"| **مخزون المنتج التام 30/9** | **{fg_val:,.2f}** |")
    w(f"| **إجمالى المخزون 30/9** | **{raw_val + fg_val:,.2f}** |\n")
    w(f"أُدخل {len(changed)} سعرًا بملف سبتمبر (عمود H للخامات و L للمنتج التام). عمود M (قيمة المنتج التام) كان فارغًا "
      f"بالملف الأصلى فأُضيفت له =K×L فى {len(added)} صفًا مُسعّرًا فقط. لم تتغير أى خانة أخرى — لا قيمة ولا معادلة ولا "
      "تنسيق (تحقق آلى خانة بخانة).\n")
    sh, fa = scn["مفروم فراخ شعبي"], scn["مفروم فراخ فاخر"]
    w("## ٣. مفروم الفراخ — تصحيح الفاتورة 1111115217\n")
    w(CORRECTION_NOTE + ". بناءً عليه سُعّر الصنفان بالمتوسط بعد التصحيح، وأُعيد تسعير الوصفات التى تستخدمهما:\n")
    w("| السطر | الكمية | السعر المسجّل | الإجمالى المسجّل | السعر المعتمد | الإجمالى المعتمد |\n|---|---:|---:|---:|---:|---:|")
    for p in purchases:
        if p.get("corrected"):
            w(f"| {p['idx']} — كود {p['code']} «{p['name']}» | {p['qty']:,.0f} | {p['price_recorded']:g} | "
              f"{p['total_recorded']:,.2f} | {p['price']:g} | {p['total']:,.2f} |")
    w("\n| القراءة | مفروم فراخ شعبي | مفروم فراخ فاخر |\n|---|---:|---:|")
    for lab in [k for k in sh if not k.startswith("_")]:
        b = "**" if lab == SCN_APPLIED else ""
        w(f"| {b}متوسط سبتمبر — {lab}{b} | {b}{sh[lab]:.4f}{b} | {b}{fa[lab]:.4f}{b} |")
    w(f"| متوسط أغسطس | {sh['_aug']:.4f} | {fa['_aug']:.4f} |")
    w(f"| السعر الذى كان بملف جرد سبتمبر | {sh['_orig']} | {fa['_orig']} |")
    w(f"| كمية جرد 30/9 | {sh['_close_q']:,.2f} | {fa['_close_q']:,.2f} |")
    w(f"| قيمة 30/9 بالسعر المعتمد | {sh['_close_q'] * sh['_file']:,.2f} | {fa['_close_q'] * fa['_file']:,.2f} |\n")
    w("الرتبة B (كود 371، 1,800 كجم بسعر 34) حُمّلت على «شعبي» لعدم وجود بند B بالجرد — داخلة فى القراءات الثلاث.\n")
    w("## ٤. جدول المراجعة — الأصناف التى عليها مشتريات فى سبتمبر\n")
    w("| كود الشراء | الصنف | كمية أول المدة | قيمة أول المدة | كمية المشتريات | قيمة المشتريات | إجمالى الكمية | "
      "إجمالى القيمة | المتوسط 30/9 | السعر بالملف | كمية جرد 30/9 | قيمة 30/9 |")
    w("|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|")
    codes = defaultdict(set)
    for p in purchases:
        codes[key(p["item"])].add(p["code"])
    rows = [(POOL, pool, sum(items[key(m)]["qty"] for m in POOL_MEMBERS), pool["wac"])] + \
           [(x["name"], x, x["qty"], x["file_price"]) for x in items.values() if x["pq"] > 0 and x["name"] not in POOL_MEMBERS]
    for name, x, cq, fp in rows:
        oq = x["open_q"] if x["open_p"] is not None else 0.0
        ov = oq * x["open_p"] if x["open_p"] is not None else 0.0
        w(f"| {'، '.join(sorted(codes.get(key(name), [])))} | {name} | {x['open_q']:,.2f} | {ov:,.2f} | {x['pq']:,.2f} | "
          f"{x['pv']:,.2f} | {oq + x['pq']:,.2f} | {ov + x['pv']:,.2f} | {x['wac']:,.4f} | "
          f"{fp if fp is not None else '—'} | {cq:,.2f} | {cq * (fp or 0):,.2f} |")
    w("\nالأصناف بلا مشتريات فى سبتمبر احتفظت بسعر أغسطس — التفصيل الكامل بورقة «احتساب متوسط سبتمبر».\n")
    big = sorted(((n, x) for n, x, _, _ in rows if x["open_p"] and abs(x["wac"] / x["open_p"] - 1) >= 0.10),
                 key=lambda t: t[1]["wac"] / t[1]["open_p"])
    if big:
        w("**أصناف تغيّر متوسطها 10% أو أكثر عن أغسطس (للمراجعة):**\n")
        w("| الصنف | متوسط أغسطس | متوسط سبتمبر | التغير | أسعار الشراء بسبتمبر |\n|---|---:|---:|---:|---|")
        for n, x in big:
            pr = sorted({p["price"] for p in purchases if p["item"] == n})
            w(f"| {n} | {x['open_p']:.4f} | {x['wac']:.4f} | {x['wac'] / x['open_p'] - 1:+.1%} | "
              f"{' / '.join(f'{v:g}' for v in pr)} |")
        w("")
    w("## ٥. أمثلة محسوبة يدويًا\n")
    for nm in ("بصل", "دهن بلدي", "فلانك"):
        x = items[key(nm)]
        ov = x["open_q"] * x["open_p"]
        w(f"**{nm}:** أول المدة {x['open_q']:,.2f} × {x['open_p']:.4f} = {ov:,.4f} ؛ المشتريات {x['pq']:,.2f} بقيمة "
          f"{x['pv']:,.2f} ؛ ({ov:,.4f} + {x['pv']:,.2f}) ÷ ({x['open_q']:,.2f} + {x['pq']:,.2f}) = "
          f"{ov + x['pv']:,.4f} ÷ {x['open_q'] + x['pq']:,.2f} = **{x['wac']:.4f}** (أغسطس {x['open_p']:.4f}، "
          f"{x['wac'] / x['open_p'] - 1:+.2%}). قيمة 30/9 = {x['qty']:,.2f} × {x['wac']:.4f} = {x['qty'] * x['wac']:,.2f}.\n")
    w("## ٦. سعر كجم الوصفات\n")
    w("| المنتج | أغسطس | سبتمبر | التغير |\n|---|---:|---:|---:|")
    for rn in kg_sep:
        w(f"| {rn} | {kg_aug[rn]:.4f} | {kg_sep[rn]:.4f} | {kg_sep[rn] / kg_aug[rn] - 1:+.2%} |")
    w("\nالوصفات التى تحتوى مفروم الفراخ محسوبة بمتوسط سبتمبر للمفروم بعد تصحيح الفاتورة.\n")
    w("## ٧. المطابقة النهائية\n")
    w("| الفحص | النتيجة |\n|---|---|")
    for lab, res in rec:
        w(f"| {lab} | {res} |")
    w("| كل متوسط = إعادة احتساب مستقلة سطرًا بسطر من أسطر الشراء (4 أرقام عشرية) | ✔ مطابق |")
    w("| إعادة احتساب الملفين بـ LibreOffice — لا أخطاء، والقيمة = الكمية × السعر لكل صنف | ✔ |\n")
    w("## ٨. البنود المُعلَّمة للمراجعة\n")
    w("| النوع | الصنف / البند | التفصيل |\n|---|---|---|")
    for t in flags:
        w("| " + " | ".join(str(v).replace("|", "/") for v in t) + " |")
    OUT_MD.write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
