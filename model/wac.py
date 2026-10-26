#!/usr/bin/env python3
"""متوسط التكلفة المرجّح (WAC) لأى شهر — نموذج واحد بدل سكربت لكل شهر.

كل شهر مجلد ثابت الأسماء تحت monthly/YYYY-MM/ :
    month.json                  استثناءات الشهر (أكواد جديدة، أسماء بديلة، تصحيحات مؤكَّدة) — مطلوب
    purchases.pdf               تقرير المشتريات من البرنامج — مطلوب
    stock_count.xlsx            الجرد الفعلى آخر الشهر — مطلوب
    stock_count_revised.xlsx    الجرد المعدَّل إن صدر (يصبح رصيد أول المدة للشهر التالى) — اختيارى
    valuation_issued.xlsx       تقييم صادر للشهر من خارج النموذج (للمطابقة ولأسعار ما لا يُحتسب) — اختيارى
ومرجع ثابت لكل الأشهر تحت monthly/reference/ : recipes.xlsx (الوصفات) و june_prices.xlsx (ربط الخامات).

السلسلة: كل شهر يبدأ من رصيد آخر الشهر السابق بمتوسطه؛ أول السلسلة شهر «bootstrap» (يوليو 2026)
تُقرأ أرصدته وأسعاره وربط أكواده من ملف التقييم الصادر له.

    متوسط الشهر = (كمية أول المدة × سعر أول المدة + قيمة مشتريات الشهر) ÷ (كمية أول المدة + كمية المشتريات)
    لا مشتريات ⇒ السعر كما هو؛ أول مدة بلا سعر ⇒ المتوسط من المشتريات وحدها؛
    مجمّع اللحوم الهندى (كتف/رقبه/سن/أرباع) متوسط واحد.

القواعد هى نفسها بـ wac_august_2026.py و wac_september_2026.py، وtests/test_wac_engine.py يتحقق أن هذا
النموذج يعيد أسعار أغسطس وسبتمبر الحالية نفسها. السكربتان باقيان كما هما (مصنّفات الاحتساب التفصيلية).

التشغيل:
    python3 model/wac.py 2026-10            يحتسب الشهر ويكتب wac.json و REPORT.md (و stock_count_priced.xlsx
                                            إن كان بالجرد عمود «السعر»)
    python3 model/wac.py --check 2026-10    قائمة استلام ملفات الشهر فقط: يسمّى كل ملف ناقص
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import OrderedDict, defaultdict
from functools import lru_cache
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import wac_august_2026 as A  # noqa: E402
from wac_august_2026 import NO_MATCH, POOL, POOL_MEMBERS, clean, key, num, recipe_cost  # noqa: E402

ROOT = A.ROOT
MONTHLY = ROOT / "monthly"
REFERENCE = MONTHLY / "reference"

# ---------------------------------------------------------------- قائمة الاستلام
REQUIRED = OrderedDict([
    ("month.json", "استثناءات الشهر (أكواد شراء جديدة، أسماء بديلة، تصحيحات مؤكَّدة)"),
    ("purchases.pdf", "تقرير المشتريات التفصيلى للشهر من البرنامج (PDF)"),
    ("stock_count.xlsx", "الجرد الفعلى آخر الشهر (خامات + منتج تام)"),
])
OPTIONAL = OrderedDict([
    ("stock_count_revised.xlsx", "الجرد المعدَّل آخر الشهر إن صدر — يصبح رصيد أول المدة للشهر التالى"),
    ("valuation_issued.xlsx", "تقييم صادر للشهر من خارج النموذج — للمطابقة ولأسعار المنتج خارج الوصفات"),
])
BOOTSTRAP_REQUIRED = OrderedDict([
    ("month.json", "إعداد شهر البداية"),
    ("valuation_issued.xlsx", "ملف تقييم شهر البداية (أرصدة وأسعار وربط أكواد وقواعد المنتج التام)"),
])
REFERENCE_REQUIRED = OrderedDict([
    ("recipes.xlsx", "وصفات التصنيع (BOM)"),
    ("june_prices.xlsx", "أسعار يونيه لربط خامات الوصفات بأصناف الجرد"),
])


# ---------------------------------------------------------------- وصفات الخامة الواحدة
# تأكيد الإدارة 26/10/2026: المعيار = خلطة التصنيع بمتوسط آخر الشهر. المنتج التام: (الخامة، كجم، تعبئة).
# مفروم صافي 350 = 0.32 كجم لحم هندى + 6.25؛ كبده اسكندراني (400 بالمبيعات) = 0.34 كجم كبده + 6.25.
# تسرى من MATERIAL_FROM: تقييمات أغسطس وسبتمبر صدرت قبل الوصفتين فتبقى كما صدرت.
MATERIAL_RULES = {
    "مفروم ابو هاشم 350": (POOL, 0.32, 6.25),
    "كبده اسكندراني 370": ("كبده", 0.34, 6.25),
}
MATERIAL_FROM = "2026-10"


def material_rule(month: str, name: str):
    if month < MATERIAL_FROM:
        return None
    return next((r for n, r in MATERIAL_RULES.items() if key(n) == key(name)), None)


def material_price(rule, prices: dict, pool_wac):
    item, kg, pkg = rule
    p = pool_wac if item == POOL else prices.get(key(item))
    return None if p is None else round(kg * p + pkg, 4)


class MissingFiles(Exception):
    def __init__(self, month, missing):
        self.month, self.missing = month, missing
        lines = [f"ملفات ناقصة لشهر {month}:"] + [f"  - {p}  ({d})" for p, d in missing]
        super().__init__("\n".join(lines))


def month_dir(month: str) -> Path:
    return MONTHLY / month


def prev_month(month: str) -> str:
    y, m = map(int, month.split("-"))
    return f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"


def check_month(month: str):
    """قائمة الاستلام: [(المسار النسبى، الوصف)] لكل ملف مطلوب غير موجود (فارغة = الشهر مكتمل)."""
    d = month_dir(month)
    missing = [(f"monthly/reference/{n}", t) for n, t in REFERENCE_REQUIRED.items() if not (REFERENCE / n).exists()]
    if not (d / "month.json").exists():
        return missing + [(f"monthly/{month}/{n}", t) for n, t in REQUIRED.items()]
    req = BOOTSTRAP_REQUIRED if config(month).get("bootstrap") else REQUIRED
    return missing + [(f"monthly/{month}/{n}", t) for n, t in req.items() if not (d / n).exists()]


def config(month: str) -> dict:
    return json.loads((month_dir(month) / "month.json").read_text(encoding="utf-8"))


def require(month: str):
    missing = check_month(month)
    if missing:
        raise MissingFiles(month, missing)


# ---------------------------------------------------------------- القراءة
def read_count(path: Path):
    """الجرد: صف 4 عناوين؛ عمود «اسم الصنف» الثانى يبدأ قسم المنتج التام. عمود «السعر» إن وُجد."""
    ws = openpyxl.load_workbook(path).worksheets[0]
    hdr = {c: clean(ws.cell(4, c).value) for c in range(1, ws.max_column + 1)}
    names = [c for c, h in hdr.items() if h == "اسم الصنف"]
    assert names and names[0] == 1, f"{path.name}: عمود «اسم الصنف» الأول ليس A"
    fg_col = names[1] if len(names) > 1 else None
    price = [c for c, h in hdr.items() if h.startswith("السعر")]
    raw_price = next((c for c in price if fg_col is None or c < fg_col), None)
    fg_price = next((c for c in price if fg_col and c > fg_col), None)
    raw, fg = [], []
    for r in range(5, ws.max_row + 1):
        n = clean(ws.cell(r, 1).value)
        if n and not n.startswith(("إجمالى", "اجمالي", "إجمالي")):
            locs = [num(ws.cell(r, c).value) or 0.0 for c in range(2, 7)]
            g, qty, issue = ws.cell(r, 7).value, sum(locs), None
            if isinstance(g, str) and g.replace(" ", "") != f"=SUM(B{r}:F{r})" and abs(sum(locs[:2]) - qty) > 1e-9:
                issue = f"معادلة الإجمالى «{g}» تُسقط مواقع بها رصيد — اعتُمد مجموع المواقع {qty:g}"
            elif not isinstance(g, str) and abs((num(g) or 0.0) - qty) > 1e-9:
                issue = f"الإجمالى مُدخل يدويًا ({num(g) or 0:g}) بينما مجموع المواقع {qty:g} — اعتُمد مجموع المواقع"
            raw.append({"row": r, "name": n, "qty": qty, "issue": issue,
                        "orig_price": ws.cell(r, raw_price).value if raw_price else None})
        if fg_col:
            f = ws.cell(r, fg_col).value
            if f not in (None, 0) and clean(f) not in ("", "0", "اسم الصنف"):
                fg.append({"row": r, "name": clean(f), "qty": num(ws.cell(r, fg_col + 1).value) or 0.0,
                           "orig_price": ws.cell(r, fg_price).value if fg_price else None})
    cols = {"raw_price": raw_price, "fg_col": fg_col, "fg_price": fg_price}
    return raw, fg, cols


def read_issued(path: Path):
    """تقييم صادر بنفس تصميم ملفات «تسعير جرد اخر شهر …»: ورقة الخامات وورقة المنتج التام."""
    wb = openpyxl.load_workbook(path, data_only=True)
    s_raw = next(n for n in wb.sheetnames if n.startswith("جرد الخامات"))
    s_fg = next(n for n in wb.sheetnames if n.startswith("جرد المنتج التام"))
    raw, fg, total = {}, {}, None
    for r in wb[s_raw].iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        n = clean(r[0])
        if n.startswith(("إجمالى", "اجمالي")):
            if n.startswith("إجمالى قيمة مخزون الخامات"):
                total = num(r[10])
            continue
        raw[key(n)] = {"name": n, "wac": num(r[8]), "qty": num(r[9]) or 0.0}
    for r in wb[s_fg].iter_rows(min_row=4, values_only=True):
        if r[0] and not clean(r[0]).startswith(("إجمالى", "اجمالي")):
            fg[key(r[0])] = {"name": clean(r[0]), "price": num(r[3]), "note": clean(r[6])}
    return raw, fg, total


@lru_cache(maxsize=None)
def recipes():
    return A.load_recipes(A.load_june_prices(REFERENCE / "june_prices.xlsx"), REFERENCE / "recipes.xlsx")


def kg_of(prices):
    return {rn: recipe_cost(rc, lambda it: prices.get(key(POOL_MEMBERS[0] if it == POOL else it)))[1]
            for rn, rc in recipes().items()}


def wac(open_q, open_p, pq, pv):
    if pq > 0:
        return round(pv / pq, 4) if open_p is None else round((open_q * open_p + pv) / (open_q + pq), 4)
    return open_p


# ---------------------------------------------------------------- شهر البداية
def bootstrap(month: str) -> dict:
    jul, codes, jul_fg, _kg = A.load_july(month_dir(month) / "valuation_issued.xlsx")
    cfg = config(month)
    total = read_issued(month_dir(month) / "valuation_issued.xlsx")[2]     # إجمالى الخامات المطبوع بالملف
    if cfg.get("raw_total") is not None:
        assert abs(total - cfg["raw_total"]) < 0.01, f"{month}: إجمالى الخامات {total} ≠ {cfg['raw_total']}"
    closing = OrderedDict((k, {"name": j["name"], "qty": j["qty"], "price": j["wac"], "src": "ملف شهر البداية"})
                          for k, j in jul.items())
    return {"month": month, "bootstrap": True, "codes": codes, "unmatched": {},
            "rules": {k: f["rule"] for k, f in jul_fg.items() if f["rule"]},
            "fg_aliases": {}, "items": OrderedDict((k, {"name": j["name"], "wac": j["wac"]}) for k, j in jul.items()),
            "closing": closing, "issued_fg": {k: f["price"] for k, f in jul_fg.items()},
            "fg_last": {k: f["price"] for k, f in jul_fg.items() if f["price"] is not None}, "raw_total": total}


# ---------------------------------------------------------------- الشهر
@lru_cache(maxsize=None)
def run(month: str) -> dict:
    require(month)
    cfg = config(month)
    if cfg.get("bootstrap"):
        return bootstrap(month)
    prev = run(prev_month(month))
    d = month_dir(month)
    raw, fg, cols = read_count(d / "stock_count.xlsx")
    purchases, printed = A.load_purchases(d / "purchases.pdf")

    # تصحيحات الأسعار المؤكَّدة (يبقى المسجَّل بالتقرير للمطابقة)
    corr = {(c["invoice"], c["code"]): c for c in cfg.get("corrections", [])}
    for p in purchases:
        p["total_recorded"] = p["total"]
        c = corr.get((p["invoice"], p["code"]))
        if c:
            assert abs(p["qty"] * p["price"] - p["total"]) < 0.01, p
            p.update(price_recorded=p["price"], price=c["price"], total=round(p["qty"] * c["price"], 2),
                     corrected=True, note=c.get("note"))
    assert sum(p.get("corrected", False) for p in purchases) == len(corr), f"{month}: تصحيح لا يقابله سطر شراء"
    assert abs(sum(p["total_recorded"] for p in purchases) - printed) < 0.005

    # ربط الأكواد: الأقدم يغلب (يوليو ← أغسطس ← …) ثم أكواد الشهر الجديدة
    codes = {c: (v["item"], v["confidence"], v["basis"]) for c, v in cfg.get("new_codes", {}).items()}
    codes.update(prev["codes"])
    unmatched = {**prev["unmatched"], **cfg.get("unmatched", {})}
    outliers = {(o["idx"], o["code"]): o["note"] for o in cfg.get("outliers", [])}
    targets = {key(x["name"]): x for x in raw}
    unknown, missing_target = [], []
    for p in purchases:
        if p["code"] in unmatched:
            p.update(item=NO_MATCH, conf="—", basis=unmatched[p["code"]])
            continue
        if p["code"] not in codes:
            unknown.append(p)
            continue
        item, conf, basis = codes[p["code"]]
        if item != POOL:
            if key(item) not in targets:
                missing_target.append((p["code"], item))
                continue
            item = targets[key(item)]["name"]
        note = outliers.get((p["idx"], p["code"])) or p.get("note")
        p.update(item=item, conf=conf, basis=basis + (f" | ⚠ {note}" if note else ""), outlier=bool(note))
    if unknown or missing_target:
        msg = [f"{month}: أكواد شراء تحتاج ربطًا فى monthly/{month}/month.json:"]
        msg += [f"  - كود {p['code']} «{p['name']}» (سطر {p['idx']}) — أضفه إلى new_codes أو unmatched"
                for p in {p["code"]: p for p in unknown}.values()]
        msg += [f"  - كود {c} مربوط بـ «{i}» غير الموجود بجرد الشهر" for c, i in sorted(set(missing_target))]
        raise ValueError("\n".join(msg))

    # أول المدة = رصيد آخر الشهر السابق بسعره، عبر الأسماء البديلة
    aliases = cfg.get("aliases_prev", {})
    items = OrderedDict()
    for c in raw:
        o = prev["closing"].get(key(aliases.get(c["name"], c["name"])))
        items[key(c["name"])] = {**c, "open_q": o["qty"] if o else 0.0, "open_p": o["price"] if o else None,
                                 "open_name": o["name"] if o else None, "pq": 0.0, "pv": 0.0, "n": 0}
    pool = {"name": POOL, "pq": 0.0, "pv": 0.0, "n": 0, "open_q": sum(items[key(m)]["open_q"] for m in POOL_MEMBERS),
            "open_p": items[key(POOL_MEMBERS[0])]["open_p"]}
    assert len({items[key(m)]["open_p"] for m in POOL_MEMBERS}) == 1, f"{month}: أسعار مجمّع اللحوم غير موحدة"
    for p in purchases:
        if p["item"] == NO_MATCH:
            continue
        t = pool if p["item"] == POOL else items[key(p["item"])]
        t["pq"] += p["qty"]
        t["pv"] += p["total"]
        t["n"] += 1
        t["corrected"] = t.get("corrected", False) or p.get("corrected", False)
    pool["wac"] = wac(pool["open_q"], pool["open_p"], pool["pq"], pool["pv"])
    unresolved = cfg.get("unresolved", {})
    for x in items.values():
        x["wac"] = pool["wac"] if x["name"] in POOL_MEMBERS else wac(x["open_q"], x["open_p"], x["pq"], x["pv"])
        x["unresolved"] = unresolved.get(x["name"])
        x["file_price"] = x["orig_price"] if x["unresolved"] else (x["wac"] if x["wac"] is not None else x["orig_price"])
        x["rule"] = rule_text(x)

    # احتساب مستقل من أسطر الشراء مباشرة
    agg = defaultdict(lambda: [0.0, 0.0])
    for p in purchases:
        if p["item"] != NO_MATCH:
            agg[key(p["item"])][0] += p["qty"]
            agg[key(p["item"])][1] += p["total"]
    pq, pv = agg[key(POOL)]
    pool_ind = wac(pool["open_q"], pool["open_p"], pq, pv)
    for k, x in items.items():
        q_, v_ = agg.get(k, (0.0, 0.0))
        exp = pool_ind if x["name"] in POOL_MEMBERS else wac(x["open_q"], x["open_p"], q_, v_)
        assert exp == x["wac"], (month, x["name"], exp, x["wac"])

    # الوصفات والمنتج التام
    prices = {k: (x["open_p"] if x["unresolved"] else x["wac"]) for k, x in items.items()}
    kg, kg_open = kg_of(prices), kg_of({k: x["open_p"] for k, x in items.items()})
    fg_aliases = {**prev["fg_aliases"], **cfg.get("fg_aliases", {})}
    fg_rows = []
    for f in fg:
        alias = fg_aliases.get(f["name"])
        ks = [key(f["name"])] + ([key(alias)] if alias else [])
        rule = next((prev["rules"][k] for k in ks if k in prev["rules"]), None)
        row = {**f, "price": None, "rule": rule, "flag": False}
        mrule = material_rule(month, f["name"])
        mprice = material_price(mrule, prices, pool["wac"]) if mrule else None
        if mprice is not None:
            item, w, pkg = mrule
            row.update(price=mprice, note=f"من وصفة الإدارة: {w:g} كجم «{item.lstrip('@')}» × متوسط الشهر"
                       f" + تعبئة {pkg:.2f} ج.م")
        elif rule:
            rn, w, pkg = rule
            row.update(price=round(w * kg[rn] + pkg, 4), note=f"من وصفة «{rn}»: {w:g} كجم × سعر كجم الشهر"
                       + (f" + تعبئة {pkg:.2f} ج.م" if pkg else " بدون تكلفة تعبئة (غير محددة بالوصفة)"))
        else:
            issued = next((prev["issued_fg"][k] for k in ks if prev["issued_fg"].get(k) is not None), None)
            last = next((prev["fg_last"][k] for k in ks if k in prev["fg_last"]), None)
            if issued is not None:
                row.update(price=issued, note=f"خارج الوصفات — آخر سعر بتقييم {prev['month']} الصادر")
            elif last is not None:
                row.update(price=last, note="خارج الوصفات — آخر سعر متاح بالسلسلة")
            else:
                row.update(note="بلا وصفة ولا سعر سابق — يلزم توفير سعره", flag=f["qty"] > 0)
        fg_rows.append(row)
    fg_last = {**prev["fg_last"], **{key(r["name"]): r["price"] for r in fg_rows if r["price"] is not None}}

    # رصيد آخر الشهر للشهر التالى: الجرد المعدَّل إن وُجد، ومطابقة التقييم الصادر
    issued_raw, issued_fg, issued_total = ({}, {}, None)
    if (d / "valuation_issued.xlsx").exists():
        issued_raw, issued_fg, issued_total = read_issued(d / "valuation_issued.xlsx")
    rev_aliases, iss_aliases = cfg.get("revised_aliases", {}), cfg.get("issued_aliases", {})
    revised = (d / "stock_count_revised.xlsx").exists()
    closing_rows = read_count(d / "stock_count_revised.xlsx")[0] if revised else raw
    closing, checks = OrderedDict(), []
    for c in closing_rows:
        mine = items.get(key(rev_aliases.get(c["name"], c["name"])), {}).get("wac")
        iss = issued_raw.get(key(iss_aliases.get(c["name"], c["name"])))
        if mine is not None and iss and iss["wac"] is not None and abs(mine - iss["wac"]) > 0.00005:
            checks.append(("سعر", c["name"], mine, iss["wac"]))
        price = mine if mine is not None else (iss["wac"] if iss else None)
        src = "متوسط الشهر" if mine is not None else ("التقييم الصادر" if price is not None else "")
        closing[key(c["name"])] = {"name": c["name"], "qty": c["qty"], "price": price, "src": src}
    assert not checks, (month, "متوسط النموذج ≠ التقييم الصادر", checks)

    raw_exact = sum(x["qty"] * (x["file_price"] or 0) for x in items.values())
    fg_exact = sum(r["qty"] * (r["price"] or 0) for r in fg_rows)
    raw_value, fg_value, total_value = round(raw_exact, 2), round(fg_exact, 2), round(raw_exact + fg_exact, 2)
    open_value = round(sum(x["open_q"] * (x["open_p"] or 0) for x in items.values()), 2)
    return {"month": month, "bootstrap": False, "cfg": cfg, "prev": prev["month"], "codes": codes,
            "unmatched": unmatched, "rules": prev["rules"], "fg_aliases": fg_aliases, "items": items, "pool": pool,
            "purchases": purchases, "printed": printed, "kg": kg, "kg_open": kg_open, "fg": fg_rows,
            "fg_last": fg_last, "issued_fg": {k: v["price"] for k, v in issued_fg.items()}, "closing": closing,
            "revised": revised, "issued_total": issued_total, "cols": cols, "raw_value": raw_value,
            "fg_value": fg_value, "total_value": total_value, "open_value": open_value}


def rule_text(x):
    if x.get("unresolved"):
        return "⚠ لم يُغيَّر السعر — بيانات شراء متعارضة: " + x["unresolved"]
    if x["name"] in POOL_MEMBERS:
        return "متوسط مجمّع اللحوم الهندى"
    if x["wac"] is None:
        return "بلا سعر — لا سعر أول مدة ولا مشتريات" + (" (السعر الموجود بالجرد تُرك)" if x["orig_price"] else "")
    if x["pq"] > 0 and x["open_p"] is None:
        return "المتوسط من مشتريات الشهر فقط (أول المدة بلا سعر)" if x["open_q"] else "المتوسط = متوسط سعر الشراء"
    if x.get("corrected"):
        return "قياسى: (أول المدة + المشتريات) بعد تصحيح سعر مؤكَّد"
    if x["pq"] > 0:
        return "قياسى: (أول المدة + المشتريات)"
    return "لا مشتريات فى الشهر ⇒ السعر كما هو"


# ---------------------------------------------------------------- المخرجات
def write_outputs(R: dict, out_dir: Path | None = None) -> list:
    d, out = out_dir or month_dir(R["month"]), []
    data = {
        "month": R["month"], "opening_from": R["prev"],
        "purchases": {"lines": len(R["purchases"]), "printed_total": R["printed"],
                      "corrected_total": round(sum(p["total"] for p in R["purchases"]), 2),
                      "unmatched_total": round(sum(p["total"] for p in R["purchases"] if p["item"] == NO_MATCH), 2)},
        "values": {"opening": R["open_value"], "raw": R["raw_value"], "fg": R["fg_value"],
                   "total": R["total_value"]},
        "pool": {"open_q": round(R["pool"]["open_q"], 4), "open_p": R["pool"]["open_p"], "pq": round(R["pool"]["pq"], 4),
                 "pv": round(R["pool"]["pv"], 2), "wac": R["pool"]["wac"]},
        "raw": [{"name": x["name"], "row": x["row"], "open_q": round(x["open_q"], 4), "open_p": x["open_p"],
                 "purchase_q": round(x["pq"], 4), "purchase_v": round(x["pv"], 2), "lines": x["n"], "wac": x["wac"],
                 "close_q": round(x["qty"], 4), "file_price": x["file_price"], "rule": x["rule"], "count_issue": x["issue"]}
                for x in R["items"].values()],
        # سعر الكجم بستة أرقام: sum() فى بايثون ≥ 3.12 يختلف عن الأقدم فى آخر رقم عشرى (~1e-14)
        "kg": {rn: {"open": round(R["kg_open"][rn], 6), "month": round(R["kg"][rn], 6)} for rn in R["kg"]},
        "fg": [{"name": f["name"], "row": f["row"], "qty": round(f["qty"], 4), "price": f["price"], "note": f["note"],
                "flag": f["flag"]} for f in R["fg"]],
    }
    p = d / "wac.json"
    p.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    out.append(p)
    if R["cols"]["raw_price"]:
        out.append(write_priced_count(R, d / "stock_count_priced.xlsx"))
    out.append(write_report(R, d / "REPORT.md"))
    return out


def write_priced_count(R: dict, dst: Path) -> Path:
    """نسخة من الجرد بأسعار الشهر: تُكتب خانات السعر فقط فى XML الورقة، وكل خانة أخرى كما هى."""
    src = month_dir(R["month"]) / "stock_count.xlsx"
    from openpyxl.utils import get_column_letter as L
    rp, fq, fp = R["cols"]["raw_price"], R["cols"]["fg_col"] + 1, R["cols"]["fg_price"]
    sets, adds = {}, {}
    for x in R["items"].values():
        if not x["unresolved"] and x["wac"] is not None and x["orig_price"] != x["wac"]:
            sets[f"{L(rp)}{x['row']}"] = x["wac"]
    for f in R["fg"]:
        if fp and f["price"] is not None and f["orig_price"] != f["price"]:
            sets[f"{L(fp)}{f['row']}"] = f["price"]
            adds[f"{L(fp + 1)}{f['row']}"] = f"{L(fq)}{f['row']}*{L(fp)}{f['row']}"
    zin = zipfile.ZipFile(src)
    parts = {n: zin.read(n) for n in zin.namelist()}
    sh = parts["xl/worksheets/sheet1.xml"].decode("utf-8")

    def cell_re(ref):
        return re.compile(r'<c r="%s"((?: [a-z]+="[^"]*")*)(?:/>|>(.*?)</c>)' % ref)

    def style(attrs):
        m = re.search(r' s="(\d+)"', attrs)
        return f' s="{m.group(1)}"' if m else ""

    for ref, val in sets.items():
        m = cell_re(ref).search(sh)
        assert m and "<f>" not in (m.group(2) or ""), ref
        sh = sh[:m.start()] + f'<c r="{ref}"{style(m.group(1))}><v>{val!r}</v></c>' + sh[m.end():]
    for ref, formula in list(adds.items()):
        m = cell_re(ref).search(sh)
        if not m or (m.group(2) or "").strip():      # خانة القيمة بها شىء أو غير موجودة ⇒ لا تُلمس
            adds.pop(ref)
            continue
        sh = sh[:m.start()] + f'<c r="{ref}"{style(m.group(1))}><f>{formula}</f></c>' + sh[m.end():]
    sh = re.sub(r"(<f>[^<]*</f>)<v>[^<]*</v>", r"\1", sh)
    parts["xl/worksheets/sheet1.xml"] = sh.encode("utf-8")
    wbx = parts["xl/workbook.xml"].decode("utf-8")
    parts["xl/workbook.xml"] = re.sub(r"<calcPr([^>]*?)/>", lambda m: "<calcPr" + m.group(1).replace(
        ' fullCalcOnLoad="1"', "") + ' fullCalcOnLoad="1"/>', wbx).encode("utf-8")
    parts.pop("xl/calcChain.xml", None)
    parts["xl/_rels/workbook.xml.rels"] = re.sub(rb'<Relationship [^>]*calcChain[^>]*/>', b"",
                                                 parts["xl/_rels/workbook.xml.rels"])
    parts["[Content_Types].xml"] = re.sub(rb'<Override [^>]*calcChain[^>]*/>', b"", parts["[Content_Types].xml"])
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            if n in parts:
                zi = zin.getinfo(n)
                zi.date_time = (2026, 1, 1, 0, 0, 0)          # ملف قابل للتكرار بايتًا ببايت
                zout.writestr(zi, parts[n])
    a, b = openpyxl.load_workbook(src).worksheets[0], openpyxl.load_workbook(dst).worksheets[0]
    assert a.dimensions == b.dimensions
    for row in a.iter_rows():
        for c in row:
            v = b[c.coordinate].value
            if c.coordinate in sets:
                assert v == sets[c.coordinate]
            elif c.coordinate in adds:
                assert v == "=" + adds[c.coordinate]
            else:
                assert v == c.value, (c.coordinate, c.value, v)
            assert b[c.coordinate].style_id == c.style_id, c.coordinate
    return dst


def write_report(R: dict, p: Path) -> Path:
    m, items = R["month"], list(R["items"].values())
    unpriced = [x for x in items if x["wac"] is None and x["qty"] > 0]
    L = [f"# متوسط التكلفة المرجّح — {m}", "",
         f"مولَّد بـ `python3 model/wac.py {m}` من ملفات `monthly/{m}/`؛ رصيد أول المدة = آخر {R['prev']}.", "",
         "| البند | القيمة |", "|---|---|",
         f"| أسطر المشتريات | {len(R['purchases'])} |",
         f"| إجمالى المشتريات المطبوع | {R['printed']:,.2f} |",
         f"| بعد التصحيحات المؤكَّدة | {sum(p['total'] for p in R['purchases']):,.2f} |",
         f"| أسطر بلا صنف مقابل | {sum(p['item'] == NO_MATCH for p in R['purchases'])} |",
         f"| قيمة أول المدة | {R['open_value']:,.2f} |",
         f"| قيمة الخامات آخر الشهر | {R['raw_value']:,.2f} |",
         f"| قيمة المنتج التام آخر الشهر | {R['fg_value']:,.2f} |",
         f"| **إجمالى المخزون آخر الشهر** | **{R['total_value']:,.2f}** |",
         f"| أصناف خامات / مُسعّرة | {len(items)} / {sum(x['wac'] is not None for x in items)} |",
         f"| متوسط مجمّع اللحوم الهندى | {R['pool']['wac']} |", "",
         "## سعر الكجم من الوصفات", "", "| الوصفة | أول المدة | الشهر | التغير |", "|---|---|---|---|"]
    for rn in R["kg"]:
        a, b = R["kg_open"][rn], R["kg"][rn]
        L.append(f"| {rn} | {a:.4f} | {b:.4f} | {((b / a - 1) * 100 if a else 0):+.1f}% |")
    L += ["", "## يحتاج متابعة", ""]
    flags = ([f"- خامة بلا سعر وبها رصيد: {x['name']} ({x['qty']:g})" for x in unpriced]
             + [f"- منتج تام بلا سعر وبه رصيد: {f['name']} ({f['qty']:g})" for f in R["fg"] if f["flag"]]
             + [f"- جرد: {x['name']} — {x['issue']}" for x in items if x["issue"]]
             + [f"- شراء بلا صنف مقابل: كود {p['code']} «{p['name']}» {p['total']:,.2f} — {p['basis']}"
                for p in R["purchases"] if p["item"] == NO_MATCH]
             + [f"- {p['basis'].split('⚠ ', 1)[1]} (سطر {p['idx']})" for p in R["purchases"] if p.get("outlier")])
    L += flags or ["- لا شىء"]
    p.write_text("\n".join(L) + "\n", encoding="utf-8")
    return p


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="متوسط التكلفة المرجّح لأى شهر من monthly/YYYY-MM/")
    ap.add_argument("month", help="YYYY-MM")
    ap.add_argument("--check", action="store_true", help="قائمة استلام ملفات الشهر فقط")
    a = ap.parse_args(argv)
    if not re.fullmatch(r"\d{4}-\d{2}", a.month):
        ap.error("الشهر بصيغة YYYY-MM")
    if a.check:
        missing = check_month(a.month)
        for path, desc in missing:
            print(f"ناقص: {path}  ({desc})")
        print("مكتمل" if not missing else f"{len(missing)} ملف ناقص")
        return 1 if missing else 0
    try:
        R = run(a.month)
    except MissingFiles as e:
        print(e, file=sys.stderr)
        return 1
    if R["bootstrap"]:
        print(f"{a.month} شهر البداية — إجمالى الخامات {R['raw_total']:,.2f}")
        return 0
    for p in write_outputs(R):
        print("كُتب", p.relative_to(ROOT))
    print(f"{a.month}: مشتريات {len(R['purchases'])} سطرًا {R['printed']:,.2f} | أول المدة {R['open_value']:,.2f} | "
          f"خامات {R['raw_value']:,.2f} | منتج تام {R['fg_value']:,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
