#!/usr/bin/env python3
"""متوسط التكلفة المرجّح (WAC) للخامات آخر أغسطس 2026.

المعادلة (دورية — Periodic WAC، نفس سلسلة يونيه ← يوليو):
    متوسط أغسطس = (كمية جرد 31/7 × متوسط يوليو + قيمة مشتريات أغسطس)
                 ÷ (كمية جرد 31/7 + كمية مشتريات أغسطس)

المدخلات (داخل المستودع):
    تسعير جرد اخر شهر ٧ سنة ٢٠٢٦.xlsx          رصيد أول المدة (جرد 31/7 + WAC يوليو) وربط أكواد يوليو
    تقرير المشتريات شهر ٨ سنة ٢٠٢٦.pdf          تقرير المشتريات 01/08–31/08
    جرد شهر ٨ سنة ٢٠٢٦.xlsx                     الجرد الفعلى 31/8 (خامات + منتج تام)
    ABOHASHEM_تسعير_المنتجات_بمتوسط_التكلفة_يونيه2026.xlsx   وصفات التصنيع (BOM)
    تسعير جرد اخر شهر ٦ سنة ٢٠٢٦.xlsx          أسعار يونيه (لربط خامات الوصفات بأصناف الجرد)

المخرجات:
    تسعير جرد اخر شهر ٨ سنة ٢٠٢٦.xlsx   الاحتساب الكامل بمعادلات Excel حيّة
    WAC_AUGUST_2026_REPORT.md            تقرير المراجعة

التشغيل:  python3 model/wac_august_2026.py   (يتطلب openpyxl و pdftotext)
"""
from __future__ import annotations

import re
import subprocess
import unicodedata
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
F_JUL = ROOT / "تسعير جرد اخر شهر ٧ سنة ٢٠٢٦.xlsx"
F_JUN = ROOT / "تسعير جرد اخر شهر ٦ سنة ٢٠٢٦.xlsx"
F_CNT = ROOT / "جرد شهر ٨ سنة ٢٠٢٦.xlsx"
F_PUR = ROOT / "تقرير المشتريات شهر ٨ سنة ٢٠٢٦.pdf"
F_BOM = ROOT / "ABOHASHEM_تسعير_المنتجات_بمتوسط_التكلفة_يونيه2026.xlsx"
OUT_AUG = ROOT / "تسعير جرد اخر شهر ٨ سنة ٢٠٢٦.xlsx"
OUT_MD = ROOT / "WAC_AUGUST_2026_REPORT.md"

JUL_RAW_TOTAL = 2016665.61      # إجمالى قيمة مخزون الخامات آخر يوليو كما بملف يوليو
POOL = "@مجمع اللحوم الهندى"
POOL_MEMBERS = ["لحم كتف", "لحم رقبه", "لحم سن", "لحمه ارباع"]
NO_MATCH = "— بلا مقابل —"

# ---------------------------------------------------------------- ربط أكواد أغسطس الجديدة
# الأكواد التى لم تظهر فى مشتريات يوليو: (الصنف المقابل بالجرد، درجة الثقة، أساس المطابقة)
NEW_CODES = {
    "87": ("جلتومات", "عالية", "جلوتامات الصوديوم = «جلتومات/مونوصوديم» بالوصفات؛ سعر 88 مقابل 104.00 ليوليو"),
    "380": ("كراتين اسبشيال احمر", "عالية", "نفس الكرتونة (اسبشيال أحمر)؛ سعر 23.36 ≈ 22.59 ليوليو"),
    "381": ("كراتين s.p ابيض", "عالية", "«اسباشيل» = s.p ، كرتونة بيضاء؛ الصنف بلا سعر سابق فيُسعَّر من الشراء"),
    "395": ("زيت سيارات", "متوسطة", "«عربيات» = سيارات؛ وحدة الشراء «كيلو» بسعر 1350 للعبوة 5 كيلو بينما الجرد بالعدد — يُراجَع"),
    "222": ("طعم كفته", "عالية", "تطابق اسمى تام"),
    "178": ("اكياس برجر الهنا ك", "عالية", "أكياس تغليف برجر الهنا — «كبير» = ك"),
    "179": ("اكياس سجق الهنا ص", "عالية", "أكياس تغليف سجق الهنا — «صغير» = ص"),
    "177": ("اكياس برجر الهنا ص", "عالية", "أكياس تغليف برجر الهنا — «صغير» = ص"),
    "181": ("اكياس كفتة الهنا ص", "عالية", "عبوة 350 جم = كيس كفته الهنا الصغير (ص)"),
    "299": ("جلد", "عالية", "جلد الفراخ هو «جلد» بالوصفات؛ سعر 28 ≈ 29.66 ليوليو"),
    "275": ("لون طبيعي", "عالية", "كارمن = لون طبيعى؛ سعر 3000 ≈ 2919.29 ليوليو"),
    "236": ("اكياس افيز", "عالية", "سعر 6.10 ≈ 6.31 ليوليو"),
    "78": ("فوسفات", "عالية", "تطابق اسمى تام"),
    "204": ("خلنجان", "عالية", "تطابق اسمى تام"),
    "367": ("اكياس برجر s.p كبير", "عالية", "«اسباشيل» = s.p ، «ك» = كبير"),
    "353": ("كراتين الهنا ازرق", "عالية", "نفس الكرتونة (الهنا أزرق)؛ سعر 24.06 ≈ 22.69 ليوليو"),
}
UNMATCHED_CODES = {
    "453": "بابريكا مدخنة صنف مختلف عن «بابريكا» العادية ولا يوجد له بند بالجرد — لم تُحمَّل على أى صنف",
}
# ملف يوليو حمّل بعض خامات الوصفات على صنف آخر تساوى سعره فى يونيه؛ تُستخدم فقط لإثبات أن
# محرك الوصفات يعيد إنتاج أرقام يوليو، أما أغسطس فيُسعَّر بالصنف الصحيح.
JULY_RECIPE_QUIRKS = {"سكر": "كربوناتو"}
# منتج تام بالجرد باسم مختلف عن يوليو لنفس الصنف
FG_ALIASES = {"كفته أبو هاشم 1 ك": ("كفتة ابو هاشم 1ك", "نفس الصنف «كفتة ابو هاشم 1ك» بفرق إملائى")}

# ---------------------------------------------------------------- نصوص
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
BIDI = re.compile("[‎‏‪-‮⁦-⁩]")


def clean(s) -> str:
    if s is None:
        return ""
    t = unicodedata.normalize("NFKC", str(s))
    return re.sub(r"\s+", " ", BIDI.sub("", t).translate(DIGITS)).strip()


def key(s) -> str:
    t = clean(s)
    for a, b in (("ة", "ه"), ("ى", "ي"), ("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ـ", "")):
        t = t.replace(a, b)
    return t


def num(v):
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(clean(v).replace(",", ""))
    except ValueError:
        return None


# ---------------------------------------------------------------- المدخلات
def load_july(path=F_JUL):
    wb = openpyxl.load_workbook(path)
    raw = OrderedDict()
    for r in wb["جرد الخامات آخر يوليو"].iter_rows(min_row=5, values_only=True):
        n = clean(r[0])
        if not n or n.startswith("إجمالى"):
            continue
        raw[key(n)] = {"name": n, "qty": num(r[9]) or 0.0, "wac": num(r[8]), "value": num(r[10])}
    codes = {}
    for r in wb["مشتريات يوليو التفصيلية"].iter_rows(min_row=4, values_only=True):
        if r[2] and r[7] and clean(r[7]) != NO_MATCH:
            c, item = clean(r[2]), clean(r[7])
            assert codes.get(c, (item,))[0] == item, f"July code {c} maps to two items"
            codes[c] = (item, clean(r[8]), "نفس ربط الكود فى يوليو: " + clean(r[9]))
    fg = OrderedDict()
    for r in wb["جرد المنتج التام آخر يوليو"].iter_rows(min_row=4, values_only=True):
        n = clean(r[0])
        if not n or n.startswith("إجمالى"):
            continue
        note = clean(r[6])
        m = re.search(r"من وصفة «(.+?)»: ([\d.]+) كجم × سعر الكجم(?: \+ تعبئة ([\d.]+))?", note)
        fg[key(n)] = {"name": n, "price": num(r[3]), "note": note,
                      "rule": (m.group(1), float(m.group(2)), float(m.group(3) or 0)) if m else None}
    kg = {}
    for r in wb["سعر الكجم من الوصفات"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            kg[clean(r[0])] = {"brand": clean(r[1]), "val_jun": r[2], "val_jul": r[3], "kg_jun": r[4], "kg_jul": r[5]}
    return raw, codes, fg, kg


def load_purchases(path=F_PUR):
    out = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True).stdout
    lines = [BIDI.sub("", unicodedata.normalize("NFKC", x)).translate(DIGITS) for x in out.decode().splitlines()]
    row = re.compile(r"^\s*(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+(\d{4}/\d{2}/\d{2})\s+([\d,]+(?:\.\d+)?)\s+"
                     r"([\d,]+(?:\.\d+)?)\s*(\d{6,})\s+(\d+)\s+(\S+)\s+(.+?)\s+(\d+)\s*$")
    rows, printed = [], None
    for ln in lines:
        m = row.match(ln)
        if m:
            _sale, price, d, total, qty, inv, code, unit, name, idx = m.groups()
            rows.append({"idx": int(idx), "date": d, "invoice": inv, "code": code, "unit": unit, "name": clean(name),
                         "qty": float(qty.replace(",", "")), "price": float(price), "total": float(total.replace(",", ""))})
        elif re.search(r"الجمالي\s*$", ln.strip()):
            printed = float(re.findall(r"[\d,]+\.\d+", ln)[0].replace(",", ""))
    assert [r["idx"] for r in rows] == list(range(1, len(rows) + 1)), "missing purchase rows"
    assert printed is not None and abs(sum(r["total"] for r in rows) - printed) < 0.005, "purchase total mismatch"
    return rows, printed


def load_count():
    ws = openpyxl.load_workbook(F_CNT).worksheets[0]
    raw, fg = [], []
    for r in range(5, ws.max_row + 1):
        n = clean(ws.cell(r, 1).value)
        if n:
            locs = [num(ws.cell(r, c).value) or 0.0 for c in range(2, 7)]
            g = ws.cell(r, 7).value
            qty = sum(locs)
            issue = None
            if isinstance(g, str) and g.replace(" ", "") != f"=SUM(B{r}:F{r})":
                issue = f"معادلة الإجمالى بالجرد «{g}» لا تغطى كل المواقع؛ باقى المواقع صفر فالنتيجة مطابقة لمجموع المواقع"
                if abs(sum(locs[:2]) - qty) > 1e-9:
                    issue = f"معادلة الإجمالى بالجرد «{g}» تُسقط مواقع بها رصيد — اعتُمد مجموع المواقع {qty:g}"
            elif not isinstance(g, str) and abs((num(g) or 0.0) - qty) > 1e-9:
                issue = f"الإجمالى مُدخل يدويًا ({num(g) or 0:g}) بينما مجموع المواقع {qty:g} — اعتُمد مجموع المواقع"
            raw.append({"row": r, "name": n, "locs": locs, "qty": qty, "issue": issue})
        f = clean(ws.cell(r, 8).value)
        if f and f != "0" and f != "اسم الصنف":
            fg.append({"row": r, "name": f, "qty": num(ws.cell(r, 9).value) or 0.0})
    return raw, fg


def load_june_prices(path=F_JUN):
    ws = openpyxl.load_workbook(path).worksheets[0]
    by_price = defaultdict(list)
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] and isinstance(r[7], (int, float)):
            by_price[round(r[7], 4)].append(clean(r[0]))
    return by_price


def load_recipes(jun_by_price, path=F_BOM):
    """الوصفات كما بملف يونيه: خلطة فرعية (توابيل) + وصفة رئيسية + معامل الإنتاجية."""
    wb = openpyxl.load_workbook(path)
    recipes = OrderedDict()
    for sheet in ("منتجات أبو هاشم", "منتجات سبيشال", "منتجات الهنا"):
        cur, sec = None, None
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            a, b = r[0], clean(r[1])
            if isinstance(a, str) and all(v is None for v in r[1:]):
                t = clean(a)
                if t.startswith("الخلطة الفرعية"):
                    sec = "sub"
                elif t.startswith("الوصفة الرئيسية"):
                    sec = "main"
                elif t.startswith("تكلفة التعبئة"):
                    sec = "pack"
                else:
                    cur, sec = t, None
                    recipes[cur] = {"sheet": sheet, "sub": [], "main": [], "div": None}
                continue
            if b == "سعر الكجم (خامات مباشرة فقط)":
                recipes[cur]["div"] = float(str(r[3]).split("/")[1])
            elif isinstance(a, (int, float)) and b and sec in ("sub", "main") and b != "الخامة":
                p = r[3]
                if isinstance(p, str) and p.startswith("="):
                    item, basis = "@SUB", "سعر كجم الخلطة الفرعية"
                elif not isinstance(p, (int, float)):
                    item, basis = None, clean(r[5]) or "بلا سعر بالوصفة"
                else:
                    cands = jun_by_price.get(round(p, 4), [])
                    if len(cands) > 1:
                        cands = [c for c in cands if key(c) == key(b)] or cands
                    if len(cands) == 4 and set(cands) == set(POOL_MEMBERS):
                        item, basis = POOL, "سعر مجمّع اللحوم الهندى"
                    elif len(cands) == 1:
                        item, basis = cands[0], "مطابقة سعر يونيه بالوصفة مع سعر الصنف بالجرد"
                    else:
                        item, basis = ("#CONST", p), "لا يقابله صنف بالجرد — السعر ثابت كما بالوصفة"
                recipes[cur][sec].append({"ing": b, "qty": float(r[2] or 0.0), "item": item, "basis": basis})
    for n, rc in recipes.items():
        assert rc["div"], f"no yield divisor for {n}"
    return recipes


# ---------------------------------------------------------------- الاحتساب
def compute_wac(jul, codes, purchases, count_raw):
    items = OrderedDict()
    for c in count_raw:
        j = jul.get(key(c["name"]))
        items[key(c["name"])] = {"name": c["name"], "row_cnt": c["row"], "open_q": j["qty"] if j else 0.0,
                                 "open_p": j["wac"] if j else None, "new": j is None,
                                 "close_q": c["qty"], "cnt_issue": c["issue"], "pq": 0.0, "pv": 0.0, "n": 0}
    pool = {"name": POOL, "pq": 0.0, "pv": 0.0, "n": 0,
            "open_q": sum(items[key(m)]["open_q"] for m in POOL_MEMBERS)}
    pool_prices = {items[key(m)]["open_p"] for m in POOL_MEMBERS}
    assert len(pool_prices) == 1, "pool members must share one July price"
    pool["open_p"] = pool_prices.pop()

    for p in purchases:
        if p["code"] in UNMATCHED_CODES:
            p.update(item=NO_MATCH, conf="—", basis=UNMATCHED_CODES[p["code"]])
            continue
        item, conf, basis = codes.get(p["code"]) or NEW_CODES[p["code"]]
        tgt = pool if item == POOL else items[key(item)]
        p.update(item=tgt["name"], conf=conf, basis=basis)   # نفس نص الاسم بالجرد (مطلوب لـ SUMIF)
        tgt["pq"] += p["qty"]
        tgt["pv"] += p["total"]
        tgt["n"] += 1

    def wac_of(x):
        if x["pq"] > 0:
            if x["open_p"] is None:
                return round(x["pv"] / x["pq"], 4)
            return round((x["open_q"] * x["open_p"] + x["pv"]) / (x["open_q"] + x["pq"]), 4)
        return x["open_p"]

    pool["wac"] = wac_of(pool)
    for k, x in items.items():
        x["wac"] = pool["wac"] if x["name"] in POOL_MEMBERS else wac_of(x)
        x["rule"] = rule_text(x)
    pool["rule"] = "مجمّع اللحوم الهندى: متوسط واحد يُطبَّق على كتف/رقبه/سن/أرباع"
    return items, pool


def rule_text(x):
    if x["name"] in POOL_MEMBERS:
        return "سعر موحَّد لمجمّع اللحوم الهندى (كتف/رقبه/سن/أرباع) — مشترياته «تشكيلة هندى أمامى»"
    if x["pq"] > 0 and x["open_p"] is None:
        return ("رصيد أول المدة بلا سعر مُثبت ⇒ المتوسط من المشتريات فقط" if x["open_q"] > 0
                else "لا رصيد أول مدة مُسعّر ⇒ المتوسط = متوسط سعر الشراء")
    if x["pq"] > 0:
        return "قياسى: (أول المدة + المشتريات)"
    if x["open_p"] is not None:
        return "لا مشتريات فى أغسطس ⇒ السعر ثابت كما فى يوليو"
    if x["close_q"] > 0:
        return "صنف بلا سعر مُثبت (لا سعر يوليو ولا مشتريات أغسطس) — يلزم إدخال سعر لتقييمه"
    return ""


def recipe_cost(rc, price_of):
    """قيمة الوصفة وسعر الكجم بأسعار price_of(item)."""
    def unit(i):
        it = i["item"]
        if it is None:
            return 0.0
        if isinstance(it, tuple):
            return it[1]
        if it == "@SUB":
            return sub_kg
        return price_of(it) or 0.0
    sub_q = sum(i["qty"] for i in rc["sub"])
    sub_kg = (sum(i["qty"] * unit(i) for i in rc["sub"]) / sub_q) if sub_q else 0.0
    value = sum(i["qty"] * unit(i) for i in rc["main"])
    return value, value / rc["div"], sub_kg


# ---------------------------------------------------------------- تنسيق
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="1F4E78")
SUBH = PatternFill("solid", fgColor="DDEBF7")
TOT = PatternFill("solid", fgColor="FFF2CC")
FLAG = PatternFill("solid", fgColor="FCE4D6")
NUM, PRICE, PCT = "#,##0.00", "#,##0.0000", "0.0%"
RTL = Alignment(horizontal="right", vertical="center", wrap_text=True)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def sheet(wb, title, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def title_rows(ws, title, subtitle, ncol):
    ws.cell(1, 1, title).font = Font(bold=True, size=13, color="1F4E78")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    if subtitle:
        ws.cell(2, 1, subtitle).font = Font(italic=True, size=10)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        ws.cell(2, 1).alignment = RTL


def header(ws, r, labels):
    for c, t in enumerate(labels, 1):
        x = ws.cell(r, c, t)
        x.font, x.fill, x.alignment, x.border = Font(bold=True, color="FFFFFF"), HEAD, CEN, BOX
    ws.row_dimensions[r].height = 42
    ws.freeze_panes = ws.cell(r + 1, 2)


def put(ws, r, c, v, fmt=None, fill=None, bold=False):
    x = ws.cell(r, c, v)
    x.border, x.alignment = BOX, RTL if c == 1 or isinstance(v, str) and not str(v).startswith("=") else CEN
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if bold:
        x.font = Font(bold=True)
    return x


# ---------------------------------------------------------------- مصنّف أغسطس
S_RAW, S_STEP, S_PUR = "جرد الخامات آخر أغسطس", "خطوات احتساب المتوسط", "مشتريات أغسطس التفصيلية"
S_FG, S_KG, S_BOM = "جرد المنتج التام آخر أغسطس", "سعر الكجم من الوصفات", "تفصيل الوصفات بأسعار أغسطس"
S_METH, S_REC = "المنهجية والاستثناءات", "المطابقة"


def q(name):
    return f"'{name}'"


def build_august(items, pool, purchases, printed, jul, fg_rows, recipes, kg_aug, kg_jul_tab):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # --- المشتريات التفصيلية
    wp = sheet(wb, S_PUR, [5, 11, 8, 30, 7, 10, 10, 13, 24, 9, 50, 12])
    title_rows(wp, "مشتريات الخامات — الفترة 01/08/2026 حتى 31/08/2026 (المصدر: تقرير المشتريات فى الفترة)", None, 12)
    header(wp, 3, ["م", "التاريخ", "كود الصنف", "اسم الصنف بالمشتريات", "الوحدة", "الكمية", "سعر الشراء",
                   "الإجمالى", "الصنف المقابل بالجرد", "درجة الثقة", "أساس المطابقة", "رقم الفاتورة"])
    for i, p in enumerate(purchases):
        r = 4 + i
        vals = [p["idx"], p["date"], p["code"], p["name"], p["unit"], p["qty"], p["price"], p["total"],
                p["item"], p["conf"], p["basis"], p["invoice"]]
        for c, v in enumerate(vals, 1):
            put(wp, r, c, v, NUM if c in (6, 8) else PRICE if c == 7 else None,
                FLAG if p["item"] == NO_MATCH or p["conf"] == "متوسطة" else None)
    p_first, p_last = 4, 3 + len(purchases)
    p_tot = p_last + 2
    put(wp, p_tot, 4, "إجمالى مشتريات أغسطس", fill=TOT, bold=True)
    put(wp, p_tot, 8, f"=SUM(H{p_first}:H{p_last})", NUM, TOT, True)
    put(wp, p_tot + 1, 4, "الإجمالى المطبوع بالتقرير", fill=TOT)
    put(wp, p_tot + 1, 8, printed, NUM, TOT)
    put(wp, p_tot + 2, 4, "منها: مشتريات بلا مقابل بالجرد (مستبعدة من المتوسط)", fill=TOT)
    put(wp, p_tot + 2, 8, f'=SUMIF(I{p_first}:I{p_last},"{NO_MATCH}",H{p_first}:H{p_last})', NUM, TOT)
    I_RNG, Q_RNG, V_RNG = (f"{q(S_PUR)}!$I${p_first}:$I${p_last}", f"{q(S_PUR)}!$F${p_first}:$F${p_last}",
                           f"{q(S_PUR)}!$H${p_first}:$H${p_last}")

    # --- جرد الخامات (الورقة الرئيسية)
    wr = sheet(wb, S_RAW, [30, 12, 11, 13, 12, 14, 12, 14, 13, 12, 14, 10, 55])
    wb.move_sheet(wr, -(len(wb.sheetnames) - 1))
    title_rows(wr, "ABOHASHEM FOR MEAT — تقييم مخزون الخامات آخر أغسطس 2026 بمتوسط التكلفة المرجّح (WAC)",
               "متوسط السعر = (قيمة رصيد أول المدة + قيمة مشتريات أغسطس) ÷ (كمية رصيد أول المدة + كمية مشتريات أغسطس)"
               " — رصيد أول المدة = جرد 31/7 مُقيَّمًا بمتوسط يوليو", 13)
    header(wr, 4, ["اسم الصنف", "كمية أول المدة (جرد 31/7)", "سعر أول المدة (متوسط يوليو)", "قيمة أول المدة",
                   "كمية مشتريات أغسطس", "قيمة مشتريات أغسطس", "إجمالى الكمية", "إجمالى التكلفة",
                   "متوسط السعر (WAC) 31/8", "كمية الجرد 31/8", "قيمة المخزون 31/8", "التغير عن سعر يوليو %", "ملاحظة"])
    raw_row = {}
    r = 5
    for k, x in items.items():
        raw_row[k] = r
        x["row_aug"] = r
        note = "; ".join(t for t in (x["rule"], x["cnt_issue"],
                                     "صنف جديد بجرد أغسطس لم يكن بجرد يوليو" if x["new"] else None) if t)
        fl = FLAG if (x["wac"] is None and x["close_q"] > 0) or x["cnt_issue"] else None
        put(wr, r, 1, x["name"], fill=fl)
        put(wr, r, 2, x["open_q"], NUM)
        put(wr, r, 3, x["open_p"], PRICE)
        put(wr, r, 4, f'=IF(C{r}="","",B{r}*C{r})', NUM)
        put(wr, r, 5, f"=SUMIF({I_RNG},A{r},{Q_RNG})", NUM)
        put(wr, r, 6, f"=SUMIF({I_RNG},A{r},{V_RNG})", NUM)
        put(wr, r, 7, f'=IF(C{r}="",E{r},B{r}+E{r})', NUM)
        put(wr, r, 8, f"=N(D{r})+F{r}", NUM)
        if x["name"] in POOL_MEMBERS:
            put(wr, r, 9, "=POOL_WAC", PRICE)       # يُستبدل بمرجع صف المجمّع بعد بناء ورقة الخطوات
        else:
            put(wr, r, 9, f'=IF(G{r}>0,ROUND(H{r}/G{r},4),IF(C{r}="","",C{r}))', PRICE)
        put(wr, r, 10, x["close_q"], NUM)
        put(wr, r, 11, f'=IF(I{r}="","",J{r}*I{r})', NUM)
        put(wr, r, 12, f'=IF(OR(C{r}="",I{r}=""),"",I{r}/C{r}-1)', PCT)
        put(wr, r, 13, note, fill=fl)
        r += 1
    raw_last = r - 1
    raw_tot = r + 1
    put(wr, raw_tot, 1, "الإجمالى (الأصناف المُسعّرة)", fill=TOT, bold=True)
    for c in (4, 6, 8, 11):
        L = get_column_letter(c)
        put(wr, raw_tot, c, f"=SUM({L}5:{L}{raw_last})", NUM, TOT, True)

    # --- خطوات احتساب المتوسط
    ws_ = sheet(wb, S_STEP, [34, 12, 12, 14, 12, 14, 11, 12, 14, 12, 9, 40])
    title_rows(ws_, "خطوات احتساب متوسط السعر المرجّح لكل خامة — أغسطس 2026",
               "تُعرض هنا الأصناف التى حدثت عليها مشتريات خلال أغسطس فقط (باقى الأصناف يبقى سعرها كما فى يوليو).", 12)
    header(ws_, 4, ["اسم الصنف", "خطوة ١: كمية أول المدة", "× سعر أول المدة", "قيمة أول المدة",
                    "خطوة ٢: كمية المشتريات", "قيمة المشتريات (بالفواتير)", "متوسط سعر الشراء",
                    "خطوة ٣: إجمالى الكمية", "خطوة ٤: إجمالى التكلفة", "خطوة ٥: متوسط السعر",
                    "عدد فواتير الشراء", "قاعدة الاحتساب"])
    step_rows, r = {}, 5
    for k, x in [(POOL, pool)] + [(k, x) for k, x in items.items() if x["pq"] > 0 and x["name"] not in POOL_MEMBERS]:
        step_rows[k] = r
        nm = POOL if k == POOL else x["name"]
        put(ws_, r, 1, nm if k != POOL else f"{POOL} (كتف/رقبه/سن/أرباع)")
        if k == POOL:
            put(ws_, r, 2, "=" + "+".join(f"{q(S_RAW)}!B{raw_row[key(m)]}" for m in POOL_MEMBERS), NUM)
            put(ws_, r, 3, f"={q(S_RAW)}!C{raw_row[key(POOL_MEMBERS[0])]}", PRICE)
            crit = f'"{POOL}"'
        else:
            put(ws_, r, 2, f"={q(S_RAW)}!B{raw_row[k]}", NUM)
            put(ws_, r, 3, f"={q(S_RAW)}!C{raw_row[k]}", PRICE)
            crit = f"{q(S_RAW)}!A{raw_row[k]}"
        put(ws_, r, 4, f'=IF(C{r}="","",B{r}*C{r})', NUM)
        put(ws_, r, 5, f"=SUMIF({I_RNG},{crit},{Q_RNG})", NUM)
        put(ws_, r, 6, f"=SUMIF({I_RNG},{crit},{V_RNG})", NUM)
        put(ws_, r, 7, f"=ROUND(F{r}/E{r},4)", PRICE)
        put(ws_, r, 8, f'=IF(C{r}="",E{r},B{r}+E{r})', NUM)
        put(ws_, r, 9, f"=N(D{r})+F{r}", NUM)
        put(ws_, r, 10, f"=ROUND(I{r}/H{r},4)", PRICE, SUBH, True)
        put(ws_, r, 11, f"=COUNTIF({I_RNG},{crit})")
        put(ws_, r, 12, x["rule"])
        r += 1
    step_last = r - 1
    put(ws_, r + 1, 1, "الإجمالى", fill=TOT, bold=True)
    for c in (4, 6, 9):
        L = get_column_letter(c)
        put(ws_, r + 1, c, f"=SUM({L}5:{L}{step_last})", NUM, TOT, True)
    pool_ref = f"{q(S_STEP)}!J{step_rows[POOL]}"
    for m in POOL_MEMBERS:
        wr.cell(raw_row[key(m)], 9).value = "=" + pool_ref

    # --- تفصيل الوصفات بأسعار أغسطس
    wbm = sheet(wb, S_BOM, [18, 12, 22, 24, 10, 12, 13, 50])
    title_rows(wbm, "تفصيل وصفات التصنيع بأسعار خامات أغسطس (نفس كميات ومعاملات إنتاجية وصفات يونيه)", None, 8)
    header(wbm, 3, ["الوصفة", "القسم", "الخامة بالوصفة", "الصنف المقابل بالجرد", "الكمية (كجم)",
                    "سعر أغسطس", "القيمة", "ملاحظة"])
    r, bom_kg = 4, {}
    for rn, rc in recipes.items():
        sub_kg_cell = None
        for sec, label in (("sub", "خلطة فرعية (توابل)"), ("main", "الوصفة الرئيسية")):
            if not rc[sec]:
                continue
            first = r
            for i in rc[sec]:
                it = i["item"]
                put(wbm, r, 1, rn)
                put(wbm, r, 2, label)
                put(wbm, r, 3, i["ing"])
                put(wbm, r, 5, i["qty"], "0.000")
                if it is None:
                    put(wbm, r, 4, "—")
                    put(wbm, r, 6, 0, PRICE)
                elif isinstance(it, tuple):
                    put(wbm, r, 4, "— (سعر ثابت بالوصفة)")
                    put(wbm, r, 6, it[1], PRICE)
                elif it == "@SUB":
                    put(wbm, r, 4, "سعر كجم الخلطة الفرعية")
                    put(wbm, r, 6, f"={sub_kg_cell}", PRICE)
                elif it == POOL:
                    put(wbm, r, 4, POOL)
                    put(wbm, r, 6, "=" + pool_ref, PRICE)
                else:
                    put(wbm, r, 4, it)
                    put(wbm, r, 6, f"={q(S_RAW)}!I{raw_row[key(it)]}", PRICE)
                put(wbm, r, 7, f"=E{r}*F{r}", NUM)
                put(wbm, r, 8, i["basis"], fill=FLAG if it is None or isinstance(it, tuple) else None)
                r += 1
            put(wbm, r, 1, rn, fill=SUBH)
            put(wbm, r, 2, f"إجمالى {label}", fill=SUBH, bold=True)
            put(wbm, r, 5, f"=SUM(E{first}:E{r - 1})", "0.000", SUBH)
            put(wbm, r, 7, f"=SUM(G{first}:G{r - 1})", NUM, SUBH, True)
            if sec == "sub":
                put(wbm, r, 6, f"=G{r}/E{r}", PRICE, SUBH, True)
                put(wbm, r, 8, "سعر كجم الخلطة الفرعية = قيمتها ÷ كميتها", fill=SUBH)
                sub_kg_cell = f"F{r}"
            else:
                bom_kg[rn] = (f"{q(S_BOM)}!G{r}", rc["div"])
                put(wbm, r, 8, f"سعر الكجم = قيمة الوصفة ÷ {rc['div']:g} كجم (معامل الإنتاجية كما بالوصفة)", fill=SUBH)
            r += 2

    # --- سعر الكجم من الوصفات
    wk = sheet(wb, S_KG, [16, 10, 16, 16, 13, 13, 10, 14])
    title_rows(wk, "أثر تحديث أسعار الخامات على سعر كيلو الخلطة لكل منتج (خامات مباشرة فقط)", None, 8)
    header(wk, 3, ["المنتج", "العلامة", "قيمة الوصفة (يوليو)", "قيمة الوصفة (أغسطس)", "سعر الكجم (يوليو)",
                   "سعر الكجم (أغسطس)", "التغير %", "معامل الإنتاجية (كجم)"])
    kg_row, r = {}, 4
    for rn in recipes:
        kg_row[rn] = r
        t = kg_jul_tab[rn]
        put(wk, r, 1, rn)
        put(wk, r, 2, t["brand"])
        put(wk, r, 3, t["val_jul"], NUM)
        put(wk, r, 4, "=" + bom_kg[rn][0], NUM)
        put(wk, r, 5, t["kg_jul"], PRICE)
        put(wk, r, 6, f"=D{r}/H{r}", PRICE, SUBH, True)
        put(wk, r, 7, f"=F{r}/E{r}-1", PCT)
        put(wk, r, 8, recipes[rn]["div"], "0.0000")
        r += 1

    # --- جرد المنتج التام
    wf = sheet(wb, S_FG, [30, 12, 12, 14, 10, 14, 70])
    title_rows(wf, "تقييم مخزون المنتج التام آخر أغسطس 2026 (وصفات التصنيع بأسعار خامات أغسطس + تكلفة التعبئة)", None, 7)
    header(wf, 3, ["اسم الصنف", "كمية الجرد 31/8", "سعر يوليو", "سعر أغسطس المحدَّث", "التغير %",
                   "قيمة المخزون 31/8", "ملاحظة"])
    r = 4
    for f in fg_rows:
        put(wf, r, 1, f["name"], fill=FLAG if f["flag"] else None)
        put(wf, r, 2, f["qty"], NUM)
        put(wf, r, 3, f["jul_price"], PRICE)
        if f["rule"]:
            rn, w, pkg = f["rule"]
            put(wf, r, 4, f"=ROUND({w:g}*{q(S_KG)}!F{kg_row[rn]}+{pkg:g},4)", PRICE)
        else:
            put(wf, r, 4, f["aug_price"], PRICE)
        put(wf, r, 5, f'=IF(OR(C{r}="",D{r}=""),"",D{r}/C{r}-1)', PCT)
        put(wf, r, 6, f'=IF(D{r}="","",B{r}*D{r})', NUM)
        put(wf, r, 7, f["note"], fill=FLAG if f["flag"] else None)
        f["row_aug"] = r
        r += 1
    fg_last = r - 1
    fg_tot = r + 1
    put(wf, fg_tot, 1, "إجمالى قيمة مخزون المنتج التام آخر أغسطس 2026 (الأصناف المُسعّرة)", fill=TOT, bold=True)
    put(wf, fg_tot, 6, f"=SUM(F4:F{fg_last})", NUM, TOT, True)
    put(wf, fg_tot + 2, 1, "إجمالى قيمة المخزون الكلى آخر أغسطس 2026 (خامات + منتج تام)", fill=TOT, bold=True)
    put(wf, fg_tot + 2, 6, f"=F{fg_tot}+{q(S_RAW)}!K{raw_tot}", NUM, TOT, True)

    # --- المطابقة
    wc = sheet(wb, S_REC, [62, 18, 18, 14, 12])
    title_rows(wc, "المطابقة النهائية — كل فحص يجب أن يكون الفرق صفرًا", None, 5)
    header(wc, 3, ["الفحص", "القيمة المحسوبة", "القيمة المرجعية", "الفرق", "النتيجة"])
    checks = [
        ("إجمالى قيمة أول المدة = إجمالى قيمة مخزون الخامات آخر يوليو (ملف يوليو)",
         f"={q(S_RAW)}!D{raw_tot}", JUL_RAW_TOTAL, 1.0),
        ("إجمالى أسطر المشتريات = الإجمالى المطبوع بتقرير المشتريات",
         f"={q(S_PUR)}!H{p_tot}", f"={q(S_PUR)}!H{p_tot + 1}", 0.005),
        ("المشتريات المحمّلة على الأصناف + المستبعدة = إجمالى المشتريات",
         f"={q(S_RAW)}!F{raw_tot}+{q(S_STEP)}!F{step_rows[POOL]}+{q(S_PUR)}!H{p_tot + 2}",
         f"={q(S_PUR)}!H{p_tot}", 0.005),
        ("قيمة المشتريات بورقة الخطوات = المشتريات المحمّلة بورقة الجرد (شاملة المجمّع)",
         f"={q(S_STEP)}!F{step_last + 2}", f"={q(S_RAW)}!F{raw_tot}+{q(S_STEP)}!F{step_rows[POOL]}", 0.005),
        ("عدد الأصناف التى يختلف متوسطها بين ورقة الجرد وورقة الخطوات",
         "=" + "+".join(f"(ABS({q(S_RAW)}!I{raw_row[k]}-{q(S_STEP)}!J{rr})>0.00005)"
                        for k, rr in step_rows.items() if k != POOL), 0, 0),
    ]
    for i, (label, a, b, tol) in enumerate(checks):
        rr = 4 + i
        put(wc, rr, 1, label)
        put(wc, rr, 2, a, NUM)
        put(wc, rr, 3, b, NUM)
        put(wc, rr, 4, f"=B{rr}-C{rr}", NUM)
        put(wc, rr, 5, f'=IF(ABS(D{rr})<={tol},"✔ مطابق","✘ فرق")')

    wb.save(OUT_AUG)
    return {"raw_row": raw_row, "raw_tot": raw_tot, "fg_tot": fg_tot, "step_rows": step_rows}


# ---------------------------------------------------------------- منهجية + تقرير
def methodology(items, purchases, fg_rows, recipes):
    rules = [
        ("نطاق العمل", "متوسط التكلفة المرجّح (Weighted Average Cost) للخامات فقط — لم تُستخدم FIFO / LIFO / التكلفة المعيارية."),
        ("المعادلة", "متوسط أغسطس = (كمية جرد 31/7 × متوسط يوليو + قيمة مشتريات أغسطس) ÷ (كمية جرد 31/7 + كمية مشتريات أغسطس). "
                     "الطريقة دورية: متوسط واحد للشهر يُطبَّق على رصيد آخر المدة؛ يُقرَّب السعر لأربعة أرقام عشرية."),
        ("رصيد أول المدة", "كمية «جرد 31/7» وسعر «متوسط السعر (WAC) 31/7» من «تسعير جرد اخر شهر ٧ سنة ٢٠٢٦.xlsx» — استمرار لنفس السلسلة."),
        ("المشتريات", f"تقرير المشتريات 01/08/2026–31/08/2026: {len(purchases)} سطر شراء بإجمالى مطابق للإجمالى المطبوع. "
                      "قيمة الشراء المعتمدة هى «الإجمالى» بالفاتورة. عمود «سعر البيع الحالى» بالتقرير لا يدخل فى أى حساب."),
        ("ربط أكواد الشراء", "الكود المستخدم فى يوليو يُربط بنفس صنف الجرد؛ الأكواد الجديدة رُبطت بالاسم ونطاق السعر مع درجة ثقة "
                             "وأساس مطابقة مُسجَّل لكل سطر. عمود «الوحدة» بالتقرير تسمية نظام (يظهر «طبق» على أصناف بالكيلو) فلم يُعتمد للمطابقة."),
        ("مجمّع اللحوم الهندى", "بنود (لحم كتف / لحم رقبه / لحم سن / لحمه ارباع) سعرها موحَّد، ومشترياتها تُقيَّد باسم «تشكيلة هندى أمامى» — "
                                "وعاء تكلفة واحد ومتوسط واحد يُطبَّق على البنود الأربعة (نفس معالجة يوليو)."),
        ("أصناف بلا مشتريات", "يبقى سعر يوليو كما هو — المتوسط المرجّح لا يتغير بغياب حركة شراء."),
        ("أصناف برصيد بلا سعر", "رصيد أول المدة بلا سعر مُثبت يُستبعد من بسط ومقام المتوسط ويُحتسب السعر من المشتريات فقط."),
        ("كمية جرد 31/8", "مجموع أعمدة المواقع الخمسة (المصنع، النفق، النهضه، المنوفية، التغليف) بملف الجرد؛ صفوف معادلة الإجمالى المختلفة مذكورة أدناه."),
        ("المنتج التام", "سعر كل SKU = الوزن × سعر كجم الوصفة بأسعار أغسطس + تكلفة التعبئة — نفس الوصفات ومعاملات الإنتاجية "
                         "وتكاليف التعبئة المستخدمة فى يوليو دون تغيير. الأصناف خارج الوصفات التسع يُبقى سعر يوليو."),
        ("سبتمبر", "تقييم جرد آخر سبتمبر فى «model/wac_september_2026.py» (رصيد أول المدة = جرد 31/8 × متوسط أغسطس + مشتريات سبتمبر)."),
        ("خارج النطاق", "الأسعار تمثل تكلفة الخامات المباشرة (+ التعبئة للمنتج التام) ولا تشمل الأجور ولا الأعباء الصناعية غير المباشرة."),
    ]
    flags, seen = [], set()
    for p in purchases:
        if p["code"] in seen:
            continue
        seen.add(p["code"])
        if p["item"] == NO_MATCH:
            flags.append(("شراء بلا مقابل", f"كود {p['code']} «{p['name']}» — {p['total']:,.2f} ج.م", p["basis"]))
        elif p["conf"] == "متوسطة":
            flags.append(("ربط بثقة متوسطة", f"كود {p['code']} «{p['name']}» ← {p['item']}", p["basis"]))
    for x in items.values():
        if x["wac"] is None and x["close_q"] > 0:
            flags.append(("صنف بلا سعر", f"{x['name']} — كمية جرد 31/8 = {x['close_q']:,.2f}", x["rule"]))
        if x["cnt_issue"]:
            flags.append(("ملاحظة على ملف الجرد", x["name"], x["cnt_issue"]))
        if x["new"]:
            flags.append(("صنف جديد بالجرد", x["name"], "لم يكن بجرد يوليو — لا رصيد أول مدة"))
    for f in fg_rows:
        if f["flag"]:
            flags.append(("منتج تام", f"{f['name']} — كمية {f['qty']:,.2f}", f["note"]))
    for rn, rc in recipes.items():
        for i in rc["sub"] + rc["main"]:
            if isinstance(i["item"], tuple):
                flags.append(("وصفة", f"{rn}: {i['ing']}", f"{i['basis']} ({i['item'][1]:g} ج.م)"))
    for wrong_item_for, used in JULY_RECIPE_QUIRKS.items():
        users = sorted({rn for rn, rc in recipes.items() for i in rc["sub"] + rc["main"] if i["item"] == wrong_item_for})
        flags.append(("تصحيح عن يوليو", f"خامة «{wrong_item_for}» بوصفات: {'، '.join(users)}",
                      f"ملف يوليو سعّرها بسعر «{used}» ({items[key(used)]['wac']:g} ج.م) لتساوى السعرين فى يونيه؛ "
                      f"فى أغسطس سُعّرت بصنف «{wrong_item_for}» نفسه ({items[key(wrong_item_for)]['wac']:g} ج.م)"))
    return rules, flags


def add_methodology(rules, flags):
    wb = openpyxl.load_workbook(OUT_AUG)
    ws = sheet(wb, S_METH, [24, 55, 70])
    title_rows(ws, "المنهجية، الافتراضات، والاستثناءات", None, 3)
    r = 3
    for a, b in rules:
        put(ws, r, 1, a, bold=True)
        put(ws, r, 2, b)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 45
        r += 1
    r += 1
    header(ws, r, ["نوع الملاحظة", "الصنف / البند", "التفصيل"])
    ws.freeze_panes = None
    for t in flags:
        r += 1
        for c, v in enumerate(t, 1):
            put(ws, r, c, v, fill=FLAG)
    wb.save(OUT_AUG)


# ---------------------------------------------------------------- التحقق المستقل
def independent_wac(jul, purchases):
    """إعادة احتساب مستقلة مباشرة من أسطر الشراء (بدون هياكل compute_wac)."""
    agg = defaultdict(lambda: [0.0, 0.0])
    for p in purchases:
        if p["item"] != NO_MATCH:
            agg[key(p["item"])][0] += p["qty"]
            agg[key(p["item"])][1] += p["total"]
    out = {}
    pq, pv = agg.pop(key(POOL), [0.0, 0.0])
    oq = sum(jul[key(m)]["qty"] for m in POOL_MEMBERS)
    op = jul[key(POOL_MEMBERS[0])]["wac"]
    pool_w = round((oq * op + pv) / (oq + pq), 4) if pq else op
    for k, j in jul.items():
        q_, v_ = agg.get(k, (0.0, 0.0))
        if j["name"] in POOL_MEMBERS:
            out[k] = pool_w
        elif q_ and j["wac"] is not None:
            out[k] = round((j["qty"] * j["wac"] + v_) / (j["qty"] + q_), 4)
        elif q_:
            out[k] = round(v_ / q_, 4)
        else:
            out[k] = j["wac"]
    for k, (q_, v_) in agg.items():
        if k not in out:
            out[k] = round(v_ / q_, 4)
    return out


def recalc_and_verify(items, fg_rows, kg_aug, idx):
    """إعادة احتساب المعادلات بـ LibreOffice والتحقق من تطابق النتائج مع حساب Python."""
    import shutil
    import tempfile
    soffice = shutil.which("soffice")
    if not soffice:
        print("! LibreOffice غير متوفر — تخطّى فحص المعادلات")
        return None
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, str(OUT_AUG)],
                       capture_output=True, check=True, timeout=180)
        wb = openpyxl.load_workbook(Path(d) / OUT_AUG.name, data_only=True)
    wr, wf, wk, wc = wb[S_RAW], wb[S_FG], wb[S_KG], wb[S_REC]
    errs = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
    assert not errs, f"formula errors: {errs[:5]}"
    for k, x in items.items():
        got = wr.cell(x["row_aug"], 9).value
        got = None if got in ("", None) else round(got, 4)
        assert got == x["wac"], (x["name"], got, x["wac"])
    for r in range(4, 4 + len(kg_aug)):
        rn = wk.cell(r, 1).value
        assert abs(wk.cell(r, 6).value - kg_aug[rn][1]) < 1e-6, (rn, wk.cell(r, 6).value, kg_aug[rn][1])
    for f in fg_rows:
        got = wf.cell(f["row_aug"], 4).value
        got = None if got in ("", None) else round(got, 4)
        assert got == f["aug_price"], (f["name"], got, f["aug_price"])
    results = [(wc.cell(r, 1).value, wc.cell(r, 2).value, wc.cell(r, 3).value, wc.cell(r, 5).value)
               for r in range(4, wc.max_row + 1)]
    for lab, a, b, res in results:
        assert res == "✔ مطابق", (lab, a, b, res)
    raw_total = wr.cell(idx["raw_tot"], 11).value
    fg_total = wf.cell(idx["fg_tot"], 6).value
    open_total = wr.cell(idx["raw_tot"], 4).value
    return {"raw_total": raw_total, "fg_total": fg_total, "open_total": open_total, "checks": results}


def write_report(items, pool, purchases, printed, fg_rows, kg_aug, kg_jul_tab, totals, flags, codes_by_item):
    L = []
    w = L.append
    mapped = sum(p["total"] for p in purchases if p["item"] != NO_MATCH)
    w("# ABOHASHEM FOR MEAT — متوسط التكلفة المرجّح للخامات آخر أغسطس 2026\n")
    w("**الملفات:** `تسعير جرد اخر شهر ٨ سنة ٢٠٢٦.xlsx` (الاحتساب الكامل بمعادلات حيّة) — "
      "السكربت `model/wac_august_2026.py`. تقييم سبتمبر: `WAC_SEPTEMBER_2026_REPORT.md`.\n")
    w("## ١. المعادلة\n")
    w("```\nمتوسط أغسطس = (كمية جرد 31/7 × متوسط يوليو + قيمة مشتريات أغسطس)\n"
      "             ÷ (كمية جرد 31/7 + كمية مشتريات أغسطس)\n```\n")
    w("- رصيد أول المدة = جرد 31/7 مُقيَّمًا بمتوسط يوليو (استمرار سلسلة يونيه ← يوليو).\n"
      "- قيمة الشراء = «الإجمالى» بالفاتورة. كل الأسعار مقرّبة لـ 4 أرقام عشرية.\n"
      "- لا مشتريات ⇒ سعر يوليو ثابت. رصيد بلا سعر ⇒ يُستبعد من البسط والمقام.\n")
    w("## ٢. الإجماليات\n")
    w("| البند | القيمة (ج.م) |\n|---|---:|")
    w(f"| قيمة رصيد أول المدة (جرد 31/7 × متوسط يوليو) | {totals['open_total']:,.2f} |")
    w(f"| مشتريات أغسطس (إجمالى التقرير) | {printed:,.2f} |")
    w(f"| منها محمّلة على أصناف الجرد | {mapped:,.2f} |")
    w(f"| منها بلا مقابل بالجرد (مستبعدة) | {printed - mapped:,.2f} |")
    w(f"| **قيمة مخزون الخامات آخر أغسطس** | **{totals['raw_total']:,.2f}** |")
    w(f"| **قيمة مخزون المنتج التام آخر أغسطس** | **{totals['fg_total']:,.2f}** |")
    w(f"| **إجمالى المخزون آخر أغسطس** | **{totals['raw_total'] + totals['fg_total']:,.2f}** |\n")
    w("## ٣. جدول المراجعة — الأصناف التى عليها مشتريات فى أغسطس\n")
    w("| كود الشراء | الصنف | كمية أول المدة | قيمة أول المدة | كمية المشتريات | قيمة المشتريات | إجمالى الكمية | إجمالى القيمة | المتوسط 31/8 | كمية جرد 31/8 | قيمة 31/8 |")
    w("|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|")
    rows = [(POOL, pool, sum(items[key(m)]["close_q"] for m in POOL_MEMBERS))] + \
           [(x["name"], x, x["close_q"]) for x in items.values() if x["pq"] > 0 and x["name"] not in POOL_MEMBERS]
    for name, x, cq in rows:
        oq = x["open_q"] if x["open_p"] is not None else 0.0
        ov = oq * x["open_p"] if x["open_p"] is not None else 0.0
        tq, tv = oq + x["pq"], ov + x["pv"]
        w(f"| {'، '.join(sorted(codes_by_item.get(key(name), [])))} | {name} | {x['open_q']:,.2f} | {ov:,.2f} | "
          f"{x['pq']:,.2f} | {x['pv']:,.2f} | {tq:,.2f} | {tv:,.2f} | {x['wac']:,.4f} | {cq:,.2f} | {cq * x['wac']:,.2f} |")
    w("\nباقى الأصناف (بلا مشتريات فى أغسطس) احتفظت بسعر يوليو؛ التفصيل الكامل لكل الأصناف بورقة «جرد الخامات آخر أغسطس».\n")
    big = [(name, x) for name, x, _ in rows if x["open_p"] and abs(x["wac"] / x["open_p"] - 1) >= 0.10]
    if big:
        w("**أصناف تغيّر متوسطها 10% أو أكثر عن يوليو (للمراجعة):**\n")
        w("| الصنف | متوسط يوليو | متوسط أغسطس | التغير | أسعار الشراء بأغسطس |\n|---|---:|---:|---:|---|")
        for name, x in sorted(big, key=lambda t: t[1]["wac"] / t[1]["open_p"] - 1):
            prices = sorted({p["price"] for p in purchases if p["item"] == name})
            w(f"| {name} | {x['open_p']:.4f} | {x['wac']:.4f} | {x['wac'] / x['open_p'] - 1:+.1%} | "
              f"{' / '.join(f'{v:g}' for v in prices)} |")
        w("")
    w("## ٤. أمثلة محسوبة يدويًا\n")
    for nm in ("بصل", "ثوم", "مفروم فراخ شعبي"):
        x = items[key(nm)]
        ov = x["open_q"] * x["open_p"]
        w(f"**{nm}:** أول المدة {x['open_q']:,.2f} × {x['open_p']:.4f} = {ov:,.4f} ؛ "
          f"المشتريات {x['pq']:,.2f} كجم بقيمة {x['pv']:,.2f} ؛ "
          f"الإجمالى ({ov:,.4f} + {x['pv']:,.2f}) ÷ ({x['open_q']:,.2f} + {x['pq']:,.2f}) "
          f"= {ov + x['pv']:,.4f} ÷ {x['open_q'] + x['pq']:,.2f} = **{x['wac']:.4f}** ج.م "
          f"(يوليو {x['open_p']:.4f}، التغير {x['wac'] / x['open_p'] - 1:+.2%}).\n")
    w("## ٥. سعر كجم الوصفات (خامات مباشرة)\n")
    w("| المنتج | سعر الكجم يوليو | سعر الكجم أغسطس | التغير |\n|---|---:|---:|---:|")
    for rn, (val, kg) in kg_aug.items():
        kj = kg_jul_tab[rn]["kg_jul"]
        w(f"| {rn} | {kj:.4f} | {kg:.4f} | {kg / kj - 1:+.2%} |")
    w("\n## ٦. المطابقة النهائية\n")
    w("| الفحص | النتيجة |\n|---|---|")
    for lab, a, b, res in totals["checks"]:
        w(f"| {lab} | {res} |")
    w("| كل متوسط مكتوب بالمصنّف (الجرد، الخطوات) = الاحتساب المستقل بالسكربت لأربعة أرقام عشرية | ✔ مطابق |")
    w("| إعادة احتساب كل المعادلات بـ LibreOffice — لا أخطاء #DIV/0! أو #REF! | ✔ |\n")
    w("## ٧. البنود المُعلَّمة للمراجعة\n")
    w("| النوع | الصنف / البند | التفصيل |\n|---|---|---|")
    for t in flags:
        w("| " + " | ".join(str(v).replace("|", "/") for v in t) + " |")
    OUT_MD.write_text("\n".join(L) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- التشغيل
def main():
    jul, codes, jul_fg, kg_jul_tab = load_july()
    purchases, printed = load_purchases()
    count_raw, count_fg = load_count()
    jun_by_price = load_june_prices()
    recipes = load_recipes(jun_by_price)

    items, pool = compute_wac(jul, codes, purchases, count_raw)

    # التحقق: محرك الوصفات يعيد إنتاج أرقام يونيه ويوليو كما بملف يوليو
    jun_price = {n: p for p, ns in jun_by_price.items() for n in ns}
    jul_price = {k: j["wac"] for k, j in jul.items()}
    kg_jul_exact = {}
    for rn, rc in recipes.items():
        v6, kg6, _ = recipe_cost(rc, lambda it: jun_price.get(POOL_MEMBERS[0] if it == POOL else it))
        v7, kg7, _ = recipe_cost(rc, lambda it: jul_price.get(
            key(POOL_MEMBERS[0] if it == POOL else JULY_RECIPE_QUIRKS.get(it, it))))
        t = kg_jul_tab[rn]
        assert abs(v6 - t["val_jun"]) < 0.01 and abs(round(kg6, 4) - t["kg_jun"]) < 0.00015, (rn, v6, t)
        assert abs(v7 - t["val_jul"]) < 0.01 and abs(round(kg7, 4) - t["kg_jul"]) < 0.00015, (rn, v7, t)
        kg_jul_exact[rn] = kg7
    aug_price = {k: x["wac"] for k, x in items.items()}
    kg_aug = OrderedDict()
    for rn, rc in recipes.items():
        v8, kg8, _ = recipe_cost(rc, lambda it: aug_price.get(key(POOL_MEMBERS[0] if it == POOL else it)))
        kg_aug[rn] = (v8, kg8)

    # المنتج التام
    fg_rows, fg_dev = [], (0.0, "")
    for f in count_fg:
        alias = FG_ALIASES.get(f["name"])
        j = jul_fg.get(key(alias[0] if alias else f["name"]))
        row = {"name": f["name"], "qty": f["qty"], "jul_price": None, "rule": None, "aug_price": None, "flag": False}
        if j and j["rule"]:
            rn, wgt, pkg = j["rule"]
            # التحقق: قاعدة يوليو تعيد إنتاج سعر يوليو
            dev = abs(round(wgt * kg_jul_exact[rn] + pkg, 4) - j["price"])
            assert dev < 0.01, (f["name"], j)
            fg_dev = max(fg_dev, (dev, f["name"]))
            row.update(jul_price=j["price"], rule=j["rule"], aug_price=round(wgt * kg_aug[rn][1] + pkg, 4),
                       note=f"من وصفة «{rn}»: {wgt:g} كجم × سعر الكجم" + (f" + تعبئة {pkg:.2f} ج.م" if pkg else
                            " بدون تكلفة تعبئة (غير محددة بالوصفة) — السعر الفعلى أعلى قليلًا"))
        elif j and j["price"] is not None:
            row.update(jul_price=j["price"], aug_price=j["price"],
                       note="خارج نطاق الوصفات التسع — سعر يوليو مُبقًى دون تحديث")
        elif j:
            row.update(note=j["note"].split(" | ")[0] or "بلا وصفة ولا سعر سابق", flag=f["qty"] > 0)
        else:
            row.update(note="صنف جديد بجرد أغسطس بلا وصفة ولا سعر سابق — يلزم توفير سعره", flag=True)
        if alias:
            row["note"] += f" | {alias[1]}"
        fg_rows.append(row)

    # الاحتساب المستقل ومطابقته
    ind = independent_wac(jul, purchases)
    for k, x in items.items():
        if k in ind:
            assert ind[k] == x["wac"], (x["name"], ind[k], x["wac"])
        else:
            assert x["wac"] is None, x["name"]

    idx = build_august(items, pool, purchases, printed, jul, fg_rows, recipes, kg_aug, kg_jul_tab)
    rules, flags = methodology(items, purchases, fg_rows, recipes)
    add_methodology(rules, flags)

    totals = recalc_and_verify(items, fg_rows, kg_aug, idx)
    codes_by_item = defaultdict(set)
    for p in purchases:
        codes_by_item[key(p["item"])].add(p["code"])
    write_report(items, pool, purchases, printed, fg_rows, kg_aug, kg_jul_tab, totals, flags, codes_by_item)

    print(f"purchases: {len(purchases)} lines, total {printed:,.2f}")
    print(f"opening value {totals['open_total']:,.2f} (July file {JUL_RAW_TOTAL:,.2f})")
    print(f"raw stock 31/8 {totals['raw_total']:,.2f} | finished goods {totals['fg_total']:,.2f}")
    print(f"max deviation re-deriving July FG prices from recipes: {fg_dev[0]:.4f} ({fg_dev[1]})")
    print(f"items priced: {sum(1 for x in items.values() if x['wac'] is not None)}/{len(items)}; "
          f"flags {len(flags)}")


if __name__ == "__main__":
    main()
