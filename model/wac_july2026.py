# -*- coding: utf-8 -*-
"""
متوسط التكلفة المرجّح (WAC) للخامات — يوليو 2026
=================================================
المدخلات (كلها داخل المستودع):
  • «تسعير جرد اخر شهر ٦ سنة ٢٠٢٦.xlsx» → رصيد أول المدة (كمية + سعر WAC حتى يونيه).
  • «مشتريات شهر 7 سنه 2026.pdf»        → مشتريات الفترة 01/07 → 30/07/2026 (100 سطر).
  • «جرد شهر 7.xlsx»                    → كميات الجرد الفعلى آخر يوليو (المطلوب تسعيرها).
  • «ABOHASHEM_تسعير_المنتجات_بمتوسط_التكلفة_يونيه2026.xlsx» → وصفات التصنيع (BOM)
    لإعادة تسعير المنتج التام بأسعار الخامات الجديدة.

المعادلة المطبَّقة لكل صنف:
    إجمالى الكمية  = كمية أول المدة + كمية المشتريات
    إجمالى التكلفة = قيمة أول المدة + قيمة المشتريات
    متوسط السعر    = إجمالى التكلفة ÷ إجمالى الكمية
    قيمة المخزون آخر يوليو = كمية الجرد الفعلى آخر يوليو × متوسط السعر

المخرجات: «تسعير جرد اخر شهر ٧ سنة ٢٠٢٦.xlsx» + «WAC_JULY2026_REPORT.md»
"""
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
F_JUNE = ROOT / "تسعير جرد اخر شهر ٦ سنة ٢٠٢٦.xlsx"
F_JULY = ROOT / "جرد شهر 7.xlsx"
F_PURCH = ROOT / "مشتريات شهر 7 سنه 2026.pdf"
F_BOM = ROOT / "ABOHASHEM_تسعير_المنتجات_بمتوسط_التكلفة_يونيه2026.xlsx"
OUT_XLSX = ROOT / "تسعير جرد اخر شهر ٧ سنة ٢٠٢٦.xlsx"
OUT_MD = ROOT / "WAC_JULY2026_REPORT.md"

# مجمّع اللحوم الهندى: أسعار يونيه لبنوده الأربعة متطابقة (253) أى أنها تُغذَّى من
# مشتريات واحدة («تشكيلة هندى أمامى») ولا يُمسك لها سعر منفصل.
MEAT_POOL = ["لحم كتف", "لحم رقبه", "لحم سن", "لحمه ارباع"]
POOL_KEY = "@مجمع اللحوم الهندى"


# ─────────────────────────── أدوات مساعدة ───────────────────────────
def norm(s):
    """توحيد النص العربى: إزالة التشكيل/التطويل، توحيد الألف والياء والتاء المربوطة."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    s = re.sub(r"[ً-ْـ]", "", s)
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = s.replace("ى", "ي").replace("ة", "ه")
    return re.sub(r"\s+", " ", s).strip()


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def r2(x):
    return None if x is None else round(x + 0.0, 2)


def r4(x):
    return None if x is None else round(x + 0.0, 4)


# ─────────────────── ١) قراءة مشتريات يوليو من الـPDF ───────────────────
def read_purchases():
    """استخراج بنود فاتورة الشراء من تقرير المشتريات (نص PDF مخزَّن بترتيب معكوس)."""
    import pdfplumber

    def fix(t):
        return t[::-1] if any("؀" <= c <= "ۿ" for c in t) else t

    # كلمات تذييل الصفحة تتداخل مع عمود اسم الصنف عند بعض الأسطر
    FOOTER = {norm(w) for w in ["الموقع", "العنوان", "الالكتروني", "الكتروني",
                               "الايميل", "العبور", "الجمالي", "الاجمالي"]}
    recs = []
    with pdfplumber.open(F_PURCH) as pdf:
        for page in pdf.pages:
            lines = {}
            for w in page.extract_words():
                for k in lines:
                    if abs(k - w["top"]) < 4:
                        lines[k].append(w)
                        break
                else:
                    lines[w["top"]] = [w]
            merged = []
            for key in sorted(lines):
                ln = lines[key]
                if any(w["x0"] > 555 for w in ln) or not merged:
                    merged.append(list(ln))
                else:
                    merged[-1].extend(ln)
            for ln in merged:
                serial = [w for w in ln if w["x0"] > 555]
                if not serial or not re.fullmatch(r"\d+", serial[0]["text"]):
                    continue
                col = lambda lo, hi: [w for w in ln if lo <= w["x0"] < hi]
                qty, tot, prc = col(225, 255), col(180, 225), col(60, 110)
                if not (qty and tot and prc):
                    continue
                words = sorted(col(400, 555), key=lambda w: -w["x0"])
                name = " ".join(
                    fix(w["text"]) for w in words if norm(fix(w["text"])) not in FOOTER
                )
                code = col(300, 370)
                date = col(110, 180)
                recs.append(
                    dict(
                        no=int(serial[0]["text"]),
                        name=re.sub(r"\s+", " ", name).strip(),
                        code=code[0]["text"] if code else "",
                        date=date[0]["text"] if date else "",
                        qty=num(qty[0]["text"]),
                        price=num(prc[0]["text"]),
                        total=num(tot[0]["text"]),
                    )
                )
    recs.sort(key=lambda r: r["no"])
    return recs


# ───────── ٢) جدول مطابقة أسماء المشتريات بأسماء أصناف الجرد ─────────
# (الاسم كما يظهر بتقرير المشتريات) → (اسم الصنف بملف الجرد، درجة الثقة، أساس المطابقة)
PURCHASE_MAP = {
    "برغل": ("برغل", "عالية", "تطابق اسمى + سعر يونيه 21.3641 ≈ سعر الشراء 21.50"),
    "ثوم": ("ثوم", "عالية", "تطابق اسمى تام"),
    "بصل": ("بصل", "عالية", "تطابق اسمى تام"),
    "فلنك": ("فلانك", "عالية", "فرق إملائى فقط (فلنك/فلانك) + نفس نطاق السعر"),
    "مفروم فراخ C": ("مفروم فراخ شعبي", "عالية", "الرتبة C = الدرجة الشعبية (سعر 28 مقابل 34.7 ليونيه)"),
    "مفروم فراخ A": ("مفروم فراخ فاخر", "عالية", "الرتبة A = الدرجة الفاخرة (سعر 55–57 ≈ 58.50 ليونيه)"),
    "فلفل اخضر عام عام": ("فلفل بارد", "متوسطة", "الفلفل الأخضر غير الحار = «البارد»؛ سعر 12–14 ≈ 14.92 ليونيه"),
    "فلفل احمر": ("فلفل رومي حامي", "منخفضة", "البند الوحيد المتبقى من عائلة الفلفل الطازج؛ القيمة 70 ج.م (غير جوهرية)"),
    "فلفل اسمر": ("فلفل اسود", "عالية", "توابل مطحونة سعرها 340–350 ≈ 379.49 ليونيه"),
    "جوزه الطيب": ("جوزه الطيب", "عالية", "تطابق اسمى تام"),
    "سكر": ("سكر", "عالية", "تطابق اسمى تام"),
    "كسبره": ("كزبره", "عالية", "فرق نطق (كسبرة/كزبرة)"),
    "شطه": ("شطه", "عالية", "تطابق اسمى تام"),
    "حبهان": ("حبهان", "عالية", "تطابق اسمى تام"),
    "قرفه": ("قرفه", "عالية", "تطابق اسمى تام"),
    "باربيكا": ("بابريكا", "عالية", "فرق إملائى فقط"),
    "كبابه صينى": ("كبابه", "عالية", "الكبابة الصينى هى نفس بند «كبابه» بالجرد"),
    "قرنفل": ("قرنفل", "عالية", "تطابق اسمى تام"),
    "لورا": ("لورا", "عالية", "تطابق اسمى تام"),
    "ينسون": ("ينسون", "عالية", "تطابق اسمى تام"),
    "صميد": ("صميد", "عالية", "تطابق اسمى تام"),
    "نشا": ("نشا", "عالية", "تطابق اسمى تام"),
    "دقيق ابيض": ("دقيق", "عالية", "بند الدقيق الوحيد بالجرد + سعر 17 ≈ 17.19 ليونيه"),
    "كرسبى اصفر": ("بقسماط اصفر", "عالية", "سعر الشراء 38.00 = سعر يونيه 38.00 بالضبط"),
    "صويا مصرى": ("صويا لحمه", "متوسطة", "الصويا العامة تُحمَّل على «صويا لحمه» (سعر 32–40 ≈ 39.44) — نفس افتراض ملف يونيه"),
    "صويا ايزوليت": ("صويا اوزيليت", "عالية", "فرق إملائى فقط (Isolate)"),
    "فيليه فراخ": ("فيليه", "عالية", "سعر 210 ≈ 217.59 ليونيه"),
    "دوبريه": ("دوبريه", "عالية", "تطابق اسمى تام"),
    "هانك": ("هانك", "عالية", "تطابق اسمى تام"),
    "جلد": ("جلد", "عالية", "تطابق اسمى تام"),
    "دهن": ("دهن بلدي", "عالية", "بند الدهن الوحيد بالجرد"),
    "كرتونه كبده امريكى": ("كبده", "عالية", "بند الكبدة الوحيد بالجرد + سعر 158 ≈ 169.49 ليونيه"),
    "تشكيله هندى امامى": (POOL_KEY, "متوسطة", "لحوم هندى مجمَّعة: بنود (كتف/رقبه/سن/أرباع) بسعر يونيه موحَّد 253"),
    "طعم بيف": ("طعم بيف", "عالية", "تطابق اسمى تام"),
    "طعم شواي": ("طعم شواي", "عالية", "سعر الشراء 295.00 = سعر يونيه 295.00"),
    "طعم ليه": ("طعم ليه", "عالية", "تطابق اسمى تام"),
    "ام اكس": ("MX", "عالية", "الاسم العربى لبند MX"),
    "علب ثلجة اسود": ("علب سوداء", "عالية", "العلب السوداء (ثلجة) — بند واحد بالجرد"),
    "علب شفاف": ("علب بلاستيك شفاف", "متوسطة", "تطابق اسمى؛ رصيد أول المدة كان بلا سعر ⇒ السعر = سعر الشراء"),
    "فرش برانيك": ("فرش برانيك", "عالية", "تطابق اسمى تام"),
    "فرش كفته": ("فرش كفته", "عالية", "تطابق اسمى تام"),
    "فرش برجر": ("فرش البرجر 11*11", "متوسطة", "مقاس 11×11 هو الوحيد ذو رصيد؛ نفس سعر يونيه للمقاسين"),
    "اطباق نص كيلو": ("اطباق 1/2", "عالية", "طبق النصف كيلو = أطباق 1/2"),
    "اطباق3/4 عام عام": ("اطباق 3/4", "عالية", "تطابق اسمى تام"),
    "اكياس تغلبف كفته اسباشيل ص": ("اكياس كفته s.p صغير", "عالية",
                                    "اختبار الحركة: 10.5 + 514.7 − 479.15 = 46.05 استهلاك منطقى"),
    "اكياس تغليف برجر اسباشيل ص": ("اكياس برجر s.p صغير", "عالية",
                                    "اختبار الحركة: 58 + 257.3 − 288.35 = 26.95 استهلاك منطقى"),
    "استربس الهنا 1 كيلو": ("أكياس استربس الهنا ص", "منخفضة",
                            "بند أكياس الاستربس الوحيد للهنا بالجرد؛ القيمة 1,150 ج.م (غير جوهرية)"),
    # بنود مشتراة لا يقابلها صنف بالجرد (لا تدخل التقييم)
    "احمر كارموزين": (None, "—", "لون صناعى (كارموزين) لا يقابله بند بالجرد؛ لم يُدمج مع «لون طبيعي» لاختلاف الصنف"),
    "طبق ممبار": (None, "—", "صنف غير مدرَج بقائمة جرد الخامات"),
    "طبق شفاف كيلو": (None, "—", "لا يمكن الفصل بين «طبق مسطح 1/2 شفاف» و«3/4 شفاف» — تُرك بلا تسعير كما فى يونيه"),
}


# ─────────────────────────── ٣) قراءة الجرد ───────────────────────────
def read_inventory():
    ws = openpyxl.load_workbook(F_JUNE, data_only=True)["Sheet1"]
    june_raw, june_fg = {}, {}
    for r in ws.iter_rows(min_row=5, max_row=153, values_only=True):
        if r[0]:
            june_raw[norm(r[0])] = dict(label=str(r[0]).strip(), qty=num(r[6]) or 0.0,
                                        price=num(r[7]))
        if r[9]:
            june_fg[norm(r[9])] = dict(label=str(r[9]).strip(), qty=num(r[10]) or 0.0,
                                       price=num(r[11]), note=r[14] or r[13] or "")
    ws = openpyxl.load_workbook(F_JULY, data_only=True)["Sheet1"]
    july_raw, july_fg = [], []
    for r in ws.iter_rows(min_row=5, max_row=153, values_only=True):
        if r[0]:
            july_raw.append((str(r[0]).strip(), num(r[6]) or 0.0,
                             [num(r[i]) or 0.0 for i in (1, 2, 3, 4, 5)]))
        if r[7] and norm(r[7]) != "0":
            july_fg.append((str(r[7]).strip(), num(r[8]) or 0.0))
    return june_raw, june_fg, july_raw, july_fg


# ───────────────── ٤) احتساب متوسط السعر المرجّح لكل خامة ─────────────────
def compute_wac(june_raw, purchases):
    """يعيد: {مفتاح الصنف: dict(تفاصيل الاحتساب)} + قائمة البنود غير المطابَقة."""
    buckets, unmapped = {}, []
    for p in purchases:
        key = norm(p["name"])
        hit = None
        for src, (dest, conf, basis) in PURCHASE_MAP.items():
            if norm(src) == key:
                hit = (dest, conf, basis)
                break
        if hit is None:
            unmapped.append(dict(p, reason="اسم غير معرَّف بجدول المطابقة"))
            continue
        dest, conf, basis = hit
        p["dest"], p["conf"], p["basis"] = dest, conf, basis
        if dest is None:
            unmapped.append(dict(p, reason=basis))
            continue
        b = buckets.setdefault(norm(dest), dict(qty=0.0, val=0.0, lines=0))
        b["qty"] += p["qty"]
        b["val"] += p["total"]          # القيمة الفعلية بالفاتورة هى المرجع
        b["lines"] += 1

    # رصيد أول المدة: بنود مجمّع اللحوم تُدمج فى وعاء واحد
    open_q, open_v, open_priced = {}, {}, {}
    for k, v in june_raw.items():
        tgt = norm(POOL_KEY) if v["label"] in MEAT_POOL else k
        open_q[tgt] = open_q.get(tgt, 0.0) + v["qty"]
        if v["price"] is not None and v["price"] > 0:
            open_v[tgt] = open_v.get(tgt, 0.0) + v["qty"] * v["price"]
            open_priced[tgt] = True
        elif v["qty"] > 0:
            open_priced.setdefault(tgt, False)

    out = {}
    for key in set(list(open_q) + list(buckets)):
        oq = open_q.get(key, 0.0)
        ov = open_v.get(key, 0.0)
        pq = buckets.get(key, {}).get("qty", 0.0)
        pv = buckets.get(key, {}).get("val", 0.0)
        priced_open = open_priced.get(key, False)
        if priced_open or oq == 0:
            tq, tv, rule = oq + pq, ov + pv, "قياسى: (أول المدة + المشتريات)"
        else:
            # رصيد أول المدة موجود كميًّا بلا سعر ⇒ يُستبعد من المتوسط حتى لا يُخفَّض بصفر
            tq, tv, rule = pq, pv, "استُبعد رصيد أول المدة من المتوسط (بلا سعر مُثبت)"
        price = tv / tq if tq else None
        out[key] = dict(open_qty=oq, open_val=ov, open_priced=priced_open,
                        buy_qty=pq, buy_val=pv, buy_lines=buckets.get(key, {}).get("lines", 0),
                        tot_qty=tq, tot_val=tv, price=price, rule=rule)
    return out, unmapped


# ─────────── ٥) إعادة تسعير المنتج التام من وصفات التصنيع (BOM) ───────────
def reprice_bom(june_raw, wac):
    """يعيد بناء وصفات يونيه بأسعار يوليو: الربط يتم بمطابقة سعر يونيه المدوَّن بالوصفة."""
    price_index = {}                       # سعر يونيه → مفتاح الصنف بالجرد
    for k, v in june_raw.items():
        if v["price"]:
            price_index.setdefault(round(v["price"], 4), norm(POOL_KEY)
                                   if v["label"] in MEAT_POOL else k)

    def july_price(june_price):
        key = price_index.get(round(june_price, 4))
        if key and wac.get(key, {}).get("price"):
            return wac[key]["price"], key
        return june_price, None            # لا مقابل ⇒ يبقى سعر يونيه

    wb = openpyxl.load_workbook(F_BOM, data_only=True)
    products = []
    for sheet in ["منتجات أبو هاشم", "منتجات سبيشال", "منتجات الهنا"]:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        cur = None
        mode = None
        for i, r in enumerate(rows):
            a = str(r[0]).strip() if r[0] else ""
            b = str(r[1]).strip() if r[1] else ""
            if a and not re.fullmatch(r"[\d.]+", a) and a not in ("م",) and not a.startswith("تسعير"):
                if a.startswith("الوصفة الرئيسية"):
                    mode = "main"
                elif a.startswith("الخلطة الفرعية"):
                    mode = "sub"
                elif a.startswith("تكلفة التعبئة"):
                    mode = "pack"
                else:
                    cur = dict(name=a, brand=sheet, sub=[], main=[], pack=[],
                               june_sub_kg=None, june_kg=None, june_main_val=None,
                               june_main_qty=None)
                    products.append(cur)
                    mode = None
                continue
            if cur is None:
                continue
            if b.startswith("سعر الكجم للخلطة الفرعية"):
                cur["june_sub_kg"] = num(r[3])
            elif b.startswith("إجمالي الوصفة الرئيسية"):
                cur["june_main_qty"], cur["june_main_val"] = num(r[2]), num(r[4])
            elif b.startswith("سعر الكجم (خامات"):
                cur["june_kg"] = num(r[3])
            elif re.fullmatch(r"[\d.]+", a) and b:
                q, pr = num(r[2]), num(r[3])
                if mode in ("sub", "main"):
                    cur[mode].append(dict(name=b, qty=q or 0.0, june_price=pr))
                elif mode == "pack":
                    cur["pack"].append(dict(sku=b, weight=num(r[2]),
                                            june_mix=num(r[3]), packing=num(r[4]),
                                            june_total=num(r[5])))

    for p in products:
        sub_val = 0.0
        for ln in p["sub"]:
            if ln["june_price"] is None:
                ln["july_price"] = None
                ln["july_total"] = 0.0
            else:
                ln["july_price"], ln["src"] = july_price(ln["june_price"])
                ln["july_total"] = ln["qty"] * ln["july_price"]
            sub_val += ln["july_total"]
        sub_qty = sum(l["qty"] for l in p["sub"])
        p["july_sub_kg"] = sub_val / sub_qty if sub_qty else None

        main_val = 0.0
        for ln in p["main"]:
            if norm(ln["name"]) in (norm("توابيل"), norm("توابل")) and p["july_sub_kg"]:
                ln["july_price"] = p["july_sub_kg"]
            elif ln["june_price"] is None:
                ln["july_price"] = None
            else:
                ln["july_price"], ln["src"] = july_price(ln["june_price"])
            ln["july_total"] = ln["qty"] * (ln["july_price"] or 0.0)
            main_val += ln["july_total"]
        p["july_main_val"] = main_val
        # القاسم المستخدم فى يونيه (يتضمن معامل الإنتاجية إن وُجد) يُحفظ كما هو
        divisor = p["june_main_val"] / p["june_kg"] if p["june_kg"] else p["june_main_qty"]
        p["divisor"] = divisor
        p["july_kg"] = main_val / divisor if divisor else None
        for s in p["pack"]:
            s["july_mix"] = s["weight"] * p["july_kg"]
            s["july_total"] = s["july_mix"] + (s["packing"] or 0.0)
    return products


# ───────── ٥-ب) ربط أصناف المنتج التام بالوصفة المناسبة من الاسم ─────────
BRAND_TOKENS = {"ابو هاشم": "أبو هاشم", "ابوهاشم": "أبو هاشم",
                "الهنا": "الهنا", "هنا": "الهنا",
                "اسبيشيال": "سبيشال", "اسبشيال": "سبيشال", "اسبيشال": "سبيشال",
                "sp": "سبيشال"}
TYPE_TOKENS = ["سجق", "كفته", "كفتة", "برجر"]


def resolve_sku(label, june_price, products):
    """يستنتج (المنتج، الوزن، تكلفة التعبئة الفعلية) من اسم الصنف وسعره فى يونيه."""
    n = norm(label)
    brand = next((b for tok, b in BRAND_TOKENS.items() if norm(tok) in n), None)
    typ = next((t for t in TYPE_TOKENS if norm(t) in n), None)
    if not brand or not typ:
        return None
    typ = "كفته" if typ.startswith("كفت") else typ
    prod = next((p for p in products
                 if norm(typ) in norm(p["name"])
                 and p["brand"].replace("منتجات ", "") == brand), None)
    if prod is None:
        return None
    m = re.findall(r"(\d+(?:\.\d+)?)\s*(ك|كيلو)?", n)
    weight = None
    for val, kilo in m:
        v = float(val)
        if kilo or v <= 5:               # «3ك» أو «1.5»
            weight = v
        elif v >= 100:                   # جرامات
            weight = v / 1000.0
        if weight:
            break
    if not weight or weight > 6:
        return None
    if june_price is None:
        # صنف جديد لم يكن مُسعّرًا فى يونيه: تُؤخذ تكلفة التعبئة من جدول الوصفة إن وُجدت
        row = next((s for s in prod["pack"] if abs((s["weight"] or 0) - weight) < 1e-6), None)
        packing = (row["packing"] or 0.0) if row else 0.0
        new = True
    else:
        packing = june_price - weight * prod["june_kg"]
        if packing < -0.02 or packing > 25:   # لا تتسق مع (وزن × سعر الكجم + تعبئة)
            return None
        packing = max(packing, 0.0)
        new = False
    return dict(product=prod["name"], weight=weight, packing=packing, new=new,
                has_pack=bool(packing),
                july=weight * prod["july_kg"] + packing)


# ────────────────────────────── ٦) البناء ──────────────────────────────
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = PatternFill("solid", fgColor="1F4E79")
SUB = PatternFill("solid", fgColor="DDEBF7")
WARN = PatternFill("solid", fgColor="FFF2CC")
TOT = PatternFill("solid", fgColor="E2EFDA")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def finish(ws, widths, freeze="A5", ncols=None):
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze


def main():
    june_raw, june_fg, july_raw, july_fg = read_inventory()
    purchases = read_purchases()
    wac, unmapped = compute_wac(june_raw, purchases)
    products = reprice_bom(june_raw, wac)

    wb = openpyxl.Workbook()

    # ── شيت ١: جرد الخامات آخر يوليو مُسعَّرًا ──
    ws = wb.active
    ws.title = "جرد الخامات آخر يوليو"
    ws["A1"] = "ABOHASHEM FOR MEAT — تقييم مخزون الخامات آخر يوليو 2026 بمتوسط التكلفة المرجّح (WAC)"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = ("متوسط السعر = (قيمة رصيد أول المدة + قيمة مشتريات يوليو) ÷ "
                "(كمية رصيد أول المدة + كمية مشتريات يوليو)")
    ws["A2"].font = Font(size=10, italic=True, color="808080")
    head = ["اسم الصنف", "كمية أول المدة (1/7)", "سعر أول المدة", "قيمة أول المدة",
            "كمية مشتريات يوليو", "قيمة مشتريات يوليو", "إجمالى الكمية", "إجمالى التكلفة",
            "متوسط السعر (WAC) 31/7", "كمية الجرد 31/7", "قيمة المخزون 31/7",
            "التغير عن سعر يونيه %", "ملاحظة"]
    ws.append([])
    ws.append(head)
    style_header(ws, 4, len(head))
    total_val = 0.0
    unpriced = []
    for label, qty, _ in july_raw:
        key = norm(label)
        pooled = label in MEAT_POOL
        wkey = norm(POOL_KEY) if pooled else key
        w = wac.get(wkey)
        j = june_raw.get(key, {})
        jprice = j.get("price") or None          # الصفر بملف يونيه يعنى «بلا سعر» لا تكلفة صفرية
        price = w["price"] if w and w["price"] else jprice
        note = ""
        if pooled:
            note = "سعر موحَّد لمجمّع اللحوم الهندى (كتف/رقبه/سن/أرباع)"
        if w and w["buy_qty"] and not w["open_priced"] and w["open_qty"]:
            note = (note + " — " if note else "") + "رصيد أول المدة بلا سعر مُثبت ⇒ المتوسط من المشتريات فقط"
        if (not w or not w["buy_qty"]) and price:
            note = (note + " — " if note else "") + "لا مشتريات فى يوليو ⇒ السعر ثابت كما فى يونيه"
        # القيمة تُحتسب بدقة العرض (كمية مقرَّبة لخانتين × سعر مقرَّب لأربع) حتى يتطابق
        # حاصل الضرب داخل الملف مع القيمة المعروضة إلى القرش
        val = round(r2(qty) * r4(price), 2) if price else None
        if val:
            total_val += val
        if price is None and qty:
            note = ("صنف بلا سعر مُثبت (لم يكن مُسعّرًا فى يونيه ولا مشتريات له فى يوليو) — "
                    "يلزم إدخال سعر لتقييمه")
            unpriced.append((label, qty))
        chg = ((price / jprice - 1) * 100) if (price and jprice) else None
        ws.append([label,
                   r2(w["open_qty"]) if w else r2(j.get("qty")),
                   r4(jprice), r2((j.get("qty") or 0) * jprice) if jprice else None,
                   r2(w["buy_qty"]) if w else None,
                   r2(w["buy_val"]) if w else None,
                   r2(w["tot_qty"]) if w else None,
                   r2(w["tot_val"]) if w else None,
                   r4(price), r2(qty), r2(val), r2(chg), note])
    last = ws.max_row
    ws.append([])
    ws.append(["إجمالى قيمة مخزون الخامات آخر يوليو 2026 (الأصناف المُسعّرة)", None, None, None,
               None, None, None, None, None, None, r2(total_val), None, None])
    for c in range(1, len(head) + 1):
        ws.cell(row=ws.max_row, column=c).fill = TOT
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, max_row=last, max_col=len(head)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (2, 4, 5, 6, 7, 8, 10, 11):
                cell.number_format = "#,##0.00"
            elif cell.column in (3, 9):
                cell.number_format = "#,##0.0000"
            elif cell.column == 12:
                cell.number_format = "+0.0%;-0.0%"
        if row[8].value is None:
            for cell in row:
                cell.fill = WARN
    for row in ws.iter_rows(min_row=5, max_row=last, min_col=12, max_col=12):
        if row[0].value is not None:
            row[0].value = row[0].value / 100
    finish(ws, [30, 15, 13, 15, 15, 17, 14, 16, 17, 13, 16, 13, 46])

    # ── شيت ٢: خطوات احتساب متوسط السعر ──
    ws = wb.create_sheet("خطوات احتساب المتوسط")
    ws["A1"] = "خطوات احتساب متوسط السعر المرجّح لكل خامة — يوليو 2026"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A2"] = "تُعرض هنا الأصناف التى حدثت عليها مشتريات خلال يوليو فقط (باقى الأصناف يبقى سعرها كما هو)."
    ws["A2"].font = Font(size=10, italic=True, color="808080")
    head2 = ["اسم الصنف", "خطوة ١: كمية أول المدة", "× سعر أول المدة", "= قيمة أول المدة",
             "خطوة ٢: كمية المشتريات", "قيمة المشتريات (بالفواتير)", "متوسط سعر الشراء",
             "خطوة ٣: إجمالى الكمية", "خطوة ٤: إجمالى التكلفة", "خطوة ٥: متوسط السعر",
             "عدد فواتير الشراء", "قاعدة الاحتساب"]
    ws.append([])
    ws.append(head2)
    style_header(ws, 4, len(head2))
    for label, qty, _ in july_raw + [(POOL_KEY, 0, None)]:
        key = norm(POOL_KEY) if label in MEAT_POOL else norm(label)
        if label in MEAT_POOL:
            continue
        w = wac.get(key)
        if not w or not w["buy_qty"]:
            continue
        j = june_raw.get(norm(label), {})
        ws.append([POOL_KEY if label == POOL_KEY else label,
                   r2(w["open_qty"]), r4(j.get("price")), r2(w["open_val"]),
                   r2(w["buy_qty"]), r2(w["buy_val"]),
                   r4(w["buy_val"] / w["buy_qty"]) if w["buy_qty"] else None,
                   r2(w["tot_qty"]), r2(w["tot_val"]), r4(w["price"]),
                   w["buy_lines"], w["rule"]])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=len(head2)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (2, 4, 5, 6, 8, 9):
                cell.number_format = "#,##0.00"
            elif cell.column in (3, 7, 10):
                cell.number_format = "#,##0.0000"
    finish(ws, [30, 16, 14, 16, 16, 18, 15, 15, 17, 15, 12, 42])

    # ── شيت ٣: تفاصيل مشتريات يوليو ──
    ws = wb.create_sheet("مشتريات يوليو التفصيلية")
    ws["A1"] = "مشتريات الخامات — الفترة 01/07/2026 حتى 30/07/2026 (المصدر: تقرير المشتريات فى الفترة)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    head3 = ["م", "التاريخ", "كود الصنف", "اسم الصنف بالمشتريات", "الكمية", "سعر الشراء",
             "الإجمالى", "الصنف المقابل بالجرد", "درجة الثقة", "أساس المطابقة"]
    ws.append([])
    ws.append(head3)
    style_header(ws, 3, len(head3))
    for p in purchases:
        ws.append([p["no"], p["date"], p["code"], p["name"], r2(p["qty"]), r4(p["price"]),
                   r2(p["total"]), p.get("dest") or "— بلا مقابل —",
                   p.get("conf", "—"), p.get("basis", "")])
    ws.append([])
    ws.append([None, None, None, "إجمالى مشتريات يوليو", None, None,
               r2(sum(p["total"] for p in purchases))])
    for c in range(1, len(head3) + 1):
        ws.cell(row=ws.max_row, column=c).fill = TOT
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=len(head3)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (5, 7):
                cell.number_format = "#,##0.00"
            elif cell.column == 6:
                cell.number_format = "#,##0.0000"
    finish(ws, [6, 12, 10, 32, 12, 12, 14, 26, 10, 60], freeze="A4")

    # ── شيت ٤: المنتج التام ──
    ws = wb.create_sheet("جرد المنتج التام آخر يوليو")
    ws["A1"] = "تقييم مخزون المنتج التام آخر يوليو 2026 (وصفات التصنيع بأسعار خامات يوليو + تكلفة التعبئة)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    head4 = ["اسم الصنف", "كمية الجرد 31/7", "سعر يونيه", "سعر يوليو المحدَّث",
             "التغير %", "قيمة المخزون 31/7", "ملاحظة"]
    ws.append([])
    ws.append(head4)
    style_header(ws, 3, len(head4))
    fg_total = 0.0
    for label, qty in july_fg:
        j = june_fg.get(norm(label), {})
        jp = j.get("price") or None
        sku = resolve_sku(label, jp, products)
        price = sku["july"] if sku else jp
        val = round(r2(qty) * r4(price), 2) if price else None
        if val:
            fg_total += val
        chg = ((price / jp - 1) * 100) if (price and jp) else None
        note = str(j.get("note") or "")
        if jp and not sku:
            note = ((note + " | " if note else "")
                    + "خارج نطاق الوصفات التسع — سعر يونيه مُبقًى دون تحديث")
        elif sku:
            pack_txt = (f"+ تعبئة {sku['packing']:.2f} ج.م" if sku["has_pack"]
                        else "بدون تكلفة تعبئة (غير محددة بالوصفة) — السعر الفعلى أعلى قليلًا")
            note = ((note + " | " if note else "")
                    + f"من وصفة «{sku['product']}»: {sku['weight']:g} كجم × سعر الكجم {pack_txt}"
                    + (" | صنف جديد لم يكن مُسعّرًا فى يونيه" if sku["new"] else ""))
        ws.append([label, r2(qty), r4(jp), r4(price), r2(chg / 100) if chg is not None else None,
                   r2(val), note[:200]])
    last4 = ws.max_row
    ws.append([])
    ws.append(["إجمالى قيمة مخزون المنتج التام آخر يوليو 2026 (الأصناف المُسعّرة)", None, None,
               None, None, r2(fg_total), None])
    for c in range(1, len(head4) + 1):
        ws.cell(row=ws.max_row, column=c).fill = TOT
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    ws.append([])
    ws.append(["إجمالى قيمة المخزون الكلى آخر يوليو 2026 (خامات + منتج تام)", None, None,
               None, None, r2(total_val + fg_total), None])
    for c in range(1, len(head4) + 1):
        ws.cell(row=ws.max_row, column=c).fill = SUB
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True, size=12)
    for row in ws.iter_rows(min_row=4, max_row=last4, max_col=len(head4)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (2, 6):
                cell.number_format = "#,##0.00"
            elif cell.column in (3, 4):
                cell.number_format = "#,##0.0000"
            elif cell.column == 5:
                cell.number_format = "+0.0%;-0.0%"
        if row[3].value is None:
            for cell in row:
                cell.fill = WARN
    finish(ws, [30, 14, 14, 17, 10, 16, 60], freeze="A4")

    # ── شيت ٥: سعر الكيلو للمنتجات (من الوصفات) ──
    ws = wb.create_sheet("سعر الكجم من الوصفات")
    ws["A1"] = "أثر تحديث أسعار الخامات على سعر كيلو الخلطة لكل منتج (خامات مباشرة فقط)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    head5 = ["المنتج", "العلامة", "قيمة الوصفة (يونيه)", "قيمة الوصفة (يوليو)",
             "سعر الكجم (يونيه)", "سعر الكجم (يوليو)", "التغير %"]
    ws.append([])
    ws.append(head5)
    style_header(ws, 3, len(head5))
    for p in products:
        chg = (p["july_kg"] / p["june_kg"] - 1) if p["june_kg"] else None
        ws.append([p["name"], p["brand"].replace("منتجات ", ""), r2(p["june_main_val"]),
                   r2(p["july_main_val"]), r4(p["june_kg"]), r4(p["july_kg"]), r2(chg)])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=len(head5)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (3, 4):
                cell.number_format = "#,##0.00"
            elif cell.column in (5, 6):
                cell.number_format = "#,##0.0000"
            elif cell.column == 7:
                cell.number_format = "+0.0%;-0.0%"
    finish(ws, [24, 12, 18, 18, 16, 16, 10], freeze="A4")

    # ── شيت ٦: المنهجية والاستثناءات ──
    ws = wb.create_sheet("المنهجية والاستثناءات")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "المنهجية، الافتراضات، والاستثناءات"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    rows = [
        ("نطاق العمل", "متوسط التكلفة المرجّح (Weighted Average Cost) للخامات فقط — لم تُستخدم أى طريقة "
                        "أخرى (FIFO / LIFO / التكلفة المعيارية)."),
        ("المعادلة", "متوسط السعر = (قيمة أول المدة + قيمة المشتريات) ÷ (كمية أول المدة + كمية المشتريات). "
                     "الطريقة دورية (Periodic WAC): متوسط واحد للشهر كله يُطبَّق على رصيد آخر المدة."),
        ("رصيد أول المدة", "رصيد وأسعار «تسعير جرد اخر شهر ٦ سنة ٢٠٢٦.xlsx» — وهى نفسها WAC متراكم "
                            "(أول أبريل + مشتريات مايو + مشتريات يونيه)، فتكون النتيجة استمرارًا لنفس السلسلة."),
        ("المشتريات", "تقرير المشتريات فى الفترة 01/07/2026–30/07/2026، 100 سطر شراء بإجمالى مطابق "
                      "لإجمالى التقرير. قيمة الشراء المعتمدة هى «الإجمالى» بالفاتورة وليست (كمية × سعر) "
                      "تفاديًا لفروق التقريب فى سعر الوحدة."),
        ("مجمّع اللحوم الهندى", "بنود (لحم كتف / لحم رقبه / لحم سن / لحمه ارباع) سعرها بيونيه موحَّد 253 ج.م، "
                                 "ومشترياتها تُقيَّد باسم واحد «تشكيلة هندى أمامى». عوملت كوعاء تكلفة واحد "
                                 "ومتوسط واحد يُطبَّق على البنود الأربعة."),
        ("أصناف بلا مشتريات", "يبقى سعر يونيه كما هو — لأن المتوسط المرجّح لا يتغير بغياب حركة شراء."),
        ("أصناف برصيد بلا سعر", "إن وُجد رصيد أول مدة بكمية دون سعر مُثبت، يُستبعد من بسط ومقام المتوسط "
                                 "ويُحتسب السعر من المشتريات فقط — بديلًا عن إدخاله بصفر الذى يُخفّض المتوسط زورًا."),
        ("المنتج التام", "أُعيد اشتقاق سعر كل SKU من نفس وصفات التصنيع المستخدمة فى يونيه، باستبدال أسعار "
                          "الخامات فقط، مع الإبقاء على الكميات ومعاملات الإنتاجية وتكاليف التعبئة دون تغيير."),
        ("خارج النطاق", "الأسعار تمثل تكلفة الخامات المباشرة (+ التعبئة للمنتج التام) ولا تشمل الأجور المباشرة "
                         "ولا الأعباء الصناعية غير المباشرة."),
    ]
    r = 3
    for t, d in rows:
        ws.cell(row=r, column=1, value=t).font = Font(bold=True, size=11, color="1F4E79")
        c = ws.cell(row=r, column=2, value=d)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 46
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="بنود مشتراة بلا مقابل فى الجرد (لم تدخل التقييم)").font = Font(
        bold=True, size=12, color="C00000")
    r += 1
    ws.cell(row=r, column=1, value="الصنف").font = Font(bold=True)
    ws.cell(row=r, column=2, value="القيمة (ج.م)").font = Font(bold=True)
    ws.cell(row=r, column=3, value="السبب").font = Font(bold=True)
    agg = {}
    for u in unmapped:
        a = agg.setdefault(u["name"], [0.0, u["reason"]])
        a[0] += u["total"]
    for name, (v, reason) in sorted(agg.items(), key=lambda x: -x[1][0]):
        r += 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=r2(v)).number_format = "#,##0.00"
        ws.cell(row=r, column=3, value=reason).alignment = Alignment(wrap_text=True)
    r += 2
    ws.cell(row=r, column=1, value="أصناف بالجرد آخر يوليو برصيد بلا سعر").font = Font(
        bold=True, size=12, color="C00000")
    for label, qty in unpriced:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=r2(qty)).number_format = "#,##0.00"
    for col, w in zip("ABC", (34, 18, 88)):
        ws.column_dimensions[col].width = w

    wb.save(OUT_XLSX)

    # ── تقرير Markdown ──
    lines = []
    A = lines.append
    A("# متوسط سعر الخامات (WAC) — يوليو 2026 وتقييم مخزون آخر يوليو\n")
    A("**المنشأة:** ABOHASHEM FOR MEAT — Food Industries  ")
    A("**الفترة:** 01/07/2026 → 30/07/2026  ")
    A("**الطريقة:** متوسط التكلفة المرجّح الدورى (Periodic Weighted Average Cost) — دون أى طريقة أخرى.\n")
    A("---\n")
    A("## ١. المعادلة المطبَّقة\n")
    A("```")
    A("إجمالى الكميات  = كمية أول المدة + كمية المشتريات")
    A("إجمالى التكلفة = قيمة أول المدة + قيمة المشتريات")
    A("متوسط السعر    = إجمالى التكلفة ÷ إجمالى الكميات")
    A("قيمة المخزون   = كمية الجرد 31/7 × متوسط السعر")
    A("```\n")
    A("## ٢. الملخص الرقمى\n")
    open_total = sum(v["qty"] * v["price"] for v in june_raw.values()
                     if v["price"] and v["qty"])
    buy_total = sum(p["total"] for p in purchases)
    A("| البند | القيمة (ج.م) |")
    A("|---|---:|")
    A(f"| قيمة مخزون الخامات أول يوليو (رصيد أول المدة المُسعَّر) | {open_total:,.2f} |")
    A(f"| مشتريات يوليو 2026 (100 سطر شراء) | {buy_total:,.2f} |")
    A(f"| **متاح للاستخدام خلال يوليو** | **{open_total + buy_total:,.2f}** |")
    A(f"| قيمة مخزون الخامات آخر يوليو (بالمتوسط الجديد) | {total_val:,.2f} |")
    A(f"| تكلفة الخامات المستنفدة خلال يوليو (بالفرق) | {open_total + buy_total - total_val:,.2f} |")
    A(f"| قيمة مخزون المنتج التام آخر يوليو | {fg_total:,.2f} |")
    A(f"| **إجمالى قيمة المخزون آخر يوليو 2026** | **{total_val + fg_total:,.2f}** |")
    A("")
    A("## ٣. أكبر ١٥ صنفًا من حيث قيمة المخزون آخر يوليو\n")
    rows_val = []
    for label, qty, _ in july_raw:
        key = norm(POOL_KEY) if label in MEAT_POOL else norm(label)
        w = wac.get(key)
        price = w["price"] if w and w["price"] else june_raw.get(norm(label), {}).get("price")
        if price and qty:
            rows_val.append((label, qty, price, qty * price,
                             june_raw.get(norm(label), {}).get("price")))
    rows_val.sort(key=lambda x: -x[3])
    A("| # | الصنف | كمية 31/7 | متوسط السعر يوليو | القيمة | سعر يونيه | التغير |")
    A("|---:|---|---:|---:|---:|---:|---:|")
    for i, (lb, q, pr, v, jp) in enumerate(rows_val[:15], 1):
        ch = f"{(pr/jp-1)*100:+.1f}%" if jp else "—"
        A(f"| {i} | {lb} | {q:,.2f} | {pr:,.4f} | {v:,.2f} | {jp:,.4f} | {ch} |" if jp
          else f"| {i} | {lb} | {q:,.2f} | {pr:,.4f} | {v:,.2f} | — | — |")
    A("")
    A("## ٤. خطوات الاحتساب للأصناف التى تحرّك سعرها\n")
    moved = []
    for label, qty, _ in july_raw:
        key = norm(POOL_KEY) if label in MEAT_POOL else norm(label)
        w = wac.get(key)
        jp = june_raw.get(norm(label), {}).get("price")
        if w and w["buy_qty"] and w["price"] and jp:
            moved.append((abs(w["price"] / jp - 1) * w["tot_val"], label, w, jp))
    moved.sort(key=lambda x: -x[0])
    seen = set()
    for _, label, w, jp in moved[:12]:
        if label in MEAT_POOL:
            if POOL_KEY in seen:
                continue
            seen.add(POOL_KEY)
            label = f"{POOL_KEY} (كتف/رقبه/سن/أرباع)"
        A(f"**{label}**  ")
        A(f"- أول المدة: {w['open_qty']:,.2f} × {jp:,.4f} = {w['open_val']:,.2f} ج.م  ")
        A(f"- المشتريات: {w['buy_qty']:,.2f} بقيمة {w['buy_val']:,.2f} ج.م "
          f"(متوسط شراء {w['buy_val']/w['buy_qty']:,.4f})  ")
        A(f"- إجمالى الكمية = {w['open_qty']:,.2f} + {w['buy_qty']:,.2f} = {w['tot_qty']:,.2f}  ")
        A(f"- إجمالى التكلفة = {w['open_val']:,.2f} + {w['buy_val']:,.2f} = {w['tot_val']:,.2f}  ")
        A(f"- **متوسط السعر = {w['tot_val']:,.2f} ÷ {w['tot_qty']:,.2f} = "
          f"{w['price']:,.4f} ج.م/وحدة** ({(w['price']/jp-1)*100:+.1f}% عن يونيه)\n")
    A("## ٥. أثر التحديث على سعر كيلو الخلطة للمنتجات\n")
    A("| المنتج | العلامة | سعر الكجم يونيه | سعر الكجم يوليو | التغير |")
    A("|---|---|---:|---:|---:|")
    for p in products:
        A(f"| {p['name']} | {p['brand'].replace('منتجات ','')} | {p['june_kg']:,.4f} | "
          f"{p['july_kg']:,.4f} | {(p['july_kg']/p['june_kg']-1)*100:+.1f}% |")
    A("")
    A("## ٦. أسعار المنتج التام بعد التحديث (الأصناف ذات الرصيد)\n")
    A("| الصنف | كمية 31/7 | سعر يونيه | سعر يوليو | التغير | قيمة المخزون |")
    A("|---|---:|---:|---:|---:|---:|")
    for label, qty in july_fg:
        if not qty:
            continue
        j = june_fg.get(norm(label), {})
        jp = j.get("price") or None
        sku = resolve_sku(label, jp, products)
        pr = sku["july"] if sku else jp
        if not pr:
            A(f"| {label} | {qty:,.2f} | — | — | — | بلا سعر |")
            continue
        ch = f"{(pr/jp-1)*100:+.1f}%" if jp else "صنف جديد"
        jps = f"{jp:,.4f}" if jp else "—"
        A(f"| {label} | {qty:,.2f} | {jps} | {pr:,.4f} | {ch} | {qty*pr:,.2f} |")
    A("")
    A("## ٧. الاستثناءات والافتراضات الجوهرية\n")
    A("| البند | القيمة (ج.م) | المعالجة |")
    A("|---|---:|---|")
    for name, (v, reason) in sorted(agg.items(), key=lambda x: -x[1][0]):
        A(f"| {name} | {v:,.2f} | {reason} |")
    A("")
    A("### أصناف برصيد فعلى آخر يوليو بلا سعر مُثبت (خارج التقييم)\n")
    A("هذه الأصناف لم تكن مُسعّرة فى يونيه ولم تُشترَ فى يوليو، فلا يوجد أساس لتسعيرها. "
      "قيمتها = صفر فى الإجمالى أعلاه، ويلزم إدخال آخر سعر شراء لها لاستكمال التقييم:\n")
    A("| الصنف | الكمية آخر يوليو |")
    A("|---|---:|")
    for lb, q in sorted(unpriced, key=lambda x: -x[1]):
        A(f"| {lb} | {q:,.2f} |")
    A("")
    A("**اختبار الاتساق (كمية أول المدة + المشتريات ≥ جرد آخر المدة):** ")
    bad = []
    for label, qty, _ in july_raw:
        key = norm(POOL_KEY) if label in MEAT_POOL else norm(label)
        w = wac.get(key)
        if w and qty - w["tot_qty"] > 0.5 and label not in MEAT_POOL:
            bad.append((label, w["tot_qty"], qty))
    if bad:
        A("رُصدت الحالات التالية حيث يتجاوز رصيد آخر المدة المتاحَ للاستخدام "
          "(يرجَّح أنها إنتاج داخلى أو تحويل بين مخازن وليست شراءً):\n")
        A("| الصنف | المتاح (أول المدة + مشتريات) | جرد 31/7 | الفرق |")
        A("|---|---:|---:|---:|")
        for lb, av, q in sorted(bad, key=lambda x: -(x[2] - x[1]))[:20]:
            A(f"| {lb} | {av:,.2f} | {q:,.2f} | {q-av:+,.2f} |")
    else:
        A("لم تُرصد مخالفات.")
    A("")
    A("---")
    A("> النطاق: تكلفة خامات مباشرة فقط (+ تعبئة للمنتج التام). لا تشمل أجورًا مباشرة "
      "ولا أعباء صناعية غير مباشرة.")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print(f"مشتريات: {len(purchases)} سطر بإجمالى {buy_total:,.2f}")
    print(f"قيمة مخزون الخامات آخر يوليو : {total_val:,.2f}")
    print(f"قيمة المنتج التام آخر يوليو  : {fg_total:,.2f}")
    print(f"الإجمالى                     : {total_val + fg_total:,.2f}")
    print(f"بنود شراء بلا مقابل: {len(agg)} | أصناف بلا سعر برصيد: {len(unpriced)}")
    print("✔", OUT_XLSX.name, "|", OUT_MD.name)


if __name__ == "__main__":
    main()
