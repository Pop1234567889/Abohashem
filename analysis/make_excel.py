# -*- coding: utf-8 -*-
"""Combine the analysis CSV/JSON outputs into one styled Excel workbook."""
import json, os, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, 'output')
XLSX = os.path.join(OUT, 'abu_hashem_sales_debt_analysis.xlsx')

HEAD_FILL = PatternFill('solid', fgColor='C00000')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color='C00000')
THIN = Side(style='thin', color='D0D0D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY_COLS = ('EGP', 'sales', 'debt', 'bonus', 'price', 'arrears', 'monthly', 'remaining', 'paid', 'qty', 'total')

def load_csv(name):
    with open(os.path.join(OUT, name), encoding='utf-8-sig') as f:
        r = csv.reader(f)
        return list(r)

def add_sheet(wb, title, header_caption, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    # caption row
    ws.append([header_caption])
    ws['A1'].font = TITLE_FONT
    if rows:
        ws.append([])  # spacer
        header = rows[0]
        ws.append(header)
        hr = ws.max_row
        for c in range(1, len(header) + 1):
            cell = ws.cell(hr, c)
            cell.fill = HEAD_FILL; cell.font = HEAD_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = BORDER
        for row in rows[1:]:
            ws.append(row)
            r = ws.max_row
            for c, val in enumerate(row, 1):
                cell = ws.cell(r, c)
                cell.border = BORDER
                # numeric formatting
                hdr = header[c - 1].lower()
                try:
                    num = float(str(val).replace(',', ''))
                    cell.value = num
                    if 'pct' in hdr or '%' in hdr:
                        cell.number_format = '0.0'
                    elif any(k in hdr for k in MONEY_COLS):
                        cell.number_format = '#,##0.00'
                except (ValueError, TypeError):
                    pass
        # column widths + autofilter + freeze
        for c in range(1, len(header) + 1):
            letter = get_column_letter(c)
            maxlen = max([len(str(rows[i][c-1])) for i in range(len(rows))] + [10])
            ws.column_dimensions[letter].width = min(max(maxlen + 2, 12), 40)
        ws.freeze_panes = ws.cell(hr + 1, 1)
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(header))}{ws.max_row}"
    return ws

wb = Workbook()
summary = json.load(open(os.path.join(OUT, 'summary.json'), encoding='utf-8'))

# --- Summary sheet (key:value) ---
LABELS = {
    'period_start': 'بداية الفترة / Period start',
    'period_end': 'نهاية الفترة / Period end',
    'period_days': 'عدد الأيام / Days',
    'invoices': 'عدد الفواتير / Invoices',
    'zero_invoices': 'فواتير صفرية / Zero-value invoices',
    'line_items': 'بنود الفواتير / Line items',
    'distinct_items': 'عدد الأصناف / Distinct items',
    'distinct_customers_invoiced': 'عملاء لهم مبيعات / Customers invoiced',
    'salespeople': 'عدد المناديب / Salespeople',
    'total_sales_EGP': 'إجمالى المبيعات (ج) / Total sales',
    'total_qty': 'إجمالى الكميات / Total quantity',
    'total_paid_EGP': 'المحصّل نقدًا (ج) / Collected',
    'total_remaining_EGP': 'المتبقّى آجل (ج) / On credit',
    'total_outstanding_debt_EGP': 'إجمالى المديونية (ج) / Total debt',
}
srows = [['المؤشر / Metric', 'القيمة / Value']]
for k, lab in LABELS.items():
    srows.append([lab, summary.get(k)])
add_sheet(wb, 'الملخص Summary', 'ملخص عام / Overview — أبو هاشم للحوم', srows, first=True)

add_sheet(wb, 'المديونية المستحقة AR',
          'المديونية المستحقة لكل عميل بمعدل ائتمان شهرى (انكشاف بالشهور، بدون فائدة) / '
          'Due indebtedness per customer with monthly credit exposure (months, no interest)',
          load_csv('receivables_by_customer.csv'))
add_sheet(wb, 'اعمار الديون Aging',
          'تقسيم أعمار الديون لمراحل لكل عميل (0-30 / 31-60 / 61-90 / +90 يوم + رصيد مرحّل أقدم) / '
          'A/R aging buckets per customer (0-30 / 31-60 / 61-90 / 90+ days + older carryover)',
          load_csv('aging_by_customer.csv'))
add_sheet(wb, 'الاستحقاق خلال شهر Due1mo',
          'المبالغ المستحقة لكل عميل بأجل شهر واحد (تاريخ الفاتورة + شهر) / '
          'Amounts due per customer under a one-month credit term (invoice date + 1 month)',
          load_csv('due_within_month_by_customer.csv'))
add_sheet(wb, 'المناديب Salespeople',
          'مديونية كل مندوب والانكشاف الائتمانى / Debt & credit exposure per salesperson',
          load_csv('salespeople_summary.csv'))
add_sheet(wb, 'العملاء Customers',
          'مبيعات ومديونية ومكافأة كل عميل / Sales, debt & bonus per customer',
          load_csv('customers_summary.csv'))
add_sheet(wb, 'الأصناف Items',
          'الكمية ومتوسط سعر البيع لكل صنف / Quantity & average selling price per item',
          load_csv('items_summary.csv'))
add_sheet(wb, 'الفواتير الصفرية Zero',
          'الفواتير صفرية القيمة / Zero-value invoices',
          load_csv('zero_invoices.csv'))

wb.save(XLSX)
print('Wrote', XLSX)
print('Sheets:', wb.sheetnames)
